"""Semantic collation, safe document import and local publication exporters.

Document order, not the visual order of tracks, is authoritative. Optional
converters report their availability and never manufacture a successful result.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import html
import io
import json
import mimetypes
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from urllib.parse import unquote, urlparse
import uuid
import zipfile
import xml.etree.ElementTree as ET
from .page_geometry import page_geometry
from .page_numbering import validate_numbering, format_page_number, section_numbering, read_docx_numbering

MAX_IMPORT_BYTES = 100 * 1024 * 1024
DOCUMENT_FORMAT = json.loads(Path(__file__).with_name('document-format.json').read_text(encoding='utf-8'))
MAX_IMAGE_BYTES = 20 * 1024 * 1024
SUPPORTED_FORMATS = ("txt", "md", "html", "docx", "pdf", "epub", "azw3")
MIMES = {"txt": "text/plain", "md": "text/markdown", "html": "text/html",
         "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
         "pdf": "application/pdf", "epub": "application/epub+zip", "azw3": "application/vnd.amazon.mobi8-ebook"}
ALIASES = {"bulletList": "bullet_list", "orderedList": "ordered_list", "listItem": "list_item",
           "hardBreak": "hard_break", "horizontalRule": "horizontal_rule", "codeBlock": "code_block",
           "tableRow": "table_row", "tableCell": "table_cell", "tableHeader": "table_header", "pageBreak": "page_break"}


def _kind(node):
    return ALIASES.get(node.get("type"), node.get("type"))


def _warn(warnings, message):
    if message not in warnings:
        warnings.append(message)


def plain_text(node: dict) -> str:
    kind = _kind(node)
    if kind == "text":
        return node.get("text", "")
    if kind == "hard_break":
        return "\n"
    if kind in ("image", "inline_image"):
        return ""
    separator = "\n" if kind in ("doc", "blockquote", "bullet_list", "ordered_list", "list_item", "table", "table_cell", "table_header") else "\t" if kind == "table_row" else ""
    return separator.join(plain_text(child) for child in node.get("content", []))


def text_document(text: str) -> dict:
    return {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": part}] if part else []}
                                        for part in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]}


def _order(obj, fallback):
    value = obj.get("order", fallback)
    return value if isinstance(value, (int, float)) else fallback


def project_document(project: dict) -> list[dict]:
    """Resolve included placements, frozen snapshots and accepted variants.

    A project with no placements has an empty collation. Draft drafts are never
    silently added to an output. Muted/solo track states affect audition only.
    """
    from .book import publication_project
    project = publication_project(project)
    clips = {c["id"]: c for c in project.get("clips", [])}
    sections = sorted(enumerate(project.get("sections", [])), key=lambda x: (_order(x[1], x[0]), x[0]))
    result = []
    placements = project.get("placements", [])
    known_sections = {s["id"] for _, s in sections}
    for p in placements:
        if p.get("include", True) and p.get("sectionId") not in known_sections:
            raise ValueError(f"Placement {p.get('id', '?')} refers to a missing section.")
    for _, section in sections:
        selected = [(i, p) for i, p in enumerate(placements) if p.get("sectionId") == section["id"] and p.get("include", True)]
        selected.sort(key=lambda x: (_order(x[1], x[0]), x[0]))
        blocks, ids = [], []
        for _, placement in selected:
            clip = clips.get(placement.get("clipId"))
            if clip is None:
                raise ValueError(f"Placement {placement.get('id', '?')} refers to a missing clip.")
            if placement.get("frozenDocument") is not None:
                document = placement["frozenDocument"]
            elif placement.get("frozenText") is not None:
                document = text_document(placement["frozenText"])
            else:
                variant_id = placement.get("variantId", clip.get("activeVariantId"))
                if variant_id:
                    variant = next((v for v in clip.get("variants", []) if v["id"] == variant_id), None)
                    if variant is None:
                        raise ValueError(f"Clip {clip.get('title', clip['id'])} has a missing chosen variant.")
                    document = variant.get("document") or text_document(variant.get("text", ""))
                else:
                    document = clip.get("document") or text_document(clip.get("text", ""))
            if not isinstance(document, dict) or document.get("type") != "doc":
                raise ValueError("A publication draft must contain a structured document.")
            blocks.extend(copy.deepcopy(document.get("content", [])))
            ids.append(clip["id"])
        result.append({"id": section["id"], "title": section.get("title", ""), "role": section.get("role", "chapter"),
                       "blocks": blocks, "text": plain_text({"type": "doc", "content": blocks}), "clipIds": ids,
                       **({"rawMarkdown":clips[ids[0]]["rawSource"]["text"]} if len(selected)==1 and not selected[0][1].get("frozenDocument") and not selected[0][1].get("frozenText") and not selected[0][1].get("variantId") and not clips[ids[0]].get("activeVariantId") and clips[ids[0]].get("rawSource",{}).get("format")=="markdown" and not clips[ids[0]]["rawSource"].get("error") else {})})
    return result


def _number(value, default, low, high):
    try:
        result = float(value)
        return min(high, max(low, result)) if result == result else default
    except (TypeError, ValueError):
        return default


def _publication_sections(project, warnings):
    from .document_bookmarks import publication_bookmarks
    sections = project_document(project)
    from .document_fields import refresh_fields
    refresh_fields([{'type':'doc','content':section['blocks']} for section in sections],project,warnings)
    from .running_matter import project_running_stories, stories
    running = list(project_running_stories(project))
    for section in sections:
        for node in section['blocks']:
            if node['type']=='section_break': running.extend(stories(node.get('attrs',{}).get('runningMatter')))
    refresh_fields(running,project,warnings)
    from .running_headings import refresh_body_headings
    refresh_body_headings(sections,warnings)
    from .document_contents import refresh_contents
    refresh_contents([{'type':'doc','content':section['blocks']} for section in sections], {s['id']:s for s in project.get('styles', [])})
    from .document_references import publication_references
    publication_references(sections, warnings)
    publication_bookmarks(sections, warnings)
    from .heading_numbering import number_headings
    number_headings(sections)
    from .outline_lists import paint_outline_labels
    for section in sections:paint_outline_labels({'type':'doc','content':section['blocks']})
    return sections


def _prepare_publication(project, options=None, warnings=None):
    """Resolve styles and optional generated matter on an isolated snapshot."""
    warnings = warnings if warnings is not None else []
    from .book import publication_project
    prepared = copy.deepcopy(publication_project(project))
    view = (options or {}).get('reviewView', 'markup')
    if view not in ('final', 'original', 'markup'): raise ValueError('Unsupported review publication view.')
    if view != 'markup':
        from .document_revisions import review_document
        from .document_comments import walk as review_walk
        def projected(document):
            if any(n.get('type') == 'revision_break' or n.get('attrs', {}).get('revisionProperties') or any(m.get('type', '').startswith('revision_') for m in n.get('marks', [])) for n in review_walk(document)):
                _warn(warnings, f'This publication uses the {view} review view. The authored Alder document retains its pending changes.')
            return review_document(document, view)
        for clip in prepared.get('clips', []):
            clip['document'] = projected(clip['document'])
            for variant in clip.get('variants', []): variant['document'] = projected(variant['document'])
        for placement in prepared.get('placements', []):
            if placement.get('frozenDocument'): placement['frozenDocument'] = projected(placement['frozenDocument'])
    styles = {s["id"]: s for s in prepared.get("styles", [])}
    resolved, visiting = {}, set()

    def meaningful(values):
        return {key: value for key, value in values.items() if value is not None and value != ""}

    def canonical_color(value):
        from PIL import ImageColor
        try:
            red, green, blue, alpha = ImageColor.getcolor(str(value), "RGBA")
        except (ValueError, TypeError) as exc:
            raise ValueError(f"Unsupported publication colour: {value!r}.") from exc
        if alpha != 255:
            _warn(warnings, "Translucent text colours were flattened against white for consistent document output.")
            red, green, blue = (round(channel * alpha / 255 + 255 - alpha) for channel in (red, green, blue))
        return f"#{red:02x}{green:02x}{blue:02x}"

    def resolve(style_id):
        if style_id in resolved:
            return resolved[style_id]
        if style_id in visiting:
            raise ValueError(f"Style inheritance contains a cycle at '{style_id}'.")
        if style_id not in styles:
            raise ValueError(f"Publication refers to missing style '{style_id}'.")
        visiting.add(style_id)
        style = styles[style_id]
        base = resolve(style["basedOn"]) if style.get("basedOn") else {}
        resolved[style_id] = {**base, **meaningful(style), "id": style_id, "kind": style.get("kind") or "paragraph"}
        if resolved[style_id].get("color"):
            resolved[style_id]["color"] = canonical_color(resolved[style_id]["color"])
        visiting.remove(style_id)
        return resolved[style_id]

    for style_id in styles:
        resolve(style_id)
    prepared["styles"] = list(resolved.values())
    metadata = {"id", "name", "kind", "basedOn", "nextStyle"}

    def apply(attrs):
        attrs = attrs or {}
        base = {key: value for key, value in resolve(attrs["styleId"]).items() if key not in metadata} if attrs.get("styleId") else {}
        result = {**base, **meaningful(attrs)}
        if result.get("color"):
            result["color"] = canonical_color(result["color"])
        return result

    def walk(node):
        if not isinstance(node, dict):
            return
        if (options or {}).get('includeReviewComments', prepared.get('settings', {}).get('includeReviewComments', True)) is False:
            if 'content' in node:
                node['content'] = [child for child in node['content'] if child.get('type') != 'comment_anchor']
        if node.get("attrs"):
            node["attrs"] = apply(node["attrs"])
        if node.get('type') == 'revision_break' and node['attrs'].get('kind') == 'deletion':
            original = {'type': node['attrs']['rightType'], 'attrs': node['attrs']['rightAttrs']}
            walk(original)
            node['attrs']['rightAttrs'] = original['attrs']
        if _kind(node) == "heading":
            attrs = node.get("attrs") or {}
            level = int(_number(attrs.get("level"), 1, 1, 6))
            node["attrs"] = {**DOCUMENT_FORMAT['heading'], 'fontSize': DOCUMENT_FORMAT['headingSizes'][level-1], **attrs}
        if node.get('attrs', {}).get('revisionProperties'):
            log = node['attrs']['revisionProperties']
            walk(log['base'])
            for record in log['records']:
                walk(record['before']); walk(record['after'])
        for mark in node.get("marks", []):
            if mark.get("type") in ("text_style", "textStyle", "fontFamily", "fontSize", "color"):
                mark["attrs"] = apply(mark.get("attrs"))
        for child in node.get("content", []):
            walk(child)

    from .list_formatting import resolve_list_orders
    for clip in prepared.get("clips", []):
        walk(clip.get("document"))
        if clip.get("document"): resolve_list_orders(clip["document"])
        for variant in clip.get("variants", []):
            walk(variant.get("document"))
            if variant.get("document"): resolve_list_orders(variant["document"])
    for placement in prepared.get("placements", []):
        walk(placement.get("frozenDocument"))
        if placement.get("frozenDocument"): resolve_list_orders(placement["frozenDocument"])
    from .running_matter import project_running_stories
    for story in project_running_stories(prepared): walk(story)

    settings = _settings(prepared, options)
    if settings.get("includeGlossary"):
        entries = [entry for entry in prepared.get("dictionary", [])
                   if str(entry.get("word") or "").strip() and str(entry.get("definition") or "").strip()]
        omitted = len(prepared.get("dictionary", [])) - len(entries)
        if omitted:
            _warn(warnings, f"Glossary omitted {omitted} dictionary entry/entries without both a word and an authored definition.")
        if not entries:
            _warn(warnings, "No authored dictionary definitions were available for the requested glossary.")
        else:
            blocks = []
            for entry in sorted(entries, key=lambda item: str(item["word"]).casefold()):
                blocks.append({"type": "heading", "attrs": {"level": 3}, "content": [{"type": "text", "text": str(entry["word"]).strip()}]})
                details = " · ".join(str(entry.get(key) or "").strip() for key in ("ipa", "partOfSpeech") if entry.get(key))
                if not details and entry.get("pos"):
                    details = str(entry["pos"])
                if details:
                    blocks.append({"type": "paragraph", "content": [{"type": "text", "text": details, "marks": [{"type": "em"}]}]})
                blocks.extend(text_document(str(entry["definition"]).strip())["content"])
            occupied = {item["id"] for collection in ("sections", "clips", "placements") for item in prepared.get(collection, [])}
            identifier = "alder-generated-glossary"
            while identifier in occupied:
                identifier += "-output"
            prepared.setdefault("sections", []).append({"id": identifier, "title": str(settings.get("glossaryTitle") or "Glossary"), "role": "glossary",
                "order": max([_order(s, i) for i, s in enumerate(prepared.get("sections", []))] or [0]) + 1})
            prepared.setdefault("clips", []).append({"id": identifier, "title": "Generated glossary", "document": {"type": "doc", "content": blocks}})
            prepared.setdefault("placements", []).append({"id": identifier, "clipId": identifier, "sectionId": identifier, "include": True, "order": 0})
    return prepared


def _settings(project, options=None):
    values = {"author": "", "description": "", "pageSize": "A4", "marginMm": 22, "fontFamily": "Cambria",
              "fontSize": 12, "lineHeight": 1.6, "header": "", "footer": True, "includeTitle": True, "includeToc": False, "pageNumberFormat": "decimal"}
    values.update(project.get("settings", {}))
    values.update(options or {})
    values["fontSize"] = _number(values.get("fontSize"), 12, 6, 72)
    values["lineHeight"] = _number(values.get("lineHeight"), 1.6, 0.8, 4)
    values["marginMm"] = _number(values.get("marginMm"), 22, 0, 100)
    values["fontFamily"] = re.sub(r"[^\w ,'-]", "", str(values.get("fontFamily", "Cambria")))[:100] or "Cambria"
    values["pageSize"] = values.get("pageSize") if values.get("pageSize") in ("A4", "A5", "Letter", "Legal", "6x9", "Custom") else "A4"
    values["orientation"] = "landscape" if values.get("orientation") == "landscape" else "portrait"
    values["firstPageNumber"] = int(_number(values.get("firstPageNumber"), 1, 1, 9999))
    values['pageNumberFormat'] = values.get('pageNumberFormat') or 'decimal'
    validate_numbering(values['firstPageNumber'], values['pageNumberFormat'])
    return values


def _safe_link(value):
    value = str(value or "").strip()
    if re.search(r"[\x00-\x20]", value):
        return ""
    parsed = urlparse(value)
    return value if parsed.scheme.lower() in ("http", "https", "mailto") or value.startswith("#") else ""


def _xml(data):
    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)", data, re.I):
        raise ValueError("External entities and document types are not accepted in imported XML.")
    return ET.fromstring(data)


def _zip(path):
    archive = zipfile.ZipFile(path)
    infos = archive.infolist()
    if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_IMPORT_BYTES:
        archive.close()
        raise ValueError("This document exceeds the 100 MB expanded import limit.")
    names = [i.filename for i in infos]
    if len(set(names)) != len(names):
        archive.close()
        raise ValueError("Document archive contains duplicate entries.")
    for name in names:
        if PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts or "\\" in name:
            archive.close()
            raise ValueError("Document archive contains an unsafe resource path.")
    return archive


class _Resources:
    def __init__(self, project, options, warnings):
        self.bookmark_names = set()
        self.project = project
        self.options = options or {}
        self.warnings = warnings
        self.images = {}
        self._image_sources = {}

    def image(self, attrs, apply_crop=True):
        """Read only embedded raster data or assets inside the controlled root."""
        source = attrs.get("src", "") or ""
        asset_id = attrs.get("assetId") or (source[6:] if source.startswith("asset:") else None)
        cache_key = (source, asset_id, json.dumps(attrs.get('crop'), sort_keys=True) if apply_crop else None)
        if cache_key in self._image_sources:
            return self._image_sources[cache_key]
        asset = next((a for a in self.project.get("assets", []) if a.get("id") == asset_id), None)
        try:
            if asset and asset.get("data"):
                data = base64.b64decode(asset["data"], validate=True)
            elif source.startswith("data:image/"):
                if len(source) > MAX_IMAGE_BYTES * 1.4:
                    raise ValueError("image exceeds size limit")
                data = base64.b64decode(source.split(",", 1)[1], validate=True)
            elif asset and asset.get("path") and self.options.get("assetRoot"):
                root = Path(self.options["assetRoot"]).resolve()
                candidate = Path(asset["path"])
                path = (candidate if candidate.is_absolute() else root / candidate).resolve()
                if not path.is_relative_to(root):
                    raise ValueError("asset path is outside the project asset store")
                if path.stat().st_size > MAX_IMAGE_BYTES:
                    raise ValueError("image exceeds size limit")
                data = path.read_bytes()
            else:
                raise ValueError("image is not an embedded or collected project asset")
            if len(data) > MAX_IMAGE_BYTES:
                raise ValueError("image exceeds size limit")
            from PIL import Image
            with Image.open(io.BytesIO(data)) as image:
                if image.width * image.height > 40_000_000:
                    raise ValueError("image exceeds 40 megapixels")
                if apply_crop and attrs.get('crop'):
                    from .image_crop import cropped_raster
                    image = cropped_raster(image, attrs['crop'])
                # Decode to a known raster format. This also strips active data and EXIF.
                converted = io.BytesIO()
                image.convert("RGBA" if "A" in image.getbands() else "RGB").save(converted, format="PNG")
                data = converted.getvalue()
                width, height = image.size
            key = hashlib.sha256(data).hexdigest()[:24]
            info = {"name": f"image-{key}.png", "mime": "image/png", "data": data, "width": width, "height": height}
            self.images[key] = info
            self._image_sources[cache_key] = info
            return info
        except Exception as exc:
            _warn(self.warnings, f"Image {attrs.get('alt') or asset_id or source[:60] or '(unnamed)'} omitted: {exc}.")
            return None


def _block_style(node, project):
    attrs = node.get("attrs", {})
    base = next((s for s in project.get("styles", []) if s.get("id") == attrs.get("styleId")), {})
    values = {**base, **{k: v for k, v in attrs.items() if v is not None}}
    css = []
    from .paragraph_appearance import appearance_css
    css.append(appearance_css(values))
    alignment = values.get("align") or values.get("textAlign")
    if alignment in ("left", "center", "right", "justify"):
        css.append("text-align:" + alignment)
    if values.get("fontSize"):
        css.append(f"font-size:{_number(values['fontSize'], 12, 6, 72):g}pt")
    if values.get("fontFamily"):
        css.append("font-family:" + re.sub(r"[^\w ,'-]", "", str(values["fontFamily"]))[:100])
    if values.get("lineHeight"):
        css.append(f"line-height:{_number(values['lineHeight'], 1.6, 0.8, 4):g}")
    if values.get("lineSpacingMode") in ("exact", "atLeast") and values.get("lineSpacing") is not None:
        leading = f"{_number(values['lineSpacing'], 12, 1, 144):g}pt"
        leading = leading if values["lineSpacingMode"] == "exact" else f"max(1.2em,{leading})"
        css.append('line-height:' + leading + ';--alder-paragraph-leading:' + leading)
    elif values.get('lineSpacingMode') == 'multiple':
        leading = f"{_number(values.get('lineHeight'), project.get('settings', {}).get('lineHeight', 1.6), 0.8, 4):g}"
        css.append('line-height:' + leading + ';--alder-paragraph-leading:' + leading)
    for field, prop in (("spaceAfter", "margin-bottom"), ("spaceBefore", "margin-top"), ("leftIndent", "margin-left"), ("rightIndent", "margin-right"), ("firstLineIndent" if values.get("firstLineIndent") is not None else "indent", "text-indent")):
        if values.get(field) is not None:
            css.append(f"{prop}:{_number(values[field], 0, -144 if field in ('indent', 'firstLineIndent', 'leftIndent', 'rightIndent') else 0, 144):g}pt")
    for key, prop, yes in (("keepWithNext", "break-after", "avoid"), ("keepTogether", "break-inside", "avoid"), ("pageBreakBefore", "break-before", "page")):
        if type(values.get(key)) is bool:
            css.append(f"{prop}:{yes if values[key] else 'auto'}")
    if type(values.get("widowOrphan")) is bool:
        css.append("widows:2;orphans:2" if values['widowOrphan'] else "widows:1;orphans:1")
    if re.fullmatch(r"#[0-9a-fA-F]{3}(?:[0-9a-fA-F]{3})?", str(values.get("color") or "")):
        css.append("color:" + values["color"])
    return ";".join(css)


def _html_node(node, project, resources, warnings, epub=False):
    kind = _kind(node)
    attrs = node.get("attrs", {})
    reviews = [mark for mark in node.get('marks', []) if mark['type'].startswith('revision_')]
    if reviews:
        from .revision_formatting import format_marks
        formatting = next((m for m in reviews if m['type'] == 'revision_format'), None)
        marks = format_marks(formatting['attrs']) if formatting else [m for m in node.get('marks', []) if not m['type'].startswith('revision_')]
        value = _html_node({**node, 'marks': marks}, project, resources, warnings, epub)
        for mark in reviews:
            if mark['type'] == 'revision_format':
                value = '<span data-revision-format="' + html.escape(json.dumps(mark['attrs'], ensure_ascii=False), quote=True) + '">' + value + '</span>'
            else:
                stamp = mark['attrs']; change = mark['type'].removeprefix('revision_')
                value = '<span data-revision-' + change + '="' + html.escape(stamp['id'], quote=True) + '" data-revision-author="' + html.escape(stamp.get('author', ''), quote=True) + '" data-revision-date="' + html.escape(stamp.get('date', ''), quote=True) + '">' + value + '</span>'
        return value
    if kind == 'revision_break':
        return '<span data-revision-break="' + attrs['kind'] + '" data-revision-break-content="' + html.escape(json.dumps(attrs, ensure_ascii=False), quote=True) + '"></span>'
    if kind == "text":
        text = html.escape(node.get("text", ""))
        for mark in node.get("marks", []):
            name, a = mark.get("type"), mark.get("attrs", {})
            tag = {"strong": "strong", "bold": "strong", "em": "em", "italic": "em", "code": "code", "underline": "u", "strike": "s", "s": "s", "subscript": "sub", "superscript": "sup", "highlight": "mark"}.get(name)
            if tag:
                text = f"<{tag}>{text}</{tag}>"
            elif name == "link":
                link = _safe_link(a.get("href"))
                if link:
                    if epub and link.startswith('#'):
                        link = resources.bookmark_files.get(link[1:], '') + link
                    text = f'<a href="{html.escape(link, quote=True)}">{text}</a>'
                else:
                    _warn(warnings, "Unsafe or unsupported hyperlinks were removed.")
            elif name in ("textStyle", "text_style", "fontFamily", "fontSize", "color"):
                style = _block_style({"attrs": a}, project)
                color = a.get("color") or ""
                if re.fullmatch(r"#[0-9a-fA-F]{3}(?:[0-9a-fA-F]{3})?", color):
                    style += ";color:" + color
                if style:
                    text = f'<span style="{html.escape(style, quote=True)}">{text}</span>'
            else:
                _warn(warnings, f"Unsupported text mark '{name}' was omitted; its text was retained.")
        return text
    if kind == 'comment_anchor':
        return '<span hidden="hidden" data-comment-thread="' + html.escape(attrs['threadId'], quote=True) + '" data-comment-edge="' + attrs['edge'] + '" data-comment-content="' + html.escape(json.dumps(attrs['thread'], ensure_ascii=False), quote=True) + '"></span>'
    if kind == 'document_field':
        value = ''.join(_html_node({**child,'marks':node.get('marks',[])+child.get('marks',[])},project,resources,warnings,epub) for child in node.get('content',[]))
        return f'<span data-document-field="{attrs["fieldKind"]}" data-field-unresolved="{str(bool(attrs.get("unresolved"))).lower()}">{value}</span>'
    if kind == 'cross_reference':
        target = attrs.get('targetName') or attrs['targetId']
        value = ''.join(_html_node({**child, 'marks':node.get('marks', []) + child.get('marks', [])}, project, resources, warnings, epub) for child in node.get('content', []))
        unresolved = ' data-reference-unresolved="true"' if attrs.get('unresolved') else ''
        format = html.escape(attrs.get('referenceFormat', 'text'), quote=True)
        return f'<span data-cross-reference="{html.escape(target, quote=True)}" data-reference-format="{format}"{unresolved}>{value}</span>'
    if kind == 'sequence':
        value=''.join(_html_node({**child,'marks':node.get('marks',[])+child.get('marks',[])},project,resources,warnings,epub) for child in node.get('content',[]))
        return f'<span data-document-sequence="{attrs["sequenceName"]}">{value}</span>'
    if kind == "bookmark":
        from .document_bookmarks import bookmark_name
        name = bookmark_name(attrs["name"], resources.bookmark_names, warnings)
        marker = ' data-bookmark-range="true"' if attrs.get('range') else ''
        return f'<span id="{name}" data-document-bookmark="{name}" data-bookmark-id="{name}"{marker}></span>'
    if kind == "bookmark_end":
        if attrs.get('unresolved'): return ''
        name = html.escape(attrs.get('targetName', ''), quote=True)
        return f'<span data-bookmark-end="{name}" data-bookmark-name="{name}"></span>'
    if kind == "hard_break":
        return "<br/>"
    if kind == "horizontal_rule":
        return "<hr/>"
    if kind == "section_break":
        if epub:
            _warn(warnings, "Reflowable EPUB retains section boundaries but does not apply fixed paper dimensions or odd/even page starts.")
        from .sections import validate_section
        validate_section(attrs)
        from .running_matter import portable_matter
        return '<div class="section-break" data-section="' + html.escape(json.dumps({'kind': attrs.get('kind', 'nextPage'), 'page': attrs.get('page', {}), 'numberingStart': attrs.get('numberingStart'), 'numberingFormat': attrs.get('numberingFormat'), 'runningMatter':portable_matter(attrs.get('runningMatter'))}), quote=True) + '"></div>'
    if kind == "page_break":
        return '<div class="page-break"></div>'
    if kind in ("image", "inline_image"):
        image = resources.image(attrs)
        if not image:
            return '<span class="missing-image">' + html.escape(attrs.get("alt") or "[Image unavailable]") + "</span>"
        source = "images/" + image["name"] if epub else "data:image/png;base64," + base64.b64encode(image["data"]).decode()
        width = f' width="{_number(attrs.get("width"), image["width"], 1, 10000):g}"' if attrs.get("width") else ""
        alignment = attrs.get('align') if attrs.get('align') in ('left', 'center', 'right') else 'left'
        title = html.escape(attrs.get('title') or '', quote=True)
        alternative = html.escape(attrs.get('alt') or '', quote=True)
        decoration = f' role="presentation" data-image-decorative="true" data-image-alt="{alternative}" data-image-title="{title}"' if attrs.get('decorative') else ''
        if attrs.get('decorative'): alternative = ''; title = ''
        placement = 'inline' if kind == 'inline_image' else 'block'
        crop_metadata = ''
        if attrs.get('crop'):
            original = resources.image(attrs, apply_crop=False)
            original_source = 'images/' + original['name'] if epub else 'data:image/png;base64,' + base64.b64encode(original['data']).decode()
            crop_metadata = ' data-image-original="' + original_source + '" data-image-crop="' + html.escape(json.dumps(attrs['crop']), quote=True) + '"'
        identity = ' data-object-id="' + html.escape(attrs['blockId'], quote=True) + '"' if attrs.get('blockId') else ''
        tag = f'<img src="{source}" alt="{alternative}" title="{title}" data-image-align="{alignment}" data-image-placement="{placement}"{identity}{crop_metadata}{decoration}{width}/>'
        return tag if kind == 'inline_image' else f'<figure style="text-align:{alignment};margin:12px 0">{tag}</figure>'
    inner = "".join(_html_node(child, project, resources, warnings, epub) for child in node.get("content", []))
    if kind == 'table_of_contents':
        metadata = f' data-contents-title="{html.escape(attrs.get("title", "Contents"), quote=True)}" data-contents-depth="{attrs.get("depth", 3)}"'
        metadata += ' data-contents-pages="' + ('true' if attrs.get('pageNumbers',True) else 'false') + '"'
        metadata += ' data-contents-kind="' + html.escape(attrs.get('kind','headings'), quote=True) + '"'
        if attrs.get('entryStyle'): metadata += ' data-contents-style="' + html.escape(attrs['entryStyle'], quote=True) + '"'
        return '<div data-document-contents="true"' + metadata + '>' + inner + '</div>'
    tags = {"paragraph": "p", "blockquote": "blockquote", "bullet_list": "ul", "ordered_list": "ol", "list_item": "li",
            "table": "table", "table_row": "tr", "table_cell": "td", "table_header": "th", "code_block": "pre"}
    if kind == "doc":
        return inner
    tag = "h" + str(int(_number(attrs.get("level"), 2, 1, 6))) if kind == "heading" else tags.get(kind)
    if not tag:
        _warn(warnings, f"Unsupported block '{kind}' was flattened; its text was retained.")
        return inner
    extra = ""
    if attrs.get('revisionProperties'): extra += ' data-revision-paragraph="' + html.escape(json.dumps(attrs['revisionProperties'], ensure_ascii=False), quote=True) + '"'
    if attrs.get('blockId'): extra += ' data-block-id="' + html.escape(attrs['blockId'], quote=True) + '"'
    style = _block_style(node, project)
    if kind in ('paragraph', 'heading', 'code_block'):
        if attrs.get('headingNumbered') is not None: extra += ' data-heading-numbered="' + str(attrs['headingNumbered']).lower() + '"'
        if attrs.get('headingStart') is not None: extra += ' data-heading-start="' + str(attrs['headingStart']) + '"'
        if attrs.get('_headingLabel'): inner = '<span data-heading-marker="true">' + html.escape(attrs['_headingLabel']) + '&#160;</span>' + inner
        if attrs.get('_outlineLabel') is not None:
            inner = '<span data-outline-number="true" style="display:inline-block;white-space:pre;margin-left:-'+str(attrs['_outlineHanging'])+'pt">'+html.escape(attrs['_outlineLabel'])+'&#160;</span>'+inner
        stops = attrs.get('tabStops')
        if stops is not None:
            from .tab_stops import validate_tab_stops
            validate_tab_stops(stops)
            extra += ' data-tab-stops="' + html.escape(json.dumps(stops), quote=True) + '"'
            if stops:
                _warn(warnings, "HTML preserves custom tab definitions for Alder but other browsers use uniform tab spacing.")
        if '\t' in plain_text(node): style += ';white-space:pre-wrap;tab-size:36pt'
    if kind in ('table', 'table_row', 'table_cell', 'table_header'):
        from .table_formatting import table_css
        style += ';' + table_css(kind, attrs)
        if kind == 'table': extra += ' data-table-align="' + attrs.get('align', 'left') + '"'
    if kind == 'table' and attrs.get('blockId'): extra += ' data-object-id="' + html.escape(attrs['blockId'], quote=True) + '"'
    if attrs.get('captionLabel'):
        extra += ' data-caption-label="' + attrs['captionLabel'] + '"'
        if attrs.get('captionTargetId'): extra += ' data-caption-target="' + html.escape(attrs['captionTargetId'], quote=True) + '"'
    if kind in ('ordered_list', 'bullet_list'):
        from .list_formatting import NUMBER_FORMATS, BULLETS
        if attrs.get('listStyle') in (NUMBER_FORMATS if kind == 'ordered_list' else BULLETS):
            style += ';list-style-type:' + attrs['listStyle']
        if attrs.get('listTemplate') is not None:extra += ' data-list-template="'+html.escape(json.dumps(attrs['listTemplate']),quote=True)+'"'
        if attrs.get('_outlineIndent') is not None:style+=';list-style-type:none;padding-left:0;margin-left:'+str(attrs['_outlineIndent'])+'pt;margin-top:0;margin-bottom:0'
    if style:
        extra += ' style="' + html.escape(style, quote=True) + '"'
    if kind == "ordered_list":
        if attrs.get("continuePrevious"): extra += ' data-continue-list="true"'
        extra += f' start="{int(_number(attrs.get("order", attrs.get("start")), 1, 1, 1000000000))}"'
    if kind in ("table_cell", "table_header"):
        if attrs.get('colwidth'):
            extra += ' data-colwidth="' + ','.join(str(round(value)) for value in attrs['colwidth']) + '"'
        for field in ("colspan", "rowspan"):
            if attrs.get(field, 1) > 1:
                extra += f' {field}="{int(_number(attrs[field], 1, 1, 100))}"'
    if kind == "code_block":
        inner = "<code>" + inner + "</code>"
    if kind == "table":
        from .docx_support import table_grid
        columns, placed = table_grid(node.get('content', []))
        widths = {}
        for _, col, _, _, cell in placed:
            for offset, width in enumerate(cell.get('attrs', {}).get('colwidth') or []):
                if width: widths[col + offset] = width
        colgroup = '<colgroup>' + ''.join(f'<col style="width:{widths[col]:g}px"/>' if col in widths else '<col/>' for col in range(columns)) + '</colgroup>' if widths else ''
        if len(widths) == columns:
            extra = extra.replace('style="', f'style="width:{sum(widths.values()):g}px;')
        inner = colgroup + "<tbody>" + inner + "</tbody>"
    return f"<{tag}{extra}>{inner}</{tag}>"


def _css(settings):
    geometry = page_geometry(settings)
    margins = geometry['margins']
    return f"""@page {{ size: {geometry['width']:g}mm {geometry['height']:g}mm; margin: {margins['top']:g}mm {margins['right']:g}mm {margins['bottom']:g}mm {margins['left']:g}mm; }}
* {{ box-sizing: border-box; }}
body {{ color: #202020; background: white; font-family: {settings['fontFamily']}, serif;
font-size: {settings['fontSize']:g}pt; line-height: {settings['lineHeight']:g}; margin: 0; }}
main {{ max-width: 52em; padding: 2em; margin: 0 auto; }}
h1,h2,h3,h4,h5,h6 {{ line-height: 1.2; break-after: avoid; color: black; }}
p {{ margin: 0 0 0.7em; orphans: 2; widows: 2; }}
p span,h1 span,h2 span,h3 span,h4 span,h5 span,h6 span {{ line-height: var(--alder-paragraph-leading,inherit); }}
[data-revision-insertion] {{ text-decoration: underline; }} [data-revision-deletion] {{ text-decoration: line-through; }}
section {{ margin: 1.5em 0; }} img {{ max-width: 100%; height: auto; }}
table {{ border-collapse: collapse; width: 100%; margin: 1em 0; table-layout: fixed; }}
td,th {{ border: .75pt solid #999999; padding: 5.25pt; overflow-wrap: anywhere; vertical-align: top; }}
th {{ background: #eeeeee; }} td p,th p {{ margin: 0; }}
blockquote {{ border-left: 2px solid #aaaaaa; margin: 1em; padding-left: 1em; }}
pre {{ white-space: pre-wrap; overflow-wrap: anywhere; font-size: 0.85em; }}
a {{ color: #275268; }} .author {{ margin-bottom: 2em; }}
.page-break {{ break-before: page; }} .missing-image {{ color: #8b3a22; }}
@media print {{ main {{ max-width: none; margin: 0; padding: 0; }} nav {{ break-after: page; }} }}"""


def render_html(project: dict, options: dict | None = None) -> str:
    warnings = []
    return _render_html(_prepare_publication(project, options, warnings), options, warnings)


def _render_html(project, options, warnings):
    settings = _settings(project, options)
    resources = _Resources(project, settings, warnings)
    sections = _publication_sections(project, warnings)
    title = html.escape(project.get("name", "Untitled"))
    pieces = [f"<h1>{title}</h1>" if settings["includeTitle"] else ""]
    if settings.get("author") and settings["includeTitle"]:
        pieces.append('<p class="author">' + html.escape(str(settings["author"])) + "</p>")
    if settings["includeToc"]:
        pieces.append('<nav aria-label="Contents"><h2>Contents</h2><ol>' + "".join(f'<li><a href="#section-{i}">{html.escape(s["title"] or "Untitled section")}</a></li>' for i, s in enumerate(sections)) + "</ol></nav>")
    page_styles = []
    for i, section in enumerate(sections):
        heading = f'<h2>{html.escape(section["title"])}</h2>' if section["title"] else ""
        chunks, pending, page_style = [], [], None
        for b in section['blocks']:
            if _kind(b) == 'section_break':
                chunks.append(('<div style="' + page_style + '">' if page_style else '<div>') + ''.join(pending) + '</div>')
                pending = [_html_node(b, project, resources, warnings)]
                attrs = b.get('attrs', {})
                geo = page_geometry(attrs.get('page', {}))
                edges = geo['margins']
                name = 'alder-section-' + str(len(page_styles))
                page_styles.append(f"@page {name} {{ size:{geo['width']:g}mm {geo['height']:g}mm; margin:{edges['top']:g}mm {edges['right']:g}mm {edges['bottom']:g}mm {edges['left']:g}mm }}")
                boundary = {'oddPage':'right', 'evenPage':'left'}.get(attrs.get('kind'), 'page')
                page_style = f'page:{name};break-before:{boundary}'
                if attrs.get('kind')=='continuous':
                    page_style=''
                    _warn(warnings,'HTML retains continuous-section settings as metadata; use desktop PDF for their physical margin and running-story transitions.')
            else:
                pending.append(_html_node(b, project, resources, warnings))
        chunks.append(('<div style="' + page_style + '">' if page_style else '<div>') + ''.join(pending) + '</div>')
        content = ''.join(chunks)
        pieces.append(f'<section id="section-{i}" data-role="{html.escape(section["role"], quote=True)}">{heading}{content}</section>')
    # Used by the preview endpoint: a restrictive document policy survives use in a new tab.
    csp = "default-src 'none'; img-src data:; style-src 'unsafe-inline'; font-src 'none'; base-uri 'none'; form-action 'none'"
    language = html.escape(project.get("language", "en"), quote=True)
    metadata = ''.join(f'<meta name="{key}" content="{html.escape(str(settings.get(key,"")),quote=True)}"/>' for key in ('author','subject','description'))
    from .running_matter import html_metadata
    metadata += html_metadata(settings)
    return f'<!DOCTYPE html><html lang="{language}"><head><meta charset="utf-8"/><meta http-equiv="Content-Security-Policy" content="{csp}"/><meta name="viewport" content="width=device-width,initial-scale=1"/><title>{title}</title>{metadata}<style>{_css(settings)}{"".join(page_styles)}</style></head><body><main>{"".join(pieces)}</main></body></html>'


def _markdown(node, warnings):
    kind, attrs = _kind(node), node.get("attrs", {})
    if kind in ('document_field','sequence'):
        _warn(warnings, 'Markdown keeps the current document fields and caption numbers as ordinary text.')
        return ''.join(_markdown({**child,'marks':node.get('marks',[])+child.get('marks',[])},warnings) for child in node.get('content',[]))
    if kind == 'cross_reference':
        _warn(warnings, 'Markdown keeps the current cross-reference wording as ordinary text.')
        return ''.join(_markdown({**child, 'marks':node.get('marks', []) + child.get('marks', [])}, warnings) for child in node.get('content', []))
    if kind == "bookmark_end":
        return ''
    if kind == "bookmark":
        _warn(warnings, "Markdown omits document bookmarks while retaining all wording.")
        return ""
    if kind == "text":
        text = re.sub(r"([\\`*_{}\[\]<>])", r"\\\1", node.get("text", ""))
        for mark in node.get("marks", []):
            name = mark.get("type")
            wrapper = {"strong": "**", "bold": "**", "em": "*", "italic": "*", "code": "`", "strike": "~~"}.get(name)
            if wrapper:
                text = wrapper + text + wrapper
            elif name == "link" and _safe_link(mark.get("attrs", {}).get("href")):
                text = f'[{text}]({mark["attrs"]["href"].replace(")", "%29")})'
            else:
                _warn(warnings, f"Markdown does not preserve the '{name}' text mark.")
        return text
    if kind == "hard_break":
        return "  \n"
    if kind in ("image", "inline_image"):
        _warn(warnings, "Markdown images are exported as alt text; use HTML, DOCX, PDF or EPUB to embed images.")
        return "" if attrs.get("decorative") else attrs.get("alt") or "[Image]"
    if kind == "table":
        _warn(warnings, "Markdown tables are flattened to tab-separated rows.")
        return "\n" + plain_text(node) + "\n\n"
    if kind in ("page_break", "section_break"):
        _warn(warnings, "Markdown does not preserve print page or section breaks.")
        return "\n\n"
    if kind == "horizontal_rule":
        return "\n---\n\n"
    inner = "".join(_markdown(c, warnings) for c in node.get("content", []))
    if kind == "heading":
        if attrs.get('_headingLabel'):
            _warn(warnings, 'Markdown converts automatic heading numbers to ordinary text.')
            inner = attrs['_headingLabel'] + ' ' + inner
        return "#" * int(_number(attrs.get("level"), 2, 1, 6)) + " " + inner + "\n\n"
    if kind == "paragraph":
        return inner + "\n\n"
    if kind == "blockquote":
        return "\n".join("> " + line for line in inner.rstrip().split("\n")) + "\n\n"
    if kind in ("ordered_list", "bullet_list"):
        if attrs.get('_outlineIndent') is not None:
            _warn(warnings,'Markdown converts outline numbering to literal labels; use Alder or DOCX to retain editable numbering.')
        if attrs.get('listStyle') not in (None, 'decimal', 'disc'):
            _warn(warnings, 'Markdown converted custom list markers to ordinary numbered or bullet lists.')
        start = int(_number(attrs.get("order", attrs.get("start")), 1, 1, 1000000000))
        items = []
        for i, item in enumerate(node.get("content", [])):
            label=(item.get('content') or [{}])[0].get('attrs',{}).get('_outlineLabel')
            prefix = label+' ' if label is not None else f"{start+i}. " if kind == "ordered_list" else "- "
            lines = _markdown(item, warnings).strip().splitlines()
            items.append(prefix + (lines[0] if lines else "") + "".join("\n" + " " * len(prefix) + line for line in lines[1:]))
        return "\n".join(items) + "\n\n"
    if kind == "code_block":
        raw = plain_text(node)
        fence = "`" * max(3, 1 + max([len(m.group()) for m in re.finditer(r"`+", raw)] or [0]))
        return f"{fence}\n{raw}\n{fence}\n\n"
    return inner


def _find_tool(name, env, candidates=()):
    configured = os.environ.get(env)
    if configured:
        return str(Path(configured)) if Path(configured).is_file() else None
    found = shutil.which(name)
    return found or next((str(p) for p in map(Path, candidates) if p.is_file()), None)


def _bundled_root():
    configured = os.environ.get("ALDER_RESOURCES_DIR")
    return Path(configured) / "tools" if configured else Path(__file__).resolve().parents[2] / "work/bundle-resources/tools"


def _java():
    from .platform_runtime import resource_executable
    bundle = _bundled_root()
    candidates = [resource_executable(bundle.parent, "java"), bundle / "java/bin/java.exe", bundle / "java/bin/java", bundle / "jre/bin/java.exe", bundle / "jre/bin/java"]
    candidates.extend(bundle.glob("java/*/bin/java.exe"))
    candidates.extend(bundle.glob("jre/*/bin/java.exe"))
    bundled = next((str(p) for p in candidates if p.is_file()), None)
    if bundled and not os.environ.get("ALDER_JAVA"):
        return bundled
    if os.environ.get("JAVA_HOME"):
        candidates.append(Path(os.environ["JAVA_HOME"]) / "bin/java.exe")
    jetbrains = Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "JetBrains"
    if jetbrains.exists():
        candidates.extend(jetbrains.glob("*/jbr/bin/java.exe"))
    return _find_tool("java", "ALDER_JAVA", candidates)


def _epubcheck_command():
    bundle = _bundled_root()
    root = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / ".local/share"))) / "Alder/tools"
    configured = os.environ.get("ALDER_EPUBCHECK_JAR")
    candidates = [Path(configured)] if configured else sorted(bundle.glob("epubcheck*/epubcheck.jar")) + sorted(bundle.glob("epubcheck*/epubcheck-*/epubcheck.jar")) + sorted(root.glob("epubcheck*/epubcheck.jar")) + sorted(root.glob("epubcheck*/epubcheck-*/epubcheck.jar"))
    jar = next((p for p in candidates if p.is_file()), None)
    java = _java()
    if java and jar:
        return [java, "-jar", str(jar)]
    native = shutil.which("epubcheck")
    return [native] if native else None


def capabilities() -> dict:
    from .platform_runtime import resource_executable
    bundle = _bundled_root()
    bundled = next((str(p) for p in [resource_executable(bundle.parent, "calibre"), bundle / "calibre/ebook-convert.exe", bundle / "calibre/Calibre/ebook-convert.exe", bundle / "calibre/Calibre Portable/Calibre/ebook-convert.exe", bundle / "calibre/ebook-convert"] if p.is_file()), None)
    calibre = bundled if bundled and not os.environ.get("ALDER_EBOOK_CONVERT") else _find_tool("ebook-convert", "ALDER_EBOOK_CONVERT", [Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "Calibre2/ebook-convert.exe"])
    return {"formats": list(SUPPORTED_FORMATS[:-1]) + (["azw3"] if calibre else []), "imports": ["txt", "md", "html", "docx", "epub", "pdf", "rtf", "doc", "odt", "ods", "odp", "ppt", "pptx", "xls", "xlsx", "eml", "mobi", "azw3", "fb2", "text"],
            "epubcheck": {"available": bool(_epubcheck_command())}, "azw3": {"available": bool(calibre), "converter": calibre}}


def validate_epub(path: Path, run_epubcheck=True) -> dict:
    """Always check container, manifest, spine and XHTML; label EPUBCheck separately."""
    errors = []
    with _zip(path) as z:
        names = set(z.namelist())
        if z.infolist()[0].filename != "mimetype" or z.infolist()[0].compress_type != zipfile.ZIP_STORED:
            errors.append("The mimetype entry must be first and uncompressed.")
        if "mimetype" not in names or z.read("mimetype") != b"application/epub+zip":
            errors.append("Missing or invalid EPUB mimetype.")
        try:
            container = _xml(z.read("META-INF/container.xml"))
            rootfile = container.find(".//{*}rootfile")
            opf_path = rootfile.attrib["full-path"]
            package = _xml(z.read(opf_path))
            manifest = {i.attrib["id"]: i for i in package.findall("{*}manifest/{*}item")}
            if not manifest:
                errors.append("EPUB manifest is empty.")
            for item in manifest.values():
                target = str(PurePosixPath(opf_path).parent / unquote(item.attrib["href"]))
                if target not in names:
                    errors.append(f"Missing manifest resource: {target}")
                elif item.attrib.get("media-type") == "application/xhtml+xml":
                    _xml(z.read(target))
            spine = package.findall("{*}spine/{*}itemref")
            if not spine:
                errors.append("EPUB spine is empty.")
            for item in spine:
                if item.attrib.get("idref") not in manifest:
                    errors.append("A spine entry references a missing manifest item.")
            if not any("nav" in m.attrib.get("properties", "").split() for m in manifest.values()):
                errors.append("EPUB navigation document is missing.")
        except (KeyError, ET.ParseError, AttributeError, ValueError) as exc:
            errors.append("EPUB package is invalid: " + str(exc))
    result = {"structural": {"status": "failed" if errors else "passed", "errors": errors}, "epubcheck": {"status": "unavailable"}}
    command = _epubcheck_command() if run_epubcheck else None
    if command:
        report_path = path.with_suffix(".epubcheck.json")
        try:
            completed = subprocess.run(command + [str(path), "--json", str(report_path)], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120,
                                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            report = json.loads(report_path.read_text("utf-8")) if report_path.exists() else {}
            result["epubcheck"] = {"status": "passed" if completed.returncode == 0 else "failed", "exitCode": completed.returncode,
                                    "messages": report.get("messages", []), "output": (completed.stdout + completed.stderr)[-16000:]}
        except (OSError, subprocess.TimeoutExpired, ValueError) as exc:
            result["epubcheck"] = {"status": "error", "message": str(exc)}
        finally:
            report_path.unlink(missing_ok=True)
    result["status"] = "failed" if errors or result["epubcheck"]["status"] in ("failed", "error") else "passed" if result["epubcheck"]["status"] == "passed" else "structural-only"
    return result


def _epub(project, path, settings, warnings):
    resources = _Resources(project, settings, warnings)
    sections = _publication_sections(project, warnings)
    from .document_bookmarks import publication_bookmarks
    resources.bookmark_files = {name: f'section-{index+1}.xhtml' for name,index in publication_bookmarks(sections,warnings).items()}
    if not sections:
        sections = [{"title": "", "role": "chapter", "blocks": []}]
    language = html.escape(project.get("language", "en"), quote=True)
    title = html.escape(project.get("name", "Untitled"))
    documents = []
    manifest = []
    spine = []
    nav = []
    cover = None
    if settings.get("coverAssetId"):
        cover = resources.image({"assetId": settings["coverAssetId"], "alt": "Cover"})
    for i, section in enumerate(sections):
        name = f"section-{i+1}.xhtml"
        section_title = html.escape(section["title"] or project.get("name", "Untitled"))
        blocks = "".join(_html_node(b, project, resources, warnings, epub=True) for b in section["blocks"])
        role = section.get("role", "chapter")
        if role not in ("chapter", "part", "preface", "foreword", "introduction", "appendix", "acknowledgments", "epilogue", "afterword", "bibliography", "glossary", "index"):
            role = "chapter"
        heading = f"<h1>{section_title}</h1>" if section["title"] else ""
        if i == 0 and settings.get("includeTitle"):
            heading = f'<header><h1>{title}</h1><p>{html.escape(str(settings.get("author", "")))}</p></header>' + heading
        body = f'<section epub:type="{role}">{heading}{blocks}</section>'
        from .running_matter import html_metadata
        documents.append((name, f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="{language}" xml:lang="{language}"><head><title>{section_title}</title><link rel="stylesheet" href="style.css"/>{html_metadata(settings) if i==0 else ""}</head><body>{body}</body></html>'))
        manifest.append(f'<item id="s{i}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="s{i}"/>')
        nav.append(f'<li><a href="{name}">{section_title}</a></li>')
    if cover:
        documents.insert(0, ("cover.xhtml", f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="{language}"><head><title>Cover</title><link rel="stylesheet" href="style.css"/></head><body epub:type="cover"><img src="images/{cover["name"]}" alt="{title}"/></body></html>'))
        manifest.append('<item id="cover-page" href="cover.xhtml" media-type="application/xhtml+xml"/>')
        spine.insert(0, '<itemref idref="cover-page"/>')
    for i, image in enumerate(resources.images.values()):
        prop = ' properties="cover-image"' if cover and image["name"] == cover["name"] else ""
        manifest.append(f'<item id="image{i}" href="images/{image["name"]}" media-type="image/png"{prop}/>')
    nav_doc = f'<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="{language}" xml:lang="{language}"><head><title>Contents</title></head><body><nav epub:type="toc" id="toc"><h1>Contents</h1><ol>{"".join(nav)}</ol></nav></body></html>'
    identifier = html.escape(str(settings.get("identifier") or "urn:uuid:" + str(uuid.uuid5(uuid.NAMESPACE_URL, str(project.get("id", title))))))
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    metadata = f'<dc:identifier id="book-id">{identifier}</dc:identifier><dc:title>{title}</dc:title><dc:language>{language}</dc:language><meta property="dcterms:modified">{modified}</meta>'
    for field, tag in (("author", "creator"), ("description", "description"), ("publisher", "publisher"), ("rights", "rights"), ("subject", "subject")):
        if settings.get(field):
            metadata += f'<dc:{tag}>{html.escape(str(settings[field]))}</dc:{tag}>'
    opf = f'<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">{metadata}</metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="css" href="style.css" media-type="text/css"/>{"".join(manifest)}</manifest><spine>{"".join(spine)}</spine></package>'
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="EPUB/package.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        archive.writestr("EPUB/package.opf", opf)
        archive.writestr("EPUB/nav.xhtml", nav_doc)
        archive.writestr("EPUB/style.css", _css(settings))
        for name, content in documents:
            archive.writestr("EPUB/" + name, content)
        for image in resources.images.values():
            archive.writestr("EPUB/images/" + image["name"], image["data"])


def _docx(project, path, settings, warnings):
    from .docx_support import set_font_family, apply_paragraph_flow
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    document = Document()
    section = document.sections[0]
    from .sections import write_docx_geometry
    write_docx_geometry(section, settings)
    geometry = page_geometry(settings)
    numbering = OxmlElement("w:pgNumType")
    numbering.set(qn("w:start"), str(settings["firstPageNumber"]))
    numbering_format = settings['pageNumberFormat']
    numbering.set(qn('w:fmt'), numbering_format)
    section._sectPr.append(numbering)
    normal = document.styles["Normal"]
    set_font_family(normal.font, settings["fontFamily"])
    normal.font.size = Pt(settings["fontSize"])
    normal.paragraph_format.line_spacing = settings["lineHeight"]
    normal.paragraph_format.space_after = Pt(8)
    caption_style = document.styles['Caption']
    caption_style.base_style = normal
    set_font_family(caption_style.font, settings['fontFamily'])
    caption_style.font.size = Pt(settings['fontSize'])
    caption_style.font.bold = False; caption_style.font.italic = False
    caption_style.font.color.rgb = RGBColor(0,0,0)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        set_font_family(document.styles[name].font, settings["fontFamily"])
    for level, size in enumerate(DOCUMENT_FORMAT['headingSizes'], 1):
        heading = document.styles[f'Heading {level}']
        heading.font.size = Pt(size)
        heading.font.bold = True
        heading.paragraph_format.line_spacing = DOCUMENT_FORMAT['heading']['lineHeight']
        heading.paragraph_format.space_before = Pt(DOCUMENT_FORMAT['heading']['spaceBefore'])
        heading.paragraph_format.space_after = Pt(DOCUMENT_FORMAT['heading']['spaceAfter'])
    style_ids = {}
    character_style_ids = {}
    for style in project.get("styles", []):
        character = style.get("kind") == "character"
        name = "Alder " + ("Character " if character else "") + str(style.get("name") or style.get("id"))[:80]
        if name in document.styles:
            name += " " + str(style.get("id"))[:32]
        exported = document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER if character else WD_STYLE_TYPE.PARAGRAPH)
        exported.base_style = document.styles["Default Paragraph Font"] if character else normal
        if not character or style.get("fontFamily"):
            set_font_family(exported.font, str(style.get("fontFamily") or settings["fontFamily"]))
        if not character or style.get("fontSize"):
            exported.font.size = Pt(_number(style.get("fontSize"), settings["fontSize"], 6, 72))
        if re.fullmatch(r"#[\da-fA-F]{6}", style.get("color") or ""):
            exported.font.color.rgb = RGBColor.from_string(style["color"][1:])
        if not character:
            exported.paragraph_format.line_spacing = _number(style.get("lineHeight"), settings["lineHeight"], 0.8, 4)
            exported.paragraph_format.space_after = Pt(_number(style.get("spaceAfter"), 8, 0, 144))
            for key, attribute in (("spaceBefore", "space_before"), ("leftIndent", "left_indent"), ("rightIndent", "right_indent"), ("firstLineIndent", "first_line_indent")):
                if style.get(key) is not None:
                    setattr(exported.paragraph_format, attribute, Pt(style[key]))
            apply_paragraph_flow(exported.paragraph_format, style, warnings)
        (character_style_ids if character else style_ids)[style.get("id")] = name
    for style in project.get('styles', []):
        if style.get('id') in style_ids and style.get('nextStyle') in style_ids:
            document.styles[style_ids[style['id']]].next_paragraph_style = document.styles[style_ids[style['nextStyle']]]
    document.core_properties.title = project.get("name", "Untitled")
    document.core_properties.author = str(settings.get("author", ""))
    document.core_properties.subject = str(settings.get("subject", ""))
    document.core_properties.comments = str(settings.get("description", ""))
    resources = _Resources(project, settings, warnings)

    bookmark_names = set()
    bookmark_counter = 1000000
    bookmark_ranges = {}
    from .document_comments import WordCommentsWriter
    from .document_revisions import WordRevisionWriter, word_review_paragraphs
    revision_writer = WordRevisionWriter()
    sections = _publication_sections(project, warnings)
    sections = [{**section, 'blocks': word_review_paragraphs({'type': 'doc', 'content': section['blocks']})['content']} for section in sections]
    from .running_headings import export_heading_styles
    heading_styles = export_heading_styles(sections,settings,style_ids)
    comment_writer = WordCommentsWriter(document, sections, warnings)
    def append_field(paragraph, instruction, node):
        # CHARFORMAT refers to the first instruction character, not the cached
        # result. Paired fields give that character the authored run properties.
        start = len(paragraph._p)
        for child in node.get('content', []): inline(paragraph,{**child,'marks':node.get('marks',[])+child.get('marks',[])})
        runs = list(paragraph._p)[start:]
        properties = next((run.find(qn('w:rPr')) for run in runs if run.find(qn('w:rPr')) is not None),None)
        for run in runs: paragraph._p.remove(run)
        for kind in ('begin','instruction','separate'):
            run = OxmlElement('w:r')
            if properties is not None: run.append(copy.deepcopy(properties))
            marker = OxmlElement('w:instrText' if kind == 'instruction' else 'w:fldChar')
            if kind == 'instruction':
                marker.text = instruction; marker.set(qn('xml:space'),'preserve')
            else: marker.set(qn('w:fldCharType'),kind)
            run.append(marker);paragraph._p.append(run)
        for run in runs: paragraph._p.append(run)
        end = OxmlElement('w:r'); marker = OxmlElement('w:fldChar'); marker.set(qn('w:fldCharType'),'end');end.append(marker);paragraph._p.append(end)
    def inline(paragraph, node):
        nonlocal bookmark_counter
        if revision_writer.inline(paragraph, node, inline): return
        kind, attrs = _kind(node), node.get("attrs", {})
        if kind == 'comment_anchor':
            comment_writer.inline(paragraph, attrs)
        elif kind == 'document_field':
            from .document_fields import field_instruction
            append_field(paragraph,field_instruction(attrs['fieldKind'],heading_styles),node)
        elif kind == 'sequence':
            field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), ' SEQ ' + attrs['sequenceName'] + ' \\* ARABIC ')
            start = len(paragraph._p)
            for child in node.get('content', []): inline(paragraph, {**child,'marks':node.get('marks',[])+child.get('marks',[])})
            for run in list(paragraph._p)[start:]: field.append(run)
            paragraph._p.append(field)
        elif kind == 'cross_reference':
            name = attrs.get('targetName')
            if name and not attrs.get('unresolved'):
                append_field(paragraph,(' PAGEREF ' if attrs.get('referenceFormat') == 'page' else ' REF ') + name + (' \\w' if attrs.get('referenceFormat') == 'number' else '') + ' \\h \\* CHARFORMAT ',node)
            else:
                for child in node.get('content', []): inline(paragraph, child)
        elif kind == "bookmark":
            from .document_bookmarks import bookmark_name
            name = bookmark_name(attrs["name"], bookmark_names, warnings)
            start=OxmlElement("w:bookmarkStart");start.set(qn("w:id"),str(bookmark_counter));start.set(qn("w:name"),name)
            paragraph._p.append(start)
            if attrs.get('range'):
                bookmark_ranges[name] = str(bookmark_counter)
            else:
                end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(bookmark_counter));paragraph._p.append(end)
            bookmark_counter+=1
        elif kind == 'bookmark_end':
            target = bookmark_ranges.get(attrs.get('targetName'))
            if target and not attrs.get('unresolved'):
                end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),target);paragraph._p.append(end)
        elif kind == "hard_break":
            paragraph.add_run().add_break()
        elif kind in ("image", "inline_image"):
            image = resources.image(attrs, apply_crop=False)
            if image:
                max_mm = geometry['width'] - geometry['margins']['left'] - geometry['margins']['right']
                width_mm = min(max_mm, _number(attrs.get("width"), image["width"], 1, 10000) * 25.4 / 96)
                from .image_crop import validate_crop
                crop = validate_crop(attrs.get('crop'))
                ratio = (100-crop['top']-crop['bottom'])/(100-crop['left']-crop['right']) if crop else 1
                if not attrs.get('width') and crop:
                    width_mm *= (100-crop['left']-crop['right'])/100
                shape = paragraph.add_run().add_picture(io.BytesIO(image["data"]), width=round(width_mm*36000), height=round(width_mm*image['height']/image['width']*ratio*36000))
                if crop:
                    fill = shape._inline.graphic.graphicData.pic.blipFill
                    source_crop = OxmlElement('a:srcRect')
                    for edge, short in (('left','l'),('top','t'),('right','r'),('bottom','b')):
                        source_crop.set(short, str(round(crop[edge]*1000)))
                    fill.insert(1, source_crop)
                shape._inline.docPr.set("descr", attrs.get("alt") or "")
                # Word preserves drawing names on save. Only our explicitly named
                # standalone drawings are lifted out of their paragraph on import.
                if kind == 'image':
                    shape._inline.docPr.set("name", f"Alder block image {shape._inline.docPr.id}")
                from .image_accessibility import write_decorative
                write_decorative(shape._inline.docPr, attrs.get("decorative"))
                if attrs.get('title'):
                    shape._inline.docPr.set('title', attrs['title'])
            elif attrs.get("alt") and not attrs.get("decorative"):
                paragraph.add_run(attrs["alt"])
        elif kind == "text":
            run = paragraph.add_run(node.get("text", ""))
            for mark in node.get("marks", []):
                name, a = mark.get("type"), mark.get("attrs", {})
                if name in ("strong", "bold"):
                    run.bold = True
                elif name in ("em", "italic"):
                    run.italic = True
                elif name == "underline":
                    run.underline = True
                elif name in ("strike", "s"):
                    run.font.strike = True
                elif name in ("subscript", "superscript"):
                    setattr(run.font, name, True)
                elif name == "highlight":
                    from docx.enum.text import WD_COLOR_INDEX
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                elif name == "code":
                    set_font_family(run.font, "Consolas")
                elif name in ("textStyle", "text_style", "fontFamily", "fontSize", "color"):
                    if a.get("styleId") in character_style_ids:
                        run.style = character_style_ids[a["styleId"]]
                    if a.get("fontSize"):
                        run.font.size = Pt(_number(a["fontSize"], 12, 6, 72))
                    if a.get("fontFamily"):
                        set_font_family(run.font, str(a["fontFamily"])[:100])
                    if re.fullmatch(r"#[\da-fA-F]{6}", a.get("color") or ""):
                        run.font.color.rgb = RGBColor.from_string(a["color"][1:])
                elif name == "link":
                    href = _safe_link(a.get("href"))
                    if href:
                        link = OxmlElement("w:hyperlink")
                        if href.startswith("#"):
                            link.set(qn("w:anchor"), href[1:])
                        else:
                            link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
                        link.append(run._r)
                        paragraph._p.append(link)
                else:
                    _warn(warnings, f"DOCX omitted unsupported mark '{name}'.")
        else:
            for child in node.get("content", []):
                inline(paragraph, child)

    from .docx_support import NumberingWriter
    list_numbering = NumberingWriter(document)
    from .heading_numbering import HeadingNumberingWriter
    heading_numbering = HeadingNumberingWriter(document)
    from .outline_lists import word_outline_bindings
    outline_bindings, outline_lists = word_outline_bindings(document,sections)

    def block(container, node, indent=0, prefix="", list_info=None):
        kind, attrs = _kind(node), node.get("attrs", {})
        if kind == "table":
            from .docx_support import table_grid
            rows = node.get("content", [])
            columns, placed = table_grid(rows)
            table = container.add_table(rows=max(1, len(rows)), cols=columns)
            table.style = "Table Grid"
            from .table_formatting import write_docx_cell
            from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
            table.alignment = {'left': WD_TABLE_ALIGNMENT.LEFT, 'center': WD_TABLE_ALIGNMENT.CENTER, 'right': WD_TABLE_ALIGNMENT.RIGHT}.get(attrs.get('align'), WD_TABLE_ALIGNMENT.LEFT)
            widths = {}
            for r, col, height, span, cell in placed:
                for offset, width in enumerate(cell.get("attrs", {}).get("colwidth") or []):
                    if width: widths[col + offset] = width
            if widths:
                table.autofit = False
                for col, width in widths.items():
                    table.columns[col].width = Pt(width * 72 / 96)
                    for row in table.rows: row.cells[col].width = Pt(width * 72 / 96)
            for r, col, height, span, cell in placed:
                target = table.cell(r, col)
                if span > 1 or height > 1:
                    target = target.merge(table.cell(r + height - 1, col + span - 1))
                for child in cell.get("content", []):
                    block(target, child)
                if len(target.paragraphs) > 1 and not target.paragraphs[0].text:
                    target._element.remove(target.paragraphs[0]._p)
                if _kind(cell) == "table_header":
                    for paragraph in target.paragraphs:
                        for run in paragraph.runs: run.bold = True
                write_docx_cell(target, cell.get('attrs', {}), _kind(cell) == 'table_header')
                if height > 1:
                    # Word reads margins/borders on each vertical-merge
                    # continuation, even though python-docx returns the anchor
                    # for table.cell(). Keep the full merged outline uniform.
                    from docx.table import _Cell
                    for continuation_row in table.rows[r + 1:r + height]:
                        grid_column = 0
                        for element in continuation_row._tr.tc_lst:
                            if grid_column == col and element.vMerge == 'continue':
                                write_docx_cell(_Cell(element, table), cell.get('attrs', {}), _kind(cell) == 'table_header')
                                break
                            grid_column += element.grid_span
            for r, row in enumerate(rows):
                if not row.get('attrs', {}).get('allowSplit'):
                    table.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if row.get('attrs', {}).get('minHeight') is not None:
                    table.rows[r].height = Pt(row['attrs']['minHeight'])
                    table.rows[r].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                if row.get("content") and all(_kind(c) == "table_header" for c in row["content"]):
                    table.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            return
        if kind in ("bullet_list", "ordered_list"):
            if id(node) in outline_lists:
                for item in node.get('content',[]):
                    for child in item.get('content',[]):block(container,child)
                return
            start = int(_number(attrs.get("order", attrs.get("start")), 1, 1, 1000000000))
            number = continuation_numbers.get(id(node))
            if number is None: number = list_numbering.create(kind == "ordered_list", start, attrs.get('listStyle'))
            for item in node.get("content", []):
                for j, child in enumerate(item.get("content", [])):
                    block(container, child, indent+1, list_info=(number, indent) if j == 0 else None)
            return
        if kind == "blockquote":
            for child in node.get("content", []):
                block(container, child, indent+1)
            return
        if kind == "section_break":
            nonlocal geometry, numbering_format
            from docx.enum.section import WD_SECTION
            new_section = document.add_section({'nextPage': WD_SECTION.NEW_PAGE, 'continuous':WD_SECTION.CONTINUOUS, 'oddPage': WD_SECTION.ODD_PAGE, 'evenPage': WD_SECTION.EVEN_PAGE}[attrs.get('kind', 'nextPage')])
            write_docx_geometry(new_section, attrs.get('page', {}))
            geometry = page_geometry(attrs.get('page', {}))
            for numbering in list(new_section._sectPr.findall(qn('w:pgNumType'))):
                new_section._sectPr.remove(numbering)
            numbering_format = attrs.get('numberingFormat') or numbering_format
            numbering = OxmlElement('w:pgNumType')
            numbering.set(qn('w:fmt'), numbering_format)
            if attrs.get('numberingStart') is not None:
                numbering.set(qn('w:start'), str(attrs['numberingStart']))
            new_section._sectPr.append(numbering)
            if attrs.get('runningMatter'):
                from .running_matter import write_running_matter
                write_running_matter(new_section,attrs['runningMatter'],block,document)
            return
        if kind == 'table_of_contents':
            from .word_contents import write_contents
            write_contents(container, node, block, document, style_ids, warnings)
            return
        if kind == "page_break":
            from docx.enum.text import WD_BREAK
            container.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            return
        if kind == "horizontal_rule":
            paragraph = container.add_paragraph()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, val in (("val", "single"), ("sz", "4"), ("color", "AAAAAA")):
                bottom.set(qn("w:"+key), val)
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
            return
        if kind in ("paragraph", "heading", "code_block", "image"):
            outline_list = outline_bindings.get(id(node))
            if outline_list:list_info=outline_list[:2]
            style = style_ids.get(attrs.get("styleId"))
            if not style and attrs.get('captionLabel'): style = 'Caption'
            if not style and kind == "heading":
                style = "Heading " + str(int(_number(attrs.get("level"), 2, 1, 6)))
            paragraph = container.add_paragraph(style=style)
            if kind == 'heading':
                outline = OxmlElement('w:outlineLvl'); outline.set(qn('w:val'), str(int(attrs.get('level', 1)) - 1)); paragraph._p.get_or_add_pPr().append(outline)
                if attrs.get('headingNumbered'):
                    if list_info: raise ValueError('A heading inside a numbered list cannot carry a second heading number. Remove one numbering style before DOCX export.')
                    heading_numbering.apply(paragraph, attrs)
            if prefix:
                paragraph.add_run(prefix)
            if list_info:
                list_numbering.apply(paragraph, *list_info)
            if indent:
                paragraph.paragraph_format.left_indent = Mm(7 * indent)
            if kind == "image":
                inline(paragraph, node)
            else:
                for child in node.get("content", []):
                    inline(paragraph, child)
            if kind == "code_block":
                for run in paragraph.runs:
                    set_font_family(run.font, "Consolas")
            else:
                for run in paragraph.runs:
                    if attrs.get("fontFamily") and run.font.name is None:
                        set_font_family(run.font, str(attrs["fontFamily"])[:100])
                    if attrs.get("fontSize") and run.font.size is None:
                        run.font.size = Pt(_number(attrs["fontSize"], settings["fontSize"], 6, 72))
                    if attrs.get("color") and run.font.color.rgb is None and re.fullmatch(r"#[\da-fA-F]{6}", attrs["color"]):
                        run.font.color.rgb = RGBColor.from_string(attrs["color"][1:])
            from .docx_support import apply_paragraph_properties
            apply_paragraph_properties(paragraph, attrs, settings['lineHeight'], warnings)
            if outline_list:
                paragraph.paragraph_format.left_indent=Pt(outline_list[2]['indent']+(attrs.get('leftIndent') or 0))
                paragraph.paragraph_format.first_line_indent=Pt((attrs.get('firstLineIndent') or 0)-outline_list[2]['hanging'])
            if attrs.get('revisionBreak'): revision_writer.paragraph_break(paragraph, attrs['revisionBreak'])
            if attrs.get('revisionProperties'):
                from .paragraph_revisions import write_paragraph_revision
                write_paragraph_revision(paragraph, attrs['revisionProperties'], revision_writer, settings['lineHeight'], warnings)
            return
        _warn(warnings, f"DOCX flattened unsupported block '{kind}'.")
        for child in node.get("content", []):
            block(container, child, indent)

    if settings.get('runningMatter') is not None or any(n.get('attrs',{}).get('runningMatter') is not None for item in sections for n in item['blocks'] if n.get('type')=='section_break'):
        from .running_matter import write_running_matter, legacy_running_matter, resolve_running_matter
        write_running_matter(section,resolve_running_matter(legacy_running_matter(settings),settings.get('runningMatter')),block,document,first=True)
        rich_running_matter = True
    else:
        rich_running_matter = False
    if settings["includeTitle"]:
        document.add_paragraph(project.get("name", "Untitled"), "Title")
        if settings.get("author"):
            document.add_paragraph(str(settings["author"]), "Subtitle")
    continuation_numbers = {}
    from .list_formatting import list_sequences
    for item in sections:
        for listing, previous in list_sequences({'content':item['blocks']}):
            attrs = listing.get('attrs', {})
            continuation_numbers[id(listing)] = continuation_numbers[id(previous)] if previous is not None else list_numbering.create(True, attrs.get('order', 1), attrs.get('listStyle'))
    if settings["includeToc"]:
        document.add_heading("Contents", 1)
        for i, item in enumerate(sections):
            inline(document.add_paragraph(), {"type": "text", "text": item["title"] or "Untitled section", "marks": [{"type": "link", "attrs": {"href": f"#_AlderSection{i}"}}]})
        document.add_page_break()
    for i, item in enumerate(sections):
        if settings.get("chapterPageBreaks") and i:
            document.add_page_break()
        start_index = len(document._element.body) - 1
        heading = document.add_heading(item["title"], 1) if item["title"] else None
        for node in item["blocks"]:
            block(document, node)
        anchor = heading._p if heading is not None else next(document._element.body[start_index].iter(qn("w:p")), None)
        if anchor is None:
            continue
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), str(i))
        bookmark.set(qn("w:name"), f"_AlderSection{i}")
        anchor.insert(1 if anchor.find(qn("w:pPr")) is not None else 0, bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(i))
        anchor.append(end)
    if settings.get("header") and not rich_running_matter:
        section.header.paragraphs[0].text = str(settings["header"])
    if settings.get("footer") and not rich_running_matter:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    comment_writer.finish()
    document.save(path)


def _pdf(project, path, settings, warnings):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
    from reportlab.lib.pagesizes import A4, A5, LETTER, LEGAL
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Flowable, Paragraph, Spacer, PageBreak, CondPageBreak, KeepTogether, Table, TableStyle, Image, HRFlowable
    from .fonts import installed_faces, register_pdf_font
    faces = installed_faces()
    font_cache = {}
    def installed_font(family):
        if family not in font_cache:
            font_cache[family] = register_pdf_font(family, faces)
        return font_cache[family]
    font = installed_font(settings["fontFamily"]) or "Times-Roman"
    if font == "Times-Roman":
        configured = os.environ.get("ALDER_RESOURCES_DIR")
        bundled_fonts = (Path(configured) if configured else Path(__file__).resolve().parents[2] / "work/bundle-resources") / "fonts"
        if all((bundled_fonts / f"LiberationSerif-{style}.ttf").is_file() for style in ("Regular", "Bold", "Italic", "BoldItalic")):
            font = "AlderLiberationSerif"
            for suffix, style in (("", "Regular"), ("-Bold", "Bold"), ("-Italic", "Italic"), ("-BoldItalic", "BoldItalic")):
                if font+suffix not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(font+suffix, str(bundled_fonts / f"LiberationSerif-{style}.ttf")))
            pdfmetrics.registerFontFamily(font, normal=font, bold=font+"-Bold", italic=font+"-Italic", boldItalic=font+"-BoldItalic")
            _warn(warnings, f"PDF substituted bundled Liberation Serif for the requested font {settings['fontFamily']}.")
        else:
            _warn(warnings, "The bundled publication font is missing; PDF uses built-in Times with limited Unicode coverage.")
    def publication_font(family):
        if family == settings["fontFamily"]:
            return font
        requested = installed_font(family)
        if requested:
            return requested
        _warn(warnings, f"PDF substituted the publication font for {family}, which is unavailable or cannot be embedded.")
        return font
    geometry = page_geometry(settings)
    pagesize = (geometry['width'] * mm, geometry['height'] * mm)
    margins = {edge: value * mm for edge, value in geometry['margins'].items()}
    width = pagesize[0] - margins['left'] - margins['right']
    body_height = pagesize[1] - margins['top'] - margins['bottom']
    document = BaseDocTemplate(str(path), pagesize=pagesize, leftMargin=margins['left'], rightMargin=margins['right'], topMargin=margins['top'], bottomMargin=margins['bottom'],
                                 title=project.get("name", "Untitled"), author=str(settings.get("author", "")), subject=str(settings.get("description", "")))
    sections = _publication_sections(project, warnings)
    from .running_headings import refresh_body_headings, heading_text
    from .document_fields import is_heading_field
    first_headings = refresh_body_headings(sections, warnings)
    latest_headings, page_headings = {}, {}
    class HeadingParagraph(Paragraph):
        def draw(self):
            super().draw()
            if hasattr(self,'running_heading'):
                level,text=self.running_heading
                page_headings.setdefault(self.canv.getPageNumber(),{}).setdefault(level,text)
                latest_headings[level]=text
        def split(self,width,height):
            parts=super().split(width,height)
            if parts and hasattr(self,'running_heading'):parts[0].running_heading=self.running_heading
            return parts
    resources = _Resources(project, settings, warnings)
    base = ParagraphStyle("Body", fontName=font, fontSize=settings["fontSize"], leading=settings["fontSize"] * settings["lineHeight"], spaceAfter=8)
    templates = []
    blank_pages = set()
    number_offset = settings['firstPageNumber'] - 1
    numbering_format = settings['pageNumberFormat']
    from .running_matter import legacy_running_matter, resolve_running_matter
    running=resolve_running_matter(legacy_running_matter(settings),settings.get('runningMatter'))
    advanced_running=settings.get('runningMatter') is not None
    def add_template(page_settings, restart=None, number_format='decimal', kind='nextPage', matter=None):
        geo = page_geometry(page_settings)
        size = (geo['width'] * mm, geo['height'] * mm)
        edges = {edge: value * mm for edge, value in geo['margins'].items()}
        body_width = size[0] - edges['left'] - edges['right']
        body_h = size[1] - edges['top'] - edges['bottom']
        name = 'section-' + str(len(templates))
        from .running_matter import stories
        from .document_comments import walk as running_walk
        deferred_paint = any(n.get('type')=='document_field' and is_heading_field(n.get('attrs',{}).get('fieldKind')) for story in stories(matter) for n in running_walk(story))
        first_page = True
        first_sheet = None
        def begin_page(canvas, doc):
            nonlocal first_page, number_offset, first_sheet
            if doc.page in blank_pages:
                return
            if first_page:
                first_sheet=doc.page
                if restart is not None:
                    number_offset = section_numbering(doc.page + number_offset, kind, restart)[0] - doc.page
                first_page = False
            if not deferred_paint:paint(canvas,doc)
        def paint(canvas, doc):
            if doc.page in blank_pages:return
            is_first=doc.page==first_sheet
            canvas.saveState()
            canvas.setFont(font, 9)
            canvas.setFillColor(colors.HexColor('#555555'))
            if matter is not None:
                variant='first' if matter.get('differentFirst') and is_first else 'even' if matter.get('differentEven') and doc.page%2==0 else 'default'
                from .document_comments import walk as story_walk
                for part in ('header','footer'):
                    source=matter.get(part,{}).get(variant)
                    if source is None:continue
                    source=copy.deepcopy(source)
                    for field in story_walk(source):
                        if field.get('type')=='document_field' and field.get('attrs',{}).get('fieldKind')=='page':
                            field['content']=[{'type':'text','text':format_page_number(doc.page+number_offset,number_format)}]
                        elif field.get('type')=='document_field' and field.get('attrs',{}).get('fieldKind')=='pageCount':
                            _warn(warnings,'Direct backend PDF retains the last calculated total-page field in headers and footers; desktop PDF calculates it from the shared page layout.')
                        elif field.get('type')=='document_field' and is_heading_field(field.get('attrs',{}).get('fieldKind')):
                            level=int(field['attrs']['fieldKind'][-1])
                            value=page_headings.get(doc.page,{}).get(level,latest_headings.get(level,first_headings.get(level,'—')))
                            field['content']=[{'type':'text','text':value}] if value else []
                    paragraphs=[]
                    for node in source.get('content',[]):
                        node['attrs']={'fontSize':9,'lineHeight':1.2,'spaceBefore':0,'spaceAfter':0,**{k:v for k,v in node.get('attrs',{}).items() if v is not None},'keepTogether':False,'pageBreakBefore':False,'keepWithNext':False}
                        for item in block(node,body_width):
                            _,height=item.wrap(body_width,size[1]);paragraphs.append((item,height))
                    total=sum(height+item.getSpaceBefore()+item.getSpaceAfter() for item,height in paragraphs)
                    distance=matter.get(part+'DistanceMm',10)*mm
                    if total+distance+3>edges['top' if part=='header' else 'bottom']:
                        raise ValueError(f'The {part} needs more room. Increase its page margin, reduce its distance from the edge, or shorten the text.')
                    y=size[1]-distance if part=='header' else distance+total
                    for item,height in paragraphs:
                        y-=item.getSpaceBefore()+height;item.drawOn(canvas,edges['left'],y);y-=item.getSpaceAfter()
            elif settings.get('header'):
                header = str(settings['header'])
                while pdfmetrics.stringWidth(header, font, 9) > body_width and len(header) > 1:
                    header = header[:-2] + '\u2026'
                canvas.drawString(edges['left'], size[1] - edges['top']/2, header)
            if matter is None and settings.get('footer'):
                canvas.drawCentredString(size[0]/2, edges['bottom']/2, format_page_number(doc.page + number_offset, number_format))
            canvas.restoreState()
        frame = Frame(edges['left'], edges['bottom'], body_width, body_h, id=name, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        templates.append(PageTemplate(id=name, pagesize=size, frames=[frame], onPage=begin_page, onPageEnd=lambda canvas,doc:paint(canvas,doc) if deferred_paint else None))
        return name, body_width, body_h
    add_template(settings, number_format=numbering_format,matter=running if advanced_running else None)

    class SectionBreak(Flowable):
        def __init__(self, name, kind, restart=None):
            super().__init__()
            self.name, self.kind, self.restart = name, kind, restart
        def wrap(self, available_width, available_height):
            frame = self._doctemplateAttr('frame')
            _, blank = section_numbering(self.canv.getPageNumber() + 1 + number_offset, self.kind, self.restart)
            if blank:
                blank_pages.add(self.canv.getPageNumber() + 1)
                frame.add_generated_content(PageBreak(), PageBreak(nextTemplate=self.name))
            else:
                frame.add_generated_content(PageBreak(nextTemplate=self.name))
            return 0, 0
        def draw(self):
            pass
    story = []

    class Destination(Flowable):
        # A simple document still needs its link destination, but an empty
        # heading must not introduce a phantom line and heading spacing.
        def __init__(self, name):
            super().__init__()
            self.name = name
            self.keepWithNext = True
        def draw(self):
            self.canv.bookmarkPage(self.name)

    bookmark_names = set()
    def inline(node, available):
        kind = _kind(node)
        if kind in ('document_field','sequence'):
            return ''.join(inline({**child,'marks':node.get('marks',[])+child.get('marks',[])},available) for child in node.get('content',[]))
        if kind == 'cross_reference':
            attrs = node['attrs']
            value = ''.join(inline({**child, 'marks':node.get('marks', []) + child.get('marks', [])}, available) for child in node.get('content', []))
            return f'<link href="#{html.escape(attrs["targetName"], quote=True)}">{value}</link>' if attrs.get('targetName') and not attrs.get('unresolved') else value
        if kind == "bookmark":
            from .document_bookmarks import bookmark_name
            name = bookmark_name(node["attrs"]["name"], bookmark_names, warnings)
            return f'<a name="{name}"/>'
        if kind == "hard_break":
            return "<br/>"
        if kind == "inline_image":
            attrs = node.get('attrs', {})
            image = resources.image(attrs)
            if not image:
                return '' if attrs.get('decorative') else html.escape(attrs.get('alt') or '[Image unavailable]')
            desired = _number(attrs.get('width'), image['width'], 1, 10000) * .75
            ratio = min(desired/image['width'], available/image['width'], max(1, body_height-40)/image['height'])
            source = 'data:image/png;base64,' + base64.b64encode(image['data']).decode()
            return f'<img src="{source}" width="{image["width"]*ratio:g}" height="{image["height"]*ratio:g}" valign="bottom"/>'
        if kind == "image":
            _warn(warnings, "A legacy block image inside text was placed after its paragraph.")
            return ''
        if kind != "text":
            return "".join(inline(c, available) for c in node.get("content", []))
        text = html.escape(node.get("text", ""))
        for mark in node.get("marks", []):
            name, attrs = mark.get("type"), mark.get("attrs", {})
            tag = {"strong": "b", "bold": "b", "em": "i", "italic": "i", "underline": "u", "strike": "strike", "s": "strike", "subscript": "sub", "superscript": "super"}.get(name)
            if tag:
                text = f"<{tag}>{text}</{tag}>"
            elif name == "link":
                href = _safe_link(attrs.get("href"))
                if href:
                    text = f'<link href="{html.escape(href, quote=True)}" color="#275268">{text}</link>'
            elif name == "code":
                text = f'<font name="Courier">{text}</font>'
            elif name == "highlight":
                text = '<span backColor="#ffff00">' + text + '</span>'
            elif name in ("textStyle", "text_style", "fontFamily", "fontSize", "color"):
                attributes = []
                if attrs.get("fontSize"):
                    attributes.append(f'size="{_number(attrs["fontSize"], 12, 6, 72):g}"')
                if re.fullmatch(r"#[\da-fA-F]{6}", attrs.get("color") or ""):
                    attributes.append('color="'+attrs["color"]+'"')
                if attrs.get("fontFamily"):
                    attributes.append('name="' + publication_font(attrs["fontFamily"]) + '"')
                if attributes:
                    text = "<font " + " ".join(attributes) + ">" + text + "</font>"
            else:
                _warn(warnings, f"PDF omitted unsupported mark '{name}'.")
        return text

    def block(node, available=None, indent=0, prefix=""):
        available = width if available is None else available
        kind, attrs = _kind(node), node.get("attrs", {})
        output = []
        if kind == "table":
            from .docx_support import table_grid
            rows = node.get("content", [])
            if not rows:
                return []
            columns, placed = table_grid(rows)
            widths = [None] * columns
            for r, col, height, span, cell in placed:
                for offset, value in enumerate((cell.get("attrs", {}).get("colwidth") or [])[:span]):
                    if value: widths[col + offset] = _number(value, available/columns, 1, 10000) * .75
            missing = widths.count(None)
            remaining = max(12 * missing, available - sum(value or 0 for value in widths))
            widths = [value if value is not None else remaining / max(1, missing) for value in widths]
            if sum(widths) > available:
                scale = available / sum(widths)
                widths = [value * scale for value in widths]
                _warn(warnings, "A table wider than the printable area was scaled to fit the page width.")
            data = [[""] * columns for _ in rows]
            commands = []
            for r, col, height, span, cell in placed:
                ca = cell.get('attrs', {})
                padding = {side: ca.get('padding' + side) if ca.get('padding' + side) is not None else 5.25 for side in ('Top', 'Bottom', 'Left', 'Right')}
                data[r][col] = sum((block(child, max(1, sum(widths[col:col+span])-padding['Left']-padding['Right'])) for child in cell.get("content", [])), [])
                start, end = (col, r), (col + span - 1, r + height - 1)
                for side, value in padding.items(): commands.append((side.upper() + 'PADDING', start, end, value))
                commands.append(('VALIGN', start, end, (ca.get('verticalAlign') or 'top').upper()))
                border_width = ca.get('borderWidth') if ca.get('borderWidth') is not None else .75
                if border_width: commands.append(('BOX', start, end, border_width, colors.HexColor(ca.get('borderColor') or '#999999')))
                if span > 1 or height > 1:
                    commands.append(("SPAN", (col, r), (col + span - 1, r + height - 1)))
                background = ca.get('background') or ('#eeeeee' if _kind(cell) == 'table_header' else None)
                if background: commands.append(('BACKGROUND', start, end, colors.HexColor(background)))
            repeat = 0
            for row in rows:
                if not row.get("content") or not all(_kind(cell) == "table_header" for cell in row["content"]): break
                repeat += 1
            from .pdf_tables import RowBreakTable
            table = RowBreakTable(data, colWidths=widths, hAlign=(attrs.get('align') or 'left').upper(), splitByRow=0, splitInRow=1, repeatRows=repeat, minRowHeights=[row.get('attrs', {}).get('minHeight') or 0 for row in rows])
            table.setStyle(TableStyle(commands))
            for index, row in enumerate(rows):
                for style in table._cellStyles[index]:
                    style.alderAllowSplit = row.get('attrs', {}).get('allowSplit', False)
                    style.alderPageBodyHeight = body_height
            return [table, Spacer(1, 10)]
        if kind == "image":
            image = resources.image(attrs)
            if image:
                desired = _number(attrs.get("width"), image["width"], 1, 10000) * .75
                scale = min(available / image["width"], desired/image["width"], max(1, body_height-40)/image["height"])
                alignment = attrs.get('align', 'left')
                output.append(Image(io.BytesIO(image["data"]), width=image["width"]*scale, height=image["height"]*scale, hAlign=alignment.upper() if alignment in ('left', 'center', 'right') else 'LEFT'))
                output.append(Spacer(1, 8))
            elif attrs.get("alt") and not attrs.get("decorative"):
                output.append(Paragraph(html.escape(attrs["alt"]), base))
            return output
        if kind == "page_break":
            return [PageBreak()]
        if kind == "horizontal_rule":
            return [HRFlowable(width="100%", color=colors.HexColor("#aaaaaa")), Spacer(1, 8)]
        if kind in ("bullet_list", "ordered_list"):
            from .list_formatting import list_marker
            start = int(_number(attrs.get("order", attrs.get("start")), 1, 1, 1000000000))
            for i, item in enumerate(node.get("content", [])):
                for j, child in enumerate(item.get("content", [])):
                    marker = child.get('attrs',{}).get('_outlineLabel')
                    outlined=attrs.get('_outlineIndent') is not None
                    if marker is None:marker=list_marker(start+i, attrs.get('listStyle') or ('decimal' if kind == 'ordered_list' else 'disc'))+('.' if kind=='ordered_list' else '')
                    output.extend(block(child, available, indent+(attrs['_outlineIndent'] if outlined else 12), (marker+' ') if j == 0 else ""))
            return output
        if kind == "blockquote":
            return sum((block(child, available, indent+16) for child in node.get("content", [])), [])
        if kind in ("paragraph", "heading", "code_block"):
            named = next((s for s in project.get("styles", []) if s.get("id") == attrs.get("styleId")), {})
            merged = {**named, **{k:v for k,v in attrs.items() if v is not None}}
            if '\t' in plain_text(node):
                _warn(warnings, "PDF converts paragraph tabs to spaces; use DOCX to preserve tab alignment and leaders.")
            block_font = publication_font(merged.get("fontFamily") or settings["fontFamily"])
            size = _number(merged.get("fontSize"), base.fontSize, 6, 72)
            level = int(_number(attrs.get("level"), 2, 1, 6))
            if kind == "heading" and merged.get('fontSize') is None:
                size = DOCUMENT_FORMAT['headingSizes'][level-1]
            style = ParagraphStyle("Block", parent=base, leftIndent=indent+_number(merged.get("leftIndent"), 0, -144, 144), firstLineIndent=_number(merged.get("firstLineIndent", merged.get("indent")), 0, -144, 144), rightIndent=_number(merged.get("rightIndent"), 0, -144, 144), fontSize=size,
                                   leading=size*_number(merged.get("lineHeight"), 1.25 if kind == "heading" else settings["lineHeight"], 0.8, 4),
                                   spaceBefore=_number(merged.get("spaceBefore"), 12 if kind == "heading" else 0, 0, 144), spaceAfter=_number(merged.get("spaceAfter"), 8, 0, 144),
                                   keepWithNext=bool(merged.get("keepWithNext", kind == "heading")), fontName=block_font+"-Bold" if kind == "heading" and block_font.startswith("Alder") else block_font,
                                   alignment={"left":TA_LEFT,"center":TA_CENTER,"right":TA_RIGHT,"justify":TA_JUSTIFY}.get(merged.get("align") or merged.get("textAlign"), TA_LEFT),
                                   textColor=colors.HexColor(merged["color"]) if re.fullmatch(r"#[\da-fA-F]{6}", str(merged.get("color") or "")) else base.textColor)
            if merged.get('paragraphBackground') and merged['paragraphBackground'] != 'transparent': style.backColor = colors.HexColor(merged['paragraphBackground'])
            style.borderWidth = (merged.get('paragraphBorderWidth') or 0) if merged.get('paragraphBorderColor') != 'transparent' else 0
            style.borderColor = colors.HexColor(merged.get('paragraphBorderColor') or '#000000') if merged.get('paragraphBorderColor') != 'transparent' else None
            style.borderPadding = merged.get('paragraphPadding') or 0
            style.allowWidows = style.allowOrphans = int(merged.get("widowOrphan") is False)
            if attrs.get('_outlineHanging') is not None:style.firstLineIndent-=attrs['_outlineHanging']
            if merged.get("lineSpacingMode") in ("exact", "atLeast") and merged.get("lineSpacing") is not None:
                style.leading = _number(merged['lineSpacing'], size * 1.2, 1, 144)
                style.autoLeading = "max" if merged['lineSpacingMode'] == "atLeast" else ""
            if kind == "code_block":
                text = html.escape(plain_text(node)).replace("\n", "<br/>").replace("  ", "&#160; ")
                style.fontName = "Courier"
                style.fontSize = min(10, size)
            else:
                text = "".join(inline(c, available) for c in node.get("content", []))
            if merged.get('pageBreakBefore'):
                output.append(CondPageBreak(body_height))
            if merged.get('lineSpacingMode') != 'exact' and any(_kind(c) == 'inline_image' for c in node.get('content', [])):
                style.autoLeading = 'max'
            if attrs.get('_headingLabel'): text = html.escape(attrs['_headingLabel']) + '&#160;' + text
            paragraph = (HeadingParagraph if kind=='heading' else Paragraph)(html.escape(prefix)+text or "&#160;", style)
            if kind=='heading':paragraph.running_heading=(level,heading_text(node))
            if merged.get('keepTogether'):
                paragraph = KeepTogether([paragraph])
                paragraph.keepWithNext = style.keepWithNext
            output.append(paragraph)
            for child in node.get("content", []):
                if _kind(child) == "image":
                    output.extend(block(child, available, indent))
            return output
        _warn(warnings, f"PDF flattened unsupported block '{kind}'.")
        return sum((block(child, available, indent) for child in node.get("content", [])), [])

    if settings["includeTitle"]:
        story.extend(block({"type":"paragraph","attrs":{"fontSize":26,"spaceBefore":0,"spaceAfter":15},"content":[{"type":"text","text":project.get("name","Untitled")}]}))
        if settings.get("author"):
            story.append(Paragraph(html.escape(str(settings["author"])), base))
    if settings["includeToc"]:
        story.extend(block({"type":"heading","attrs":{"level":2},"content":[{"type":"text","text":"Contents"}]}))
        for i, item in enumerate(sections):
            story.append(Paragraph(f'<link href="#_AlderSection{i}">{html.escape(item["title"] or "Untitled section")}</link>', base))
        story.append(PageBreak())
    for i, item in enumerate(sections):
        if settings.get("chapterPageBreaks") and i:
            story.append(PageBreak())
        if item['title']:
            chapter_heading=HeadingParagraph(f'<a name="_AlderSection{i}"/>' + html.escape(item["title"]), ParagraphStyle("Section", parent=base, fontSize=18, leading=23, spaceBefore=18, spaceAfter=12, keepWithNext=True))
            chapter_heading.running_heading=(1,item['title']);story.append(chapter_heading)
        else:
            story.append(Destination(f'_AlderSection{i}'))
        for node in item["blocks"]:
            if _kind(node) == 'section_break':
                attrs = node.get('attrs', {})
                numbering_format = attrs.get('numberingFormat') or numbering_format
                kind, restart = attrs.get('kind', 'nextPage'), attrs.get('numberingStart')
                if kind=='continuous':
                    _warn(warnings,'Direct backend PDF converts continuous sections to next-page sections. Desktop PDF uses the shared continuous-section layout.')
                running=resolve_running_matter(running,attrs.get('runningMatter'))
                advanced_running=advanced_running or attrs.get('runningMatter') is not None
                name, width, body_height = add_template(attrs.get('page', {}), restart, numbering_format, kind,running if advanced_running else None)
                story.append(SectionBreak(name, kind, restart))
            else:
                story.extend(block(node))
    if not story:
        story.append(Spacer(1, 1))
    document.addPageTemplates(templates)
    document.build(story)


def build_export(project: dict, format: str, output_dir: Path, options: dict | None = None) -> dict:
    format = {"markdown": "md", "htm": "html"}.get(format.lower(), format.lower())
    if format not in SUPPORTED_FORMATS:
        raise ValueError("Unsupported export format: " + format)
    warnings = []
    options = dict(options or {})
    if format in ('txt', 'md', 'pdf', 'azw3'): options.setdefault('reviewView', 'final')
    project = _prepare_publication(project, options, warnings)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^\w .-]", "", project.get("name", "Untitled"), flags=re.UNICODE).strip(" .")[:80] or "Untitled"
    if stem.split(".", 1)[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1,10)), *(f"LPT{i}" for i in range(1,10))}:
        stem = "Alder-" + stem
    filename = f"{stem}-{uuid.uuid4().hex[:8]}.{format}"
    path = output_dir / filename
    settings = _settings(project, options)
    if page_geometry(settings)['adjusted']:
        _warn(warnings, "The document's margins were reduced proportionally to retain a 20 mm writing area, matching Write.")
    sections = _publication_sections(project, warnings)
    ids = [cid for section in sections for cid in section["clipIds"]]
    from .running_matter import project_running_stories
    if any(plain_text(doc).strip() for doc in project_running_stories(project)):
        if format in ('txt','md'):
            _warn(warnings,'Plain text and Markdown omit running headers and footers; retain Alder or DOCX for that page content.')
        elif format in ('html','epub','azw3'):
            _warn(warnings,'Reflowable output retains header/footer metadata for Alder but does not repeat it on reader pages.')
    if format != 'docx':
        from .document_comments import walk as field_walk
        if any(n.get('type')=='document_field' and n.get('attrs',{}).get('fieldKind') in ('page','pageCount') or n.get('type')=='cross_reference' and n.get('attrs',{}).get('referenceFormat')=='page' for section in sections for n in field_walk({'content':section['blocks']})):
            _warn(warnings, 'This output retains the last calculated page fields. Use Update fields before exporting; desktop PDF recalculates them from its own page layout.')
    if format in ('txt', 'md', 'pdf'):
        from .document_comments import walk as comment_walk
        if any(n.get('type') == 'comment_anchor' for section in sections for n in comment_walk({'content': section['blocks']})):
            _warn(warnings, 'Review comments are omitted from this publication format; retain the Alder or DOCX document for review.')
    omitted = len({c["id"] for c in project.get("clips", [])} - set(ids))
    if omitted:
        _warn(warnings, f"{omitted} draft clip(s) are outside the included collation and were not exported.")
    if len(ids) != len(set(ids)):
        _warn(warnings, "The collation includes repeated draft placements; they are repeated in the output.")
    validation = {"status": "passed", "format": format}
    try:
        if format == "html":
            path.write_text(_render_html(project, settings, warnings), "utf-8")
        elif format in ("txt", "md"):
            pieces = [(("# " if format == "md" else "") + project.get("name", "Untitled") + "\n\n") if settings["includeTitle"] else ""]
            if settings.get("author") and settings["includeTitle"]:
                pieces.append(str(settings["author"]) + "\n\n")
            for section in sections:
                if section["title"]:
                    pieces.append(("## " if format == "md" else "") + section["title"] + "\n\n")
                pieces.append((section["text"] + "\n\n") if format == "txt" else section["rawMarkdown"] + "\n\n" if "rawMarkdown" in section else "".join(_markdown(n, warnings) for n in section["blocks"]))
            output = "".join(pieces)
            # Remove only our final section separator, never authored trailing
            # spaces or empty paragraphs when saving an existing text file.
            if format == "txt":
                if sections: output = output[:-2]
            else:
                output = output.rstrip() + "\n"
            path.write_text(output, "utf-8", newline="")
            def has_rich_content(node):
                if _kind(node) not in ("doc", "paragraph", "text", "hard_break") or node.get("marks"):
                    return True
                if any(value is not None and value != "" and key != "blockId"
                       for key, value in (node.get("attrs") or {}).items()):
                    return True
                return any(has_rich_content(child) for child in node.get("content", []))
            if format == "txt" and any(has_rich_content(block) for section in sections for block in section["blocks"]):
                _warn(warnings, "Plain text preserves wording and order but omits rich formatting, images, links and page settings.")
        elif format == "docx":
            _docx(project, path, settings, warnings)
            with _zip(path) as archive:
                _xml(archive.read("word/document.xml"))
        elif format == "pdf":
            _pdf(project, path, settings, warnings)
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            validation["pages"] = len(reader.pages)
        elif format in ("epub", "azw3"):
            epub_path = path if format == "epub" else path.with_suffix(".epub")
            _epub(project, epub_path, settings, warnings)
            validation = validate_epub(epub_path)
            if validation["status"] == "failed":
                raise ValueError("EPUB validation failed: " + json.dumps(validation, ensure_ascii=False)[:3000])
            if validation["epubcheck"]["status"] == "unavailable":
                _warn(warnings, "EPUB container and XML checks passed. EPUBCheck is unavailable, so full EPUB conformance has not been verified.")
            if format == "azw3":
                converter = capabilities()["azw3"]["converter"]
                if not converter:
                    raise RuntimeError("The bundled AZW3 converter is missing. Repair the Alder installation to restore publishing tools.")
                completed = subprocess.run([converter, str(epub_path), str(path)], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
                                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                if completed.returncode or not path.exists() or path.stat().st_size < 100:
                    raise RuntimeError("Calibre AZW3 conversion failed: " + (completed.stderr + completed.stdout)[-3000:])
                validation = {"status": "converted", "sourceEpub": validation, "converter": "Calibre", "exitCode": 0}
                _warn(warnings, "AZW3 conversion completed through Calibre; reader layout may differ from EPUB and print output.")
                epub_path.unlink(missing_ok=True)
        return {"path": str(path.resolve()), "filename": filename, "mime": MIMES[format], "warnings": warnings, "validation": validation}
    except Exception:
        path.unlink(missing_ok=True)
        if format == "azw3":
            path.with_suffix(".epub").unlink(missing_ok=True)
        raise


def _asset_from_bytes(data, name, assets, warnings):
    probe = _Resources({"assets": []}, {}, warnings).image({"src": "data:image/png;base64," + base64.b64encode(data).decode(), "alt": name})
    if not probe:
        return None
    asset_id = "asset-" + hashlib.sha256(probe["data"]).hexdigest()[:24]
    if not any(a["id"] == asset_id for a in assets):
        assets.append({"id": asset_id, "name": Path(name).stem + ".png", "mime": "image/png", "data": base64.b64encode(probe["data"]).decode()})
    return asset_id


def _import_html(content, warnings, assets, image_reader=None, link_reader=None):
    from bs4 import BeautifulSoup, NavigableString, Tag, Comment, Doctype, ProcessingInstruction
    soup = BeautifulSoup(content, "html.parser")
    if soup.find(["script", "iframe", "object", "embed", "form"]):
        _warn(warnings, "Active HTML content was removed during import.")
    for tag in soup.find_all(["script", "style", "iframe", "object", "embed", "form", "head", "nav"]):
        tag.decompose()
    for comment in soup.find_all(string=lambda text: isinstance(text, (Comment, Doctype, ProcessingInstruction))):
        comment.extract()
    _warn(warnings, "HTML import preserves semantic blocks and inline formatting; external CSS, scripts and exact page layout are not imported.")

    def paragraph_attributes(element):
        from .tab_stops import validate_tab_stops
        css = {key.lower(): value.strip() for key, value in re.findall(r'([\w-]+)\s*:\s*([^;]+)', element.get('style', ''))}
        from .paragraph_appearance import html_appearance
        attrs = html_appearance(css)
        if element.get('data-heading-numbered') in ('true', 'false'): attrs['headingNumbered'] = element['data-heading-numbered'] == 'true'
        if element.get('data-heading-start'): attrs['headingStart'] = int(element['data-heading-start'])
        if element.has_attr('data-revision-paragraph'):
            from .paragraph_revisions import validate_paragraph_log
            log = json.loads(element['data-revision-paragraph']); validate_paragraph_log(log); attrs['revisionProperties'] = log
        if re.fullmatch(r'[A-Za-z0-9_.-]{1,128}', element.get('data-block-id', '')): attrs['blockId'] = element['data-block-id']
        if element.get('data-tab-stops') is not None:
            try: attrs['tabStops'] = validate_tab_stops(json.loads(element['data-tab-stops']))
            except (ValueError, TypeError): pass
        if css.get('text-align') in ('left', 'right', 'center', 'justify'): attrs['align'] = css['text-align']
        for prop, key, minimum, maximum in (('font-size', 'fontSize', 6, 72), ('margin-top', 'spaceBefore', 0, 144), ('margin-bottom', 'spaceAfter', 0, 144), ('margin-left', 'leftIndent', -144, 144), ('margin-right', 'rightIndent', -144, 144), ('text-indent', 'firstLineIndent', -144, 144)):
            match = re.fullmatch(r'(-?[\d.]+)(pt|px)', css.get(prop, ''))
            if match:
                try: value = float(match[1]) * (0.75 if match[2] == 'px' else 1)
                except ValueError: continue
                if minimum <= value <= maximum: attrs[key] = value
        for prop, key in (('break-after', 'keepWithNext'), ('break-inside', 'keepTogether'), ('break-before', 'pageBreakBefore')):
            value = css.get(prop, css.get('page-' + prop, ''))
            if value in ('avoid', 'page', 'always', 'auto'):
                attrs[key] = value == ('page' if key == 'pageBreakBefore' else 'avoid') or (key == 'pageBreakBefore' and value == 'always')
        if css.get('widows', '').isdigit(): attrs['widowOrphan'] = int(css['widows']) > 1
        leading = css.get('line-height', '')
        match = re.fullmatch(r'([\d.]+)(pt|px)', leading)
        minimum = re.fullmatch(r'max\(1\.2em,\s*([\d.]+)pt\)', leading)
        try:
            if match or minimum:
                value = float(match[1]) * (0.75 if match[2] == 'px' else 1) if match else float(minimum[1])
                if 1 <= value <= 144: attrs.update(lineSpacingMode='exact' if match else 'atLeast', lineSpacing=value)
            elif re.fullmatch(r'[\d.]+', leading) and 0.8 <= float(leading) <= 4:
                attrs.update(lineSpacingMode='multiple', lineHeight=float(leading))
        except ValueError: pass
        if css.get('font-family'): attrs['fontFamily'] = re.sub(r"[^\w ,'-]", '', css['font-family'])[:100]
        if re.fullmatch(r'#[\da-fA-F]{6}', css.get('color', '')): attrs['color'] = css['color']
        return attrs

    def inline(element, marks=None):
        if getattr(element, 'attrs', None) and (element.get('data-heading-marker') == 'true' or element.get('data-outline-number') == 'true'): return []
        marks = marks or []
        if isinstance(element, NavigableString):
            text = str(element)
            if not text:
                return []
            return [{"type": "text", "text": text, **({"marks": copy.deepcopy(marks)} if marks else {})}]
        if not isinstance(element, Tag):
            return []
        name = element.name.lower()
        if name == 'span':
            from .document_revisions import validate_revision, validate_break
            from .revision_formatting import validate_format_log
            if element.has_attr('data-revision-break'):
                attrs = json.loads(element.get('data-revision-break-content', 'null')); validate_break(attrs)
                return [{'type': 'revision_break', 'attrs': attrs}]
            for kind in ('insertion', 'deletion'):
                if element.has_attr('data-revision-' + kind):
                    stamp = {'id': element['data-revision-' + kind], 'author': element.get('data-revision-author', ''), 'date': element.get('data-revision-date', '')}
                    validate_revision(stamp); marks = marks + [{'type': 'revision_' + kind, 'attrs': stamp}]
            if element.has_attr('data-revision-format'):
                log = json.loads(element['data-revision-format']); validate_format_log(log)
                marks = marks + [{'type': 'revision_format', 'attrs': log}]
        if name == 'span' and element.get('data-comment-thread'):
            from .document_comments import html_anchor
            anchor = html_anchor(element)
            if not anchor:
                _warn(warnings, 'An invalid HTML comment annotation was not imported.')
            return [anchor] if anchor else []
        from .document_fields import FIELD_LABELS
        if name == 'span' and element.get('data-document-field') in FIELD_LABELS:
            parts = sum((inline(child,marks) for child in element.children),[])
            appearance = next((part.get('marks',[]) for part in parts if part['type']=='text'),[])
            return [{'type':'document_field','attrs':{'fieldKind':element['data-document-field'],'unresolved':element.get('data-field-unresolved')=='true'},'content':[{'type':'text','text':element.get_text()}] if element.get_text() else [],**({'marks':appearance} if appearance else {})}]
        if name == 'span' and re.fullmatch(r'[A-Za-z0-9_.-]{1,128}', element.get('data-cross-reference', '')):
            format = element.get('data-reference-format', 'text')
            return [{'type':'cross_reference', 'attrs':{'targetId':element['data-cross-reference'], 'referenceFormat':format if format in ('text','number','labelNumber','page') else 'text', 'unresolved':element.get('data-reference-unresolved') == 'true'}, 'content':[{'type':'text', 'text':element.get_text()}] if element.get_text() else []}]
        if name == 'span' and element.get('data-document-sequence') in ('Figure','Table') and re.fullmatch(r'[1-9][0-9]{0,5}',element.get_text()):
            return [{'type':'sequence','attrs':{'sequenceName':element['data-document-sequence']},'content':[{'type':'text','text':element.get_text()}]}]
        if name == "span" and re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}", element.get("data-document-bookmark", "")):
            return [{"type":"bookmark","attrs":{"name":element["data-document-bookmark"], 'range': element.get('data-bookmark-range') == 'true'}}]
        if name == 'span' and re.fullmatch(r'[A-Za-z0-9_.-]{1,128}', element.get('data-bookmark-end', '')):
            return [{'type':'bookmark_end','attrs':{'targetId':element['data-bookmark-end'],'name':element.get('data-bookmark-name','')}}]
        if name == "br":
            return [{"type": "hard_break"}]
        if name == "img":
            crop = None
            if element.get('data-image-original') and element.get('data-image-crop'):
                from .image_crop import validate_crop
                try:
                    crop = validate_crop(json.loads(element['data-image-crop']))
                except (ValueError, TypeError):
                    _warn(warnings, 'Invalid image crop metadata was omitted.')
            source = element['data-image-original'] if crop else element.get("src", "")
            def read_image(source):
                try:
                    if source.startswith("data:image/") and len(source) < MAX_IMAGE_BYTES*1.4:
                        return base64.b64decode(source.split(",",1)[1], validate=True)
                    if image_reader and not urlparse(source).scheme and not source.startswith("//"):
                        return image_reader(source)
                except (ValueError, KeyError, OSError):
                    pass
                return None
            data = read_image(source)
            if not data and crop:
                data = read_image(element.get('src', ''))
                crop = None
                _warn(warnings, 'The original image behind a crop was unavailable; the visible cropped image was retained where possible.')
            asset_id = _asset_from_bytes(data, element.get("alt") or "Imported image", assets, warnings) if data else None
            if asset_id:
                attrs = {"assetId": asset_id, "src": "asset:"+asset_id, "alt": element.get("data-image-alt") or element.get("alt", "")}
                if re.fullmatch(r'[A-Za-z0-9_.-]{1,128}',element.get('data-object-id','')): attrs['blockId']=element['data-object-id']
                if crop: attrs['crop'] = crop
                if element.get("data-image-decorative")=="true" or element.get("role")=="presentation": attrs["decorative"]=True
                if element.get('data-image-title') or element.get('title'):
                    attrs['title'] = element.get('data-image-title') or element['title']
                alignment = element.get('data-image-align') or paragraph_attributes(element.parent).get('align')
                if alignment in ('left', 'center', 'right'):
                    attrs['align'] = alignment
                if element.get("width"):
                    attrs["width"] = _number(element["width"], 500, 1, 10000)
                placement = element.get('data-image-placement')
                block_image = placement == 'block' or (placement != 'inline' and element.parent.name == 'figure')
                return [{"type": "image" if block_image else "inline_image", "attrs": attrs}]
            _warn(warnings, "An unavailable or external image was replaced with its alt text; no remote resource was fetched.")
            return [{"type": "text", "text": element.get("alt") or "[Image unavailable]"}]
        mark_type = {"b":"strong", "strong":"strong", "i":"em", "em":"em", "u":"underline", "s":"strike", "del":"strike", "strike":"strike", "code":"code", "sub":"subscript", "sup":"superscript", "mark":"highlight"}.get(name)
        if mark_type:
            marks = marks + [{"type": mark_type}]
        if name == "a":
            link = _safe_link(element.get("href"))
            if not link and link_reader:
                link = link_reader(element.get('href',''))
            if link:
                marks = marks + [{"type":"link", "attrs":{"href":link}}]
        if name == "span" and element.get("style"):
            attrs = {}
            for key, value in re.findall(r"([\w-]+)\s*:\s*([^;]+)", element["style"]):
                key, value = key.strip().lower(), value.strip()
                if key == "color" and re.fullmatch(r"#[\da-fA-F]{6}", value):
                    attrs["color"] = value
                elif key == "font-size" and re.fullmatch(r"[\d.]+pt", value):
                    attrs["fontSize"] = _number(value[:-2], 12, 6, 72)
                elif key == "font-family":
                    attrs["fontFamily"] = re.sub(r"[^\w ,'-]", "", value)[:100]
                elif key == "font-weight" and value in ("bold", "700", "800", "900"):
                    marks = marks + [{"type":"strong"}]
                elif key == "font-style" and value == "italic":
                    marks = marks + [{"type":"em"}]
            if attrs:
                marks = marks + [{"type":"textStyle", "attrs":attrs}]
        return sum((inline(child, marks) for child in element.children), [])

    def blocks(parent):
        output, pending = [], []
        def flush():
            nonlocal pending
            if any(n.get("type") != "text" or n.get("text", "").strip() for n in pending):
                output.append({"type":"paragraph", "content":pending})
            pending = []
        for element in parent.children:
            if isinstance(element, NavigableString):
                if str(element).strip():
                    pending.extend(inline(element))
                continue
            if not isinstance(element, Tag):
                continue
            name = element.name.lower()
            if name in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote", "ul", "ol", "li", "table", "hr", "div", "section", "article", "main", "body", "html", "header", "footer", "figure", "figcaption"):
                flush()
                if name == 'div' and element.get('data-document-contents') == 'true':
                    # Standalone HTML carries resolved entry formatting, not the
                    # project's named-style catalogue.
                    attrs = {'title':element.get('data-contents-title','Contents'), 'depth':int(element.get('data-contents-depth','3')), 'entryStyle':None, 'pageNumbers':element.get('data-contents-pages') != 'false'}
                    attrs['kind'] = element.get('data-contents-kind','headings')
                    from .document_contents import validate_contents
                    validate_contents(attrs)
                    output.append({'type':'table_of_contents','attrs':attrs,'content':blocks(element) or [{'type':'paragraph'}]})
                elif name in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "figcaption"):
                    node = {"type":"heading" if re.fullmatch("h[1-6]",name) else "paragraph", "content":sum((inline(c) for c in element.children), [])}
                    attrs = paragraph_attributes(element)
                    if element.get('data-caption-label') in ('Figure','Table'):
                        attrs['captionLabel'] = element['data-caption-label']
                        if re.fullmatch(r'[A-Za-z0-9_.-]{1,128}',element.get('data-caption-target','')):attrs['captionTargetId']=element['data-caption-target']
                    if node['type'] == 'heading': attrs['level'] = int(name[1])
                    if attrs:
                        node["attrs"] = attrs
                    output.append(node)
                elif name == "pre":
                    output.append({"type":"code_block", "attrs": paragraph_attributes(element), "content":[{"type":"text", "text":element.get_text()}]})
                elif name in ("ul", "ol"):
                    items = [{"type":"list_item", "content":blocks(child) or [{"type":"paragraph"}]} for child in element.find_all("li", recursive=False)]
                    from .list_formatting import NUMBER_FORMATS, BULLETS
                    marker = re.search(r'(?:^|;)\s*list-style-type\s*:\s*([\w-]+)', element.get('style', ''))
                    marker = marker.group(1) if marker else {'a':'lower-alpha', 'A':'upper-alpha', 'i':'lower-roman', 'I':'upper-roman', '1':'decimal'}.get(element.get('type'))
                    attrs = {'order':int(_number(element.get('start'),1,1,1000000000))}
                    if element.get('data-list-template'):
                        from .outline_lists import validate_list_template
                        template=json.loads(element['data-list-template']);validate_list_template(template);attrs['listTemplate']=template
                    if name == 'ol' and element.get('data-continue-list') == 'true': attrs['continuePrevious'] = True
                    if marker in (NUMBER_FORMATS if name == 'ol' else BULLETS): attrs['listStyle'] = marker
                    output.append({"type":"ordered_list" if name == "ol" else "bullet_list", "attrs": attrs, "content":items})
                elif name == "blockquote":
                    output.append({"type":"blockquote", "content":blocks(element)})
                elif name == "table":
                    from .table_formatting import html_table_attributes
                    rows = []
                    for row in element.find_all("tr"):
                        if row.find_parent("table") is not element:
                            continue
                        cells = []
                        for cell in row.find_all(["td", "th"], recursive=False):
                            cells.append({"type":"table_header" if cell.name == "th" else "table_cell", "attrs":{**html_table_attributes('table_cell', cell), "colspan":int(_number(cell.get("colspan"),1,1,100)), "rowspan":int(_number(cell.get("rowspan"),1,1,100))}, "content":blocks(cell) or [{"type":"paragraph"}]})
                        rows.append({"type":"table_row", "attrs":html_table_attributes('table_row', row), "content":cells})
                    table_attrs=html_table_attributes('table', element)
                    if re.fullmatch(r'[A-Za-z0-9_.-]{1,128}',element.get('data-object-id','')):table_attrs['blockId']=element['data-object-id']
                    output.append({"type":"table", "attrs":table_attrs, "content":rows})
                elif name == "hr":
                    output.append({"type":"horizontal_rule"})
                elif name == "div" and element.get("data-section"):
                    from .sections import validate_section
                    try:
                        attrs = json.loads(element['data-section'])
                        validate_section(attrs)
                        output.append({'type':'section_break', 'attrs':attrs})
                    except (ValueError, TypeError, AttributeError):
                        _warn(warnings, 'Invalid section page settings were omitted.')
                elif name == "div" and "page-break" in element.get("class", []):
                    output.append({"type":"page_break"})
                else:
                    output.extend(blocks(element))
            else:
                pending.extend(inline(element))
        flush()
        return output
    return {"type":"doc", "content":blocks(soup) or [{"type":"paragraph"}]}


def _import_docx(path, warnings, assets):
    from docx import Document
    from docx.oxml.ns import qn
    with _zip(path) as archive:
        _xml(archive.read("word/document.xml"))
        from .document_comments import WordCommentsReader
        comment_reader = WordCommentsReader(archive, _xml, warnings)
        names = set(archive.namelist())
        if any(n in names for n in ("word/footnotes.xml", "word/endnotes.xml")):
            _warn(warnings, "DOCX footnotes/endnotes are not imported; inspect the original for additional material.")
    doc = Document(path)
    from .running_headings import import_heading_styles
    heading_styles = import_heading_styles(doc)
    from .document_revisions import WordRevisionReader, restore_review_paragraphs
    from .revision_formatting import run_marks
    revision_reader = WordRevisionReader(doc._element)
    from .docx_support import NumberingReader, paragraph_attributes
    list_numbering = NumberingReader(doc)
    imported_bookmarks = {}
    imported_reference_names = {}
    _warn(warnings, "DOCX import preserves body order, text, basic marks, headings, paragraph flow, ordinary list numbering, tables, comments, supported fields, formatted running paragraphs and embedded raster images; exact layout, complex multilevel labels and unsupported fields are not retained.")
    def paragraph(p):
        content = []
        from .document_references import normalise_docx_fields
        normalise_docx_fields(p._p, warnings)
        def visit(element, inherited=None):
            inherited = inherited or []
            tag = element.tag.rsplit("}",1)[-1]
            if tag == 'pPr': return
            if tag in ('ins', 'del'):
                inherited = inherited + [revision_reader.mark(element)]
            if tag in ('commentRangeStart', 'commentRangeEnd', 'commentReference'):
                anchor = comment_reader.anchor(element)
                if anchor: content.append(anchor)
                return
            if tag == 'fldSimple':
                instruction = element.get(qn('w:instr'), '').strip()
                from .document_fields import read_field
                field_kind = read_field(instruction,heading_styles)
                if field_kind:
                    value = ''.join(t.text or '' for t in element.iter(qn('w:t')))
                    start = len(content)
                    for child in element: visit(child,inherited)
                    marks = next((child.get('marks',[]) for child in content[start:] if child['type']=='text'),[])
                    del content[start:]
                    content.append({'type':'document_field','attrs':{'fieldKind':field_kind},'content':[{'type':'text','text':value}] if value else [],**({'marks':marks} if marks else {})})
                    return
                sequence = re.fullmatch(r'SEQ\s+(Figure|Table)(?:\s+\\\*\s+ARABIC)?\s*',instruction,re.I)
                if sequence:
                    value = ''.join(t.text or '' for t in element.iter(qn('w:t'))).strip()
                    if re.fullmatch(r'[1-9][0-9]{0,5}',value):
                        content.append({'type':'sequence','attrs':{'sequenceName':sequence[1].title()},'content':[{'type':'text','text':value}]})
                        return
                match = re.fullmatch(r'(?:REF|PAGEREF)\s+"?([A-Za-z_][A-Za-z0-9_]{0,39})"?(?P<switches>(?:\s+\\[hw])*)(?:\s+\\\*\s+(?:MERGEFORMAT|CHARFORMAT))?\s*', instruction, re.I)
                if match:
                    value = ''.join(t.text or '' for t in element.iter(qn('w:t')))
                    target = imported_reference_names.get(match[1], match[1])
                    content.append({'type':'cross_reference', 'attrs':{'targetId':target, **({'referenceFormat':'page'} if instruction.upper().startswith('PAGEREF') else {'referenceFormat':'number'} if '\\w' in match['switches'].lower() else {})}, 'content':[{'type':'text','text':value}] if value else []})
                    return
                _warn(warnings, 'An unsupported Word field was imported as its displayed text.')
            if tag == "bookmarkStart":
                name = element.get(qn("w:name"), "")
                original_name = name
                if re.fullmatch(r'_(?:Ref|Toc)[A-Za-z0-9_]{0,35}', name): name = 'Alder' + name[:35]
                if re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,39}",name):
                    identity = 'bookmark_' + uuid.uuid4().hex
                    imported_reference_names[original_name] = identity
                    end = element.getnext()
                    is_range = end is None or end.tag != qn('w:bookmarkEnd') or end.get(qn('w:id')) != element.get(qn('w:id'))
                    content.append({'type':'bookmark','attrs':{'name':name,'blockId':identity,'range':is_range}})
                    imported_bookmarks[element.get(qn('w:id'))] = (identity, name, is_range)
                return
            if tag == 'bookmarkEnd':
                target = imported_bookmarks.get(element.get(qn('w:id')))
                if target and target[2]:
                    content.append({'type':'bookmark_end','attrs':{'targetId':target[0],'name':target[1]}})
                return
            if tag == "hyperlink":
                rid = element.get(qn("r:id"))
                link = p.part.rels.get(rid)
                href = _safe_link(link.target_ref) if link else _safe_link('#'+element.get(qn('w:anchor'))) if element.get(qn('w:anchor')) else ""
                if href:
                    inherited = inherited + [{"type":"link", "attrs":{"href":href}}]
            if tag == "r":
                marks = copy.deepcopy(inherited)
                props = element.find(qn("w:rPr"))
                marks.extend(run_marks(props))
                formatting = revision_reader.formatting(props)
                if formatting:
                    # Hyperlink identity lives outside rPr and is unchanged by a
                    # property revision; retain it in both projected appearances.
                    formatting['attrs']['base'] += [copy.deepcopy(m) for m in inherited if m['type'] == 'link']
                    marks.append(formatting)
                for child in element:
                    child_tag = child.tag.rsplit("}",1)[-1]
                    if child_tag in ("t", "delText") and child.text:
                        content.append({"type":"text","text":child.text, **({"marks":marks} if marks else {})})
                    elif child_tag in ("br", "cr"):
                        content.append({"type":"hard_break", **({'marks':marks} if marks else {})})
                    elif child_tag == "tab":
                        content.append({"type":"text","text":"\t", **({'marks':marks} if marks else {})})
                    elif child_tag == "drawing":
                        if next(child.iter(qn('wp:anchor')), None) is not None:
                            _warn(warnings, 'A floating Word picture was placed in line with text; its original wrapping and page position were not retained.')
                        crop = next(child.iter(qn('a:srcRect')), None)
                        from .image_crop import validate_crop
                        try:
                            image_crop = validate_crop({edge: float(crop.get(short, '0'))/1000 for edge, short in (('left','l'),('top','t'),('right','r'),('bottom','b'))}) if crop is not None else None
                        except (ValueError, TypeError):
                            image_crop = None
                            _warn(warnings, 'A Word picture with an unsupported crop/padded image area was imported uncropped.')
                        for blip in child.iter(qn("a:blip")):
                            rid = blip.get(qn("r:embed"))
                            if rid and rid in p.part.related_parts:
                                part = p.part.related_parts[rid]
                                asset_id = _asset_from_bytes(part.blob, Path(str(part.partname)).name, assets, warnings)
                                if asset_id:
                                    properties = next(child.iter(qn('wp:docPr')), None)
                                    attrs = {'assetId': asset_id, 'src': 'asset:'+asset_id, 'alt': properties.get('descr', '') if properties is not None else ''}
                                    if image_crop: attrs['crop'] = image_crop
                                    from .image_accessibility import read_decorative
                                    if read_decorative(properties): attrs['decorative'] = True
                                    if properties is not None and properties.get('title'):
                                        attrs['title'] = properties.get('title')
                                    extent = next(child.iter(qn('wp:extent')), None)
                                    if extent is not None and extent.get('cx'):
                                        attrs['width'] = _number(extent.get('cx'), 9525, 9525, 95250000) / 9525
                                    block_image = properties is not None and properties.get('name', '').startswith('Alder block image ')
                                    content.append({'type': 'image' if block_image else 'inline_image', 'attrs': attrs})
                return
            for child in element:
                visit(child, inherited)
        visit(p._p)
        style = p.style.name if p.style else ""
        heading = re.match(r"Heading (\d)", style)
        outline = p._p.pPr.find(qn('w:outlineLvl')) if p._p.pPr is not None else None
        heading_level = int(heading.group(1)) if heading else (int(outline.get(qn('w:val'))) + 1 if outline is not None else 0)
        heading = heading_level if 1 <= heading_level <= 6 else 0
        node = {"type":"heading" if heading else "paragraph", "content":content}
        attrs = paragraph_attributes(p, warnings)
        from .paragraph_revisions import read_paragraph_revision
        reviewed_properties = read_paragraph_revision(p, revision_reader, warnings)
        if reviewed_properties: attrs['revisionProperties'] = reviewed_properties
        boundary = revision_reader.paragraph_break(p._p)
        if boundary: attrs['revisionBreak'] = boundary
        if style == 'Caption':
            sequences=[node for node in content if node['type']=='sequence']
            if len(sequences)==1:attrs['captionLabel']=sequences[0]['attrs']['sequenceName']
        if heading: attrs["level"] = heading
        if p.alignment is not None:
            attrs["align"] = {0:"left",1:"center",2:"right",3:"justify"}.get(int(p.alignment),"left")
        if attrs:
            node["attrs"] = attrs
        return node
    doc_sections = list(doc.sections)
    section_index = 0
    def append_section(child, output):
        nonlocal section_index
        if child._p.pPr is not None and child._p.pPr.sectPr is not None:
            section_index += 1
            from .page_geometry import docx_page_settings
            from docx.enum.section import WD_SECTION
            following = doc_sections[section_index]
            kind = {WD_SECTION.NEW_PAGE: 'nextPage',WD_SECTION.CONTINUOUS:'continuous', WD_SECTION.ODD_PAGE: 'oddPage', WD_SECTION.EVEN_PAGE: 'evenPage'}.get(following.start_type)
            if kind is None:
                kind = 'nextPage'
                _warn(warnings, 'Column section breaks are converted to next-page sections.')
            from .running_matter import read_running_matter
            output.append({'type':'section_break', 'attrs':{'kind':kind, 'page':docx_page_settings(doc, warnings, following), 'runningMatter':read_running_matter(following,paragraph,warnings), **read_docx_numbering(following, warnings)}})
    def container(item):
        nonlocal section_index
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        output, stack = [], []
        previous_lists = {}
        root = item._element.body if item is doc else item._element
        def children(root):
            for element in root:
                if element.tag == qn('w:p'): yield Paragraph(element, item)
                elif element.tag == qn('w:tbl'): yield Table(element, item)
                elif element.tag == qn('w:sdt'):
                    from .word_contents import read_contents
                    generated = read_contents(element, item, paragraph, warnings)
                    if generated is not None: yield generated
                    else:
                        nested = element.find(qn('w:sdtContent'))
                        if nested is not None: yield from children(nested)
        for child in children(root):
            if isinstance(child, dict):
                stack = []; output.append(child); continue
            if isinstance(child, Paragraph):
                node = paragraph(child)
                info = list_numbering.info(child)
                if node['type'] == 'heading' and info:
                    from .heading_numbering import read_heading_numbering
                    numbered_heading = read_heading_numbering(list_numbering, child, info)
                    if numbered_heading:
                        node.setdefault('attrs', {}).update(numbered_heading)
                        info = None
                section_end = child._p.pPr is not None and child._p.pPr.sectPr is not None and item is doc
                if not info:
                    stack = []
                    if not section_end or node.get('content'):
                        output.append(node)
                    if section_end:
                        append_section(child, output)
                    continue
                number, level, kind, start, marker = info
                from .outline_lists import read_word_template
                outline_template = read_word_template(list_numbering,number)
                if outline_template:
                    props=node.setdefault('attrs',{});spec=outline_template[level]
                    if props.get('leftIndent') is not None:props['leftIndent']-=spec['indent']
                    if props.get('firstLineIndent') is not None:props['firstLineIndent']+=spec['hanging']
                while stack and (stack[-1][0] > level or (stack[-1][0] == level and stack[-1][1] != number)):
                    stack.pop()
                if not stack or stack[-1][0] != level:
                    target = stack[-1][2]["content"][-1]["content"] if stack else output
                    listing = {"type": kind, "content": []}
                    listing['attrs'] = {'listStyle': marker, **({'order': start} if kind == 'ordered_list' else {})}
                    if outline_template and level==0:listing['attrs']['listTemplate']=outline_template
                    scope = (id(target), marker)
                    if kind == 'ordered_list' and previous_lists.get(scope) == number:
                        listing['attrs']['continuePrevious'] = True
                    if kind == 'ordered_list': previous_lists[scope] = number
                    target.append(listing)
                    stack.append((level, number, listing))
                stack[-1][2]["content"].append({"type": "list_item", "content": [node]})
                if section_end:
                    append_section(child, output)
                    stack = []
            elif isinstance(child, Table):
                stack = []
                from .docx_support import import_table
                output.append(import_table(child, container, warnings))
        return output
    from .page_geometry import docx_page_settings
    imported = {"type":"doc", "content":container(doc) or [{"type":"paragraph"}]}
    comment_reader.finish(imported)
    restore_review_paragraphs(imported)
    def bind_fields(node):
        if node['type'] == 'cross_reference':
            attrs = node['attrs']; attrs['targetId'] = imported_reference_names.get(attrs['targetId'], attrs['targetId'])
        for mark in node.get('marks', []):
            href = mark.get('attrs', {}).get('href', '')
            if mark['type'] == 'link' and href.startswith(('#_Toc', '#_Ref')) and href[1:] in imported_reference_names:
                mark['attrs']['href'] = '#' + imported_reference_names[href[1:]]
        for child in node.get('content', []): bind_fields(child)
    bind_fields(imported)
    from .document_captions import associate_captions
    associate_captions(imported)
    settings = docx_page_settings(doc, warnings)
    from .running_matter import read_running_matter
    settings['runningMatter']=read_running_matter(doc.sections[0],paragraph,warnings,first=True)
    settings.update(author=doc.core_properties.author or '',subject=doc.core_properties.subject or '',description=doc.core_properties.comments or '')
    return imported, doc.core_properties.title or path.stem, settings


def _normalise_import_blocks(document):
    """Lift explicit standalone drawings; retain inline pictures inside their text."""
    def blocks(nodes):
        result = []
        for source in nodes:
            node = copy.deepcopy(source)
            if _kind(node) in ("paragraph", "heading") and any(_kind(n) == "image" for n in node.get("content", [])):
                pending = []
                for child in node.get("content", []):
                    if _kind(child) == "image":
                        inherited_align = node.get('attrs', {}).get('align')
                        if inherited_align in ('left', 'center', 'right'):
                            child.setdefault('attrs', {}).setdefault('align', inherited_align)
                        if pending:
                            result.append({**node,"content":pending})
                            pending = []
                        result.append(child)
                    else:
                        pending.append(child)
                if pending:
                    result.append({**node,"content":pending})
            else:
                if _kind(node) not in ("paragraph", "heading", "text", "code_block") and "content" in node:
                    node["content"] = blocks(node["content"])
                if _kind(node) == "list_item" and (not node.get("content") or _kind(node["content"][0]) != "paragraph"):
                    node["content"] = [{"type":"paragraph"}] + node.get("content", [])
                result.append(node)
        return result
    return {**document,"content":blocks(document.get("content", [])) or [{"type":"paragraph"}]}


def import_document(path: Path) -> dict:
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_IMPORT_BYTES:
        raise ValueError("The import file is missing or exceeds 100 MB.")
    extension = path.suffix.lower()
    warnings, assets = [], []
    title = path.stem
    raw_source = None
    imported_settings = None
    if extension in (".txt", ".md", ".markdown", ".html", ".htm"):
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252", errors="replace")
            _warn(warnings, "Input was not UTF-8; it was decoded as Windows-1252. Check accented characters.")
        if extension == ".txt":
            document = text_document(text)
        else:
            if extension in (".md", ".markdown"):
                raw_source = {"format":"markdown", "text":text}
                from markdown_it import MarkdownIt
                text = MarkdownIt("commonmark", {"html": False}).render(text)
            document = _import_html(text, warnings, assets)
            if extension in ('.html','.htm'):
                from bs4 import BeautifulSoup
                metadata = BeautifulSoup(text,'html.parser').head
                if metadata:
                    if metadata.title and metadata.title.get_text().strip(): title = metadata.title.get_text().strip()[:500]
                    imported_settings = {}
                    from .running_matter import read_html_metadata
                    running = read_html_metadata(text)
                    if running is not None: imported_settings['runningMatter']=running
                    for key in ('author','subject','description'):
                        entry = metadata.find('meta', attrs={'name':key})
                        if entry and entry.get('content') is not None: imported_settings[key] = entry['content'][:10000]
    elif extension == ".docx":
        document, title, imported_settings = _import_docx(path, warnings, assets)
    elif extension == ".epub":
        with _zip(path) as archive:
            container = _xml(archive.read("META-INF/container.xml"))
            opf_path = container.find(".//{*}rootfile").attrib["full-path"]
            package = _xml(archive.read(opf_path))
            manifest = {i.attrib["id"]: i.attrib for i in package.findall("{*}manifest/{*}item")}
            metadata_title = package.find("{*}metadata/{*}title")
            title = metadata_title.text if metadata_title is not None and metadata_title.text else title
            imported_settings = {}
            for tag,key in (('creator','author'),('subject','subject'),('description','description')):
                entries = [item.text for item in package.findall('{*}metadata/{*}'+tag) if item.text]
                if entries: imported_settings[key] = '; '.join(entries)[:10000]
            blocks = []
            import posixpath
            content_paths = {posixpath.normpath(str(PurePosixPath(opf_path).parent / unquote(item['href'].split('#',1)[0]))) for item in manifest.values() if item.get('media-type')=='application/xhtml+xml'}
            for ref in package.findall("{*}spine/{*}itemref"):
                if ref.attrib.get("linear", "yes") == "no":
                    _warn(warnings, "EPUB non-linear supplementary spine items were omitted.")
                    continue
                item = manifest.get(ref.attrib.get("idref"))
                if not item:
                    raise ValueError("EPUB spine references an absent manifest item.")
                resource = PurePosixPath(opf_path).parent / unquote(item["href"].split("#",1)[0])
                if ".." in resource.parts:
                    raise ValueError("EPUB content path escapes its package.")
                raw = archive.read(str(resource))
                from .running_matter import read_html_metadata
                running = read_html_metadata(raw.decode('utf-8-sig'))
                if running is not None: imported_settings['runningMatter']=running
                # HTML doctype is common and has no effect with BeautifulSoup's HTML parser.
                def image_reader(src, root=resource.parent):
                    candidate = root / unquote(src.split("#",1)[0])
                    # EPUB commonly uses ../Images/; normalize within archive, never extract.
                    parts = []
                    for part in candidate.parts:
                        if part == "..":
                            if not parts:
                                raise ValueError("Unsafe EPUB image path")
                            parts.pop()
                        elif part != ".":
                            parts.append(part)
                    return archive.read("/".join(parts))
                def link_reader(href, root=resource.parent):
                    parsed=urlparse(href)
                    if parsed.scheme or parsed.netloc or not parsed.fragment: return ''
                    destination=posixpath.normpath(str(root / unquote(parsed.path)))
                    return '#'+unquote(parsed.fragment) if destination in content_paths else ''
                blocks.extend(_import_html(raw.decode("utf-8-sig"), warnings, assets, image_reader, link_reader)["content"])
            document = {"type":"doc", "content":blocks or [{"type":"paragraph"}]}
        _warn(warnings, "EPUB was imported in spine reading order; title, author, subject and description are retained, while other publication metadata, CSS and separate section identities are not imported.")
    else:
        from .text_import import extract_text
        text, notices = extract_text(path)
        warnings.extend(notices)
        document = text_document(text)
    document = _normalise_import_blocks(document)
    from .document_captions import restore_caption_references
    restore_caption_references(document)
    from .word_contents import restore_contents_links
    restore_contents_links(document)
    return {"title":title, "document":document, "text":plain_text(document), "warnings":warnings, "assets":assets, **({"settings": imported_settings} if imported_settings else {}), **({"rawSource":raw_source, "rawFormat":"markdown"} if raw_source else {})}
