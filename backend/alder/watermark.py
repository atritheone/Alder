"""Page-centred editable text watermarks, separate from running/body wording."""
import copy,math,re
V='urn:schemas-microsoft-com:vml'
O='urn:schemas-microsoft-com:office:office'
W10='urn:schemas-microsoft-com:office:word'
PREFIX='PowerPlusWaterMarkObject'

def validate_watermark(value):
    if value is None:return
    if not isinstance(value,dict) or set(value)!={'text','fontFamily','color','widthMm','heightMm','angle'}:raise ValueError('Invalid watermark.')
    if not isinstance(value['text'],str) or not value['text'].strip() or len(value['text'])>120 or re.search(r'[\x00-\x1f\x7f]',value['text']):raise ValueError('Watermark text must be one line of 1–120 characters.')
    if not isinstance(value['fontFamily'],str) or not 0<len(value['fontFamily'])<=100 or re.search(r'[\x00-\x1f;<>"\\]',value['fontFamily']):raise ValueError('Invalid watermark font.')
    if not isinstance(value['color'],str) or not re.fullmatch(r'#[0-9a-fA-F]{6}',value['color']):raise ValueError('Invalid watermark colour.')
    for key in ('widthMm','heightMm'):
        if type(value[key]) not in (int,float) or not math.isfinite(value[key]) or not 1<=value[key]<=1000:raise ValueError('Watermark dimensions must be 1–1000 mm.')
    if type(value['angle']) is not int or not 0<=value['angle']<360:raise ValueError('Watermark angle must be a whole number from 0 to 359.')

def _style(value):
    return {k.strip().lower():v.strip() for part in (value or '').split(';') if ':' in part for k,v in [part.split(':',1)]}

def _shapes(target):
    # Resolving _element on an absent header creates it. Inspection must not
    # materialise empty stories or change link-to-previous semantics.
    while target is not None and target.is_linked_to_previous:target=target._prior_headerfooter
    if target is None:return []
    return [shape for shape in target._element.iter('{'+V+'}shape') if shape.get('id','').startswith(PREFIX)]

def write_watermarks(document,values):
    from docx.oxml import OxmlElement
    from lxml import etree
    for index,(section,value) in enumerate(zip(document.sections,values)):
        for variant in ('header','first_page_header','even_page_header'):
            target=getattr(section,variant)
            # Avoid creating unused headers when neither this section nor an
            # inherited header contains a watermark.
            if value is None and (target.is_linked_to_previous and index==0):continue
            if value is None and not _shapes(target):continue
            content=copy.deepcopy(list(target._element))
            target.is_linked_to_previous=False
            for old in list(target._element):target._element.remove(old)
            for child in content:target._element.append(child)
            for shape in _shapes(target):shape.getparent().remove(shape)
            if value is None:continue
            validate_watermark(value)
            paragraph=target.paragraphs[0] if target.paragraphs else target.add_paragraph()
            run=paragraph.add_run();run._r.get_or_add_rPr().append(OxmlElement('w:noProof'))
            pict=OxmlElement('w:pict');run._r.append(pict)
            shape=etree.SubElement(pict,'{'+V+'}shape',nsmap={'v':V,'o':O,'w10':W10})
            shape.set('id',f'{PREFIX}Alder{index}{variant}');shape.set('{'+O+'}spt','136');shape.set('coordsize','21600,21600');shape.set('path','m0,0l21600,0m0,21600l21600,21600e')
            shape.set('style',f"position:absolute;margin-left:0;margin-top:0;width:{value['widthMm']*72/25.4:.6f}pt;height:{value['heightMm']*72/25.4:.6f}pt;rotation:{value['angle']};z-index:-251658240;mso-position-horizontal:center;mso-position-horizontal-relative:page;mso-position-vertical:center;mso-position-vertical-relative:page")
            shape.set('fillcolor',value['color']);shape.set('stroked','f')
            path=etree.SubElement(shape,'{'+V+'}path');path.set('textpathok','t')
            text=etree.SubElement(shape,'{'+V+'}textpath');text.set('on','t');text.set('fitshape','t');text.set('trim','t');text.set('string',value['text']);text.set('style','font-family:"'+value['fontFamily'].replace('"','')+'";font-size:100pt;v-text-kern:t')
            wrap=etree.SubElement(shape,'{'+W10+'}wrap');wrap.set('anchorx','page');wrap.set('anchory','page')

def read_watermark(section,warnings):
    found=[]
    # A watermark can differ by header variant in Word. Alder's section value
    # represents the common value only; report unsupported differences.
    variants=['header']
    if section.different_first_page_header_footer:variants.append('first_page_header')
    if section._document_part.settings.odd_and_even_pages_header_footer:variants.append('even_page_header')
    try:
        for variant in variants:
            target=getattr(section,variant);shapes=_shapes(target)
            if not shapes:found.append(None);continue
            if len(shapes)!=1:raise ValueError()
            shape=shapes[0];style=_style(shape.get('style'));text=shape.find('{'+V+'}textpath')
            if text is None:raise ValueError()
            if any(style.get(k)!=v for k,v in {'mso-position-horizontal':'center','mso-position-horizontal-relative':'page','mso-position-vertical':'center','mso-position-vertical-relative':'page'}.items()):raise ValueError()
            if int(style.get('z-index','0'))>=0 or shape.get('stroked') not in ('f','false','0'):raise ValueError()
            fill=shape.find('{'+V+'}fill')
            if fill is not None and fill.get('opacity','1') not in ('1','65536f','100%'):raise ValueError()
            size=lambda key:float(style[key][:-2])*25.4/72 if style[key].endswith('pt') else (_ for _ in ()).throw(ValueError())
            value={'text':text.get('string',''),'fontFamily':_style(text.get('style')).get('font-family','Cambria').strip('"\''),'color':shape.get('fillcolor','#d0d0d0').lower(),'widthMm':round(size('width'),4),'heightMm':round(size('height'),4),'angle':int(float(style.get('rotation','0')))%360}
            validate_watermark(value);found.append(value)
        if any(value!=found[0] for value in found):raise ValueError()
        return found[0]
    except (ValueError,TypeError,KeyError,OverflowError):
        warnings.append('Unsupported DOCX watermark placement, effects, image or header variants were omitted; the original file is unchanged.')
        return None
