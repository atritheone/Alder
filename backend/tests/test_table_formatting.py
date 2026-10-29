from pathlib import Path
import pytest
from docx import Document
from docx.oxml.ns import qn
from alder.models import create_project, validate_project, validate_document, ValidationError
from alder.publishing import build_export, import_document


def fixture():
    project = create_project(template='blank')
    project['settings'].update(documentKind='docx', includeTitle=False, footer=False)
    appearance = dict(verticalAlign='bottom', paddingTop=2, paddingBottom=4, paddingLeft=0, paddingRight=6, borderWidth=1.5, borderColor='#123456', background='#ddeeff')
    def cell(text, width):
        return {'type': 'table_cell', 'attrs': {**appearance, 'colwidth': [width]}, 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}]}
    table = {'type': 'table', 'attrs': {'align': 'right'}, 'content': [{'type': 'table_row', 'attrs': {'minHeight': 42}, 'content': [cell('One', 100), cell('Two', 150)]}, {'type': 'table_row', 'content': [cell('Three', 100), cell('Four', 150)]}]}
    project['book'] = {'version': 1, 'chapters': [{'id': 'test-chapter', 'title': 'Test', 'text': '', 'document': {'type': 'doc', 'content': [table]}, 'include': True}]}
    return validate_project(project), appearance


@pytest.mark.parametrize('format', ['docx', 'html'])
def test_table_appearance_survives_exchange(tmp_path, format):
    project, appearance = fixture()
    exported = build_export(project, format, tmp_path)
    imported = import_document(Path(exported['path']))
    table = next(node for node in imported['document']['content'] if node['type'] == 'table')
    assert table['attrs']['align'] == 'right'
    assert table['content'][0]['attrs']['minHeight'] == 42
    for row in table['content']:
        for column, cell in enumerate(row['content']):
            for key, expected in appearance.items():
                assert cell['attrs'][key] == expected
            assert cell['attrs']['colwidth'] == [100 if column == 0 else 150]
    if format == 'docx':
        doc = Document(exported['path'])
        assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:trHeight')).get(qn('w:hRule')) == 'atLeast'


def test_table_pdf_contains_all_text_and_custom_shading(tmp_path):
    from pypdf import PdfReader
    project, _ = fixture()
    path = build_export(project, 'pdf', tmp_path)['path']
    pdf = PdfReader(path)
    assert len(pdf.pages) == 1
    text = pdf.pages[0].extract_text()
    for word in ('One', 'Two', 'Three', 'Four'): assert text.count(word) == 1
    assert b'.866667 .933333 1 rg' in pdf.pages[0].get_contents().get_data()


@pytest.mark.parametrize('format', ['docx', 'html', 'epub'])
def test_row_break_preferences_survive_exchange(tmp_path, format):
    project, _ = fixture()
    rows = project['book']['chapters'][0]['document']['content'][0]['content']
    rows[0]['attrs']['allowSplit'] = True
    rows[1]['attrs'] = {'allowSplit': False}
    exported = build_export(project, format, tmp_path)
    imported = import_document(Path(exported['path']))
    table = next(n for n in imported['document']['content'] if n['type'] == 'table')
    assert [r['attrs']['allowSplit'] for r in table['content']] == [True, False]
    if format == 'docx':
        doc = Document(exported['path'])
        assert doc.tables[0].rows[0]._tr.trPr is None or doc.tables[0].rows[0]._tr.trPr.find(qn('w:cantSplit')) is None
        assert doc.tables[0].rows[1]._tr.trPr.find(qn('w:cantSplit')) is not None
        # Word's explicit false form has the same meaning as an absent flag.
        doc.tables[0].rows[1]._tr.trPr.find(qn('w:cantSplit')).set(qn('w:val'), 'false')
        doc.save(exported['path'])
        imported = import_document(Path(exported['path']))
        assert next(n for n in imported['document']['content'] if n['type']=='table')['content'][1]['attrs']['allowSplit'] is True


def test_invalid_row_break_preference_rejected():
    with pytest.raises(ValidationError):
        validate_document({'type': 'doc', 'content': [{'type': 'table_row', 'attrs': {'allowSplit': 'false'}, 'content': []}]})


def test_pdf_row_breaks_preserve_text_and_use_remaining_page_space(tmp_path):
    import re
    from pypdf import PdfReader
    project, _ = fixture()
    project['settings'].update(pageSize='Custom',pageWidthMm=150,pageHeightMm=100,marginMm=10,fontSize=12,lineHeight=1)
    def paragraph(text):
        return {'type':'paragraph','attrs':{'spaceBefore':0,'spaceAfter':0,'widowOrphan':False},'content':[{'type':'text','text':text}]}
    for allow in (False, True):
        content = [paragraph('INTRO'+str(i)) for i in range(10)]
        content.append({'type':'table','content':[{'type':'table_row','attrs':{'allowSplit':allow},'content':[
            {'type':'table_cell','content':[paragraph(prefix+str(i)) for i in range(12)]} for prefix in ('LEFT','RIGHT')]}]})
        project['book']['chapters'][0]['document']['content'] = content
        pdf=PdfReader(build_export(project,'pdf',tmp_path/str(allow))['path'])
        pages=[page.extract_text() for page in pdf.pages]
        assert ('LEFT0' in pages[0]) is allow
        for prefix in ('LEFT','RIGHT'):
            for i in range(12):
                assert re.findall(r'\b(?:LEFT|RIGHT)\d+\b','\n'.join(pages)).count(prefix+str(i))==1


def test_merged_docx_cells_style_the_continuation_rows(tmp_path):
    from alder.table_formatting import read_docx_cell
    from docx.table import _Cell
    project, appearance = fixture()
    rows = project['book']['chapters'][0]['document']['content'][0]['content']
    rows[0]['content'][0]['attrs']['rowspan'] = 2
    rows[1]['content'].pop(0)
    doc = Document(build_export(project, 'docx', tmp_path)['path'])
    table = doc.tables[0]
    continuation = table.rows[1]._tr.tc_lst[0]
    assert continuation.vMerge == 'continue'
    attrs = read_docx_cell(_Cell(continuation, table))
    for key, value in appearance.items(): assert attrs[key] == value


@pytest.mark.parametrize('attrs', [{'paddingLeft': -1}, {'borderWidth': True}, {'background': 'red'}, {'verticalAlign': 'baseline'}, {'colwidth': [100, 100]}, {'rowspan': 0}])
def test_invalid_cell_properties_rejected(attrs):
    with pytest.raises(ValidationError):
        validate_document({'type': 'doc', 'content': [{'type': 'table_cell', 'attrs': attrs, 'content': [{'type': 'paragraph'}]}]})
