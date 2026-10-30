import copy
from pathlib import Path
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from alder.models import create_project, validate_project, validate_document, document_text, ValidationError
from alder.form_fields import properties, field_value, W14
from alder.publishing import build_export, import_document


def fixture():
    project=create_project(template='blank')
    project['settings'].update(documentKind='docx',includeTitle=False)
    content=[]
    for index,kind in enumerate(('plain','rich','checkbox','date','plain','date')):
        attrs=properties(dict(kind=kind,label='FIELD_'+str(index),tag='tag_'+str(index),
            blockId='field_'+str(index),lockContent=index==0,lockDelete=index in (0,1),
            checked=kind=='checkbox',dateValue='2026-10-30' if index==3 else ''))
        value=[{'type':'text','text':'VALUE_'+str(index)}] if index<2 else []
        if kind=='rich':value=[{'type':'text','text':'RICH_ALPHA','marks':[{'type':'strong'}]},
            {'type':'hard_break'},{'type':'text','text':'RICH_OMEGA','marks':[{'type':'em'},{'type':'link','attrs':{'href':'https://example.com/'}}]}]
        field={'type':'form_field','attrs':attrs,'content':field_value(attrs,value)}
        content.append({'type':'paragraph','content':[{'type':'text','text':kind+': '},field,{'type':'text','text':' after.'}]})
    project['book']={'version':1,'chapters':[{'id':'main','title':'Forms','include':True,'document':{'type':'doc','content':content}}]}
    return validate_project(project)


def fields(document):
    found=[]
    def walk(node):
        if node['type'] in ('form_field','form_block'):found.append(node)
        for child in node.get('content',[]):walk(child)
    walk(document)
    return found


@pytest.mark.parametrize('format',['docx','html','epub'])
def test_editable_fields_round_trip_values_formatting_and_restrictions(tmp_path,format):
    project=fixture();before=copy.deepcopy(project)
    result=build_export(project,format,tmp_path)
    restored=import_document(Path(result['path']))
    old=fields(project['book']['chapters'][0]['document']);new=fields(restored['document'])
    assert len(new)==len(old),restored['warnings']
    for expected,actual in zip(old,new):
        assert document_text(actual)==document_text(expected)
        for key in ('kind','label','tag','lockContent','lockDelete','checked','dateValue'):
            # Irrelevant checkbox/date values are not exported on other kinds.
            assert actual['attrs'][key]==expected['attrs'][key],(key,actual)
    assert 'strong' in [m['type'] for m in new[1]['content'][0]['marks']]
    assert 'link' in [m['type'] for m in new[1]['content'][-1]['marks']]
    assert project==before
    if format=='docx':
        word=Document(result['path'])
        assert len(list(word.element.iter(qn('w:sdt'))))==6
        assert len(list(word.element.iter(f'{{{W14}}}checkbox')))==1
        assert [n.get(qn('w:val')) for n in word.element.iter(qn('w:lock'))]==['sdtContentLocked','sdtLocked']


@pytest.mark.parametrize('attrs,content',[
    ({'kind':'bad'},[]),({'label':'x\n'},[]),({'lockContent':'yes'},[]),
    ({'kind':'date','dateValue':'2026-02-30'},[]),({'kind':'date','dateValue':'2026-01-01'},[]),
    ({'kind':'checkbox','checked':True},[{'type':'text','text':'not checked'}]),
    ({'kind':'plain'},[{'type':'text','text':'word','marks':[{'type':'strong'}]}]),
    ({'kind':'rich'},[{'type':'inline_image','attrs':{'src':'x'}}]),
    ({'kind':'rich'},[{'type':'form_field','attrs':{},'content':[]}]),
])
def test_invalid_fields_never_enter_projects(attrs,content):
    doc={'type':'doc','content':[{'type':'paragraph','content':[{'type':'form_field','attrs':attrs,'content':content}]}]}
    with pytest.raises(ValidationError):validate_document(doc)


@pytest.mark.parametrize('format',['rtf','odt','md'])
def test_flat_formats_report_loss_of_form_behaviour(tmp_path,format):
    result=build_export(fixture(),format,tmp_path)
    assert any('field' in w.lower() and ('restrictions' in w or 'editable' in w) for w in result['warnings'])
    imported=import_document(Path(result['path']))
    assert 'VALUE_0' in document_text(imported['document'])
    assert 'RICH_ALPHA' in document_text(imported['document'])


def test_word_placeholder_is_not_authored_text_and_binding_is_reported(tmp_path):
    result=build_export(fixture(),'docx',tmp_path)
    word=Document(result['path']);sdt=list(word.element.iter(qn('w:sdt')))[4]
    pr=sdt.find(qn('w:sdtPr'));pr.append(OxmlElement('w:showingPlcHdr'));pr.append(OxmlElement('w:dataBinding'))
    body=sdt.find(qn('w:sdtContent'));run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='Click here to enter text.';run.append(text);body.append(run)
    word.save(result['path'])
    imported=import_document(Path(result['path']))
    assert document_text(fields(imported['document'])[4])==''
    assert 'Click here' not in document_text(imported['document'])
    assert any('XML binding' in w for w in imported['warnings'])


def test_tracked_word_field_refuses_silent_review_loss(tmp_path):
    result=build_export(fixture(),'docx',tmp_path);word=Document(result['path'])
    body=next(word.element.iter(qn('w:sdtContent')));ins=OxmlElement('w:ins');ins.set(qn('w:id'),'1');ins.set(qn('w:author'),'Editor')
    ins.append(body[0]);body.append(ins);word.save(result['path'])
    with pytest.raises(ValueError,match='Tracked changes inside a form field|revision contains structure or objects'):import_document(Path(result['path']))


def block_fixture():
    import io,base64
    from PIL import Image
    p=fixture();data=io.BytesIO();Image.new('RGB',(80,40),(30,60,90)).save(data,format='PNG')
    p['assets']=[{'id':'form_picture','name':'form.png','mime':'image/png','data':base64.b64encode(data.getvalue()).decode()}]
    paragraph=lambda value:{'type':'paragraph','content':[{'type':'text','text':value}]}
    body=[paragraph('BLOCK_ALPHA'),{'type':'ordered_list','attrs':{'order':3},'content':[
        {'type':'list_item','content':[paragraph('BLOCK_LIST_ONE')]},
        {'type':'list_item','content':[paragraph('BLOCK_LIST_TWO')]}]},
        {'type':'table','content':[{'type':'table_row','content':[
            {'type':'table_cell','content':[paragraph('BLOCK_CELL_ONE')]},
            {'type':'table_cell','content':[paragraph('BLOCK_CELL_TWO')]}]}]},
        {'type':'paragraph','content':[{'type':'inline_image','attrs':{'assetId':'form_picture','width':80,'alt':'Form picture'}},
            {'type':'text','text':'BLOCK_OMEGA','marks':[{'type':'text_style','attrs':{'fontSize':14,'color':'#253f60'}}]}]}]
    p['book']['chapters'][0]['document']['content'].append({'type':'form_block','attrs':properties({'kind':'rich','blockId':'block_form','label':'Rich section','lockContent':True,'lockDelete':True}),'content':body})
    return validate_project(p)


@pytest.mark.parametrize('format',['docx','html','epub'])
def test_rich_block_control_keeps_paragraphs_lists_tables_images_and_locks(tmp_path,format):
    project=block_fixture();before=copy.deepcopy(project)
    result=build_export(project,format,tmp_path)
    imported=import_document(Path(result['path']));restored=fields(imported['document'])[-1]
    assert restored['type']=='form_block',imported['warnings']
    assert restored['attrs']['lockContent'] is True and restored['attrs']['lockDelete'] is True
    assert [n['type'] for n in restored['content']]==['paragraph','ordered_list','table','paragraph']
    for word in ('BLOCK_ALPHA','BLOCK_LIST_ONE','BLOCK_LIST_TWO','BLOCK_CELL_ONE','BLOCK_CELL_TWO','BLOCK_OMEGA'):
        assert document_text(restored).count(word)==1
    assert restored['content'][1]['attrs']['order']==3
    assert restored['content'][-1]['content'][0]['type']=='inline_image'
    assert restored['content'][-1]['content'][0]['attrs']['alt']=='Form picture'
    assert any(m['type'] in ('text_style','textStyle') and m['attrs'].get('fontSize')==14 for m in restored['content'][-1]['content'][-1]['marks'])
    assert project==before


@pytest.mark.parametrize('value,kind',[('2026-10-30','date'),('30 October 2026','plain')])
def test_word_date_without_calendar_metadata_never_loses_displayed_wording(tmp_path,value,kind):
    result=build_export(fixture(),'docx',tmp_path);word=Document(result['path'])
    control=list(word.element.iter(qn('w:sdt')))[3]
    control.find(qn('w:sdtPr')+'/'+qn('w:date')).attrib.pop(qn('w:fullDate'))
    next(control.iter(qn('w:t'))).text=value
    word.save(result['path'])
    restored=import_document(Path(result['path']));field=fields(restored['document'])[3]
    assert field['attrs']['kind']==kind
    assert document_text(field)==value
    if kind=='plain':assert any('without a calendar value' in w for w in restored['warnings'])
