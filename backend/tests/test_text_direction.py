from pathlib import Path
from zipfile import ZipFile
import pytest
from alder.models import create_project,text_document,validate_project,ValidationError,document_text
from alder.publishing import build_export,import_document

def fixture():
    p=create_project('Multilingual paragraphs',template='blank')
    p['settings'].update(documentKind='docx',includeTitle=False,footer=True,fontSize=14,lineHeight=1.4)
    samples=[('ltr','en','ENGLISH Latin and العربية 42.'),('rtl','ar','مرحبا بالعالم English 123 (2026).'),('rtl','he','שלום עולם English 456 (2026).'),('ltr','ja','日本語の文章と English 123。'),('ltr','th','ภาษาไทยและ English 789'),('ltr','en','Emoji 😀 and combining e\u0301 remain intact.'),('rtl','ar','بداية\tنهاية')]
    content=[]
    for direction,language,text in samples:
        para=text_document(text)['content'][0];para['attrs']={'textDirection':direction,'textLanguage':language,'spaceAfter':10}
        if '\t' in text:para['attrs']['tabStops']=[{'position':144,'align':'right','leader':'dot'}]
        content.append(para)
    p['book']={'version':1,'chapters':[{'id':'main','title':'Document','include':True,'document':{'type':'doc','content':content}}]}
    return validate_project(p)

@pytest.mark.parametrize('format',['html','epub','docx'])
def test_direction_and_language_survive_exchange_without_reordering_text(tmp_path,format):
    p=fixture();result=build_export(p,format,tmp_path);restored=import_document(Path(result['path']))
    expected=p['book']['chapters'][0]['document']
    assert document_text(restored['document'])==document_text(expected)
    assert [(n['attrs']['textDirection'],n['attrs']['textLanguage']) for n in restored['document']['content']]==[(n['attrs']['textDirection'],n['attrs']['textLanguage']) for n in expected['content']]
    if format=='docx':
        from lxml import etree
        with ZipFile(result['path']) as z:root=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert root.xpath('//w:body/w:p/w:pPr/w:bidi/@w:val',namespaces=ns)==['0','1','1','0','0','0','1']

def test_direction_validation_and_inherited_html(tmp_path):
    p=fixture();p['book']['chapters'][0]['document']['content'][0]['attrs']['textDirection']='sideways'
    with pytest.raises(ValidationError):validate_project(p)
    source=tmp_path/'inherited.html';source.write_text('<html lang="ar" dir="rtl"><body><p>مرحبا English</p><p dir="ltr" lang="en">English</p></body></html>',encoding='utf-8')
    restored=import_document(source)
    assert [(n['attrs']['textDirection'],n['attrs']['textLanguage']) for n in restored['document']['content']]==[('rtl','ar'),('ltr','en')]

def test_named_direction_styles_resolve_before_html_and_word_export(tmp_path):
    p=fixture();p['styles'].append({'id':'rtlstyle','name':'Arabic','kind':'paragraph','textDirection':'rtl','textLanguage':'ar'})
    p['book']['chapters'][0]['document']['content'][0]['attrs']={'styleId':'rtlstyle'}
    for format in ('html','docx'):
        restored=import_document(Path(build_export(p,format,tmp_path)['path']))
        assert restored['document']['content'][0]['attrs']['textDirection']=='rtl'

def test_unicode_resolution_includes_brackets_numbers_and_supplementary_characters():
    import unicodedata
    from alder.text_direction import embedding_levels,word_directions
    text='مرحبا بالعالم English 123 (2026).'
    assert embedding_levels(text,True)==[1]*14+[2]*18+[1]
    text='ENGLISH Latin and العربية 42.'
    assert embedding_levels(text)==[0]*18+[1]*8+[2]*2+[0]
    assert word_directions(text)==[0]*18+[1]*10+[0]
    for text in ('مرحبا بالعالم English 123 (2026).','ENGLISH Latin and العربية 42.','שלום עולם English 456','😀 e\u0301 العربية','(123) 日本語 שלום','English \u2067שלום 123\u2069 end'):
        for rtl in (True,False):
            directions=word_directions(text,rtl)
            assert len(directions)==len(text)
            for char,direction in zip(text,directions):
                if unicodedata.bidirectional(char)=='L':assert direction==0
                elif unicodedata.bidirectional(char) in ('R','AL'):assert direction==1


def test_word_direction_uses_context_across_formatting_and_link_runs():
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from alder.text_direction import resolve_word_directions,complex_script_properties
    d=Document();p=d.add_paragraph();p.add_run('English العربية ').bold=True
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'example');run=p.add_run('42');link.append(run._r);p._p.append(link);p.add_run('.')
    resolve_word_directions(d);complex_script_properties(d)
    assert ''.join(n.text for n in p._p.iter(qn('w:t')))=='English العربية 42.'
    assert link.find('./'+qn('w:r')+'/'+qn('w:rPr')+'/'+qn('w:rtl')).get(qn('w:val'))=='1'
    strong=next(r for r in p._p.iter(qn('w:r')) if 'العربية' in ''.join(t.text for t in r.iter(qn('w:t'))))
    assert strong.find('./'+qn('w:rPr')+'/'+qn('w:bCs')) is not None
    before=d.element.xml;resolve_word_directions(d);complex_script_properties(d);assert d.element.xml==before


def test_missing_bidi_runtime_fails_only_when_needed(monkeypatch,tmp_path):
    import alder.text_direction as direction
    direction._icu.cache_clear();monkeypatch.setattr(direction,'_library_candidates',lambda:iter(()))
    try:
        p=fixture();p['book']['chapters'][0]['document']=text_document('English only')
        assert Path(build_export(p,'docx',tmp_path)['path']).is_file()
        with pytest.raises(ValueError,match='ICU'):build_export(fixture(),'docx',tmp_path)
    finally:direction._icu.cache_clear()
