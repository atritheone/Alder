import io

import pytest
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

from alder.models import document_text, validate_document
from alder.publishing import import_document


def legacy_document(tmp_path, *, external=False, alternate=False, positioned=False):
    doc=Document();p=doc.add_paragraph('Before ');stream=io.BytesIO()
    Image.new('RGB',(120,64),'blue').save(stream,format='PNG');stream.seek(0)
    shape=p.add_run().add_picture(stream)
    drawing=shape._inline.getparent()
    rid=next(drawing.iter(qn('a:blip'))).get(qn('r:embed'))
    if external:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        rid=doc.part.relate_to('https://example.com/unfetched.png',RT.IMAGE,is_external=True)
    props='position:absolute;rotation:90;' if positioned else ''
    pict=parse_xml(f'<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><v:shape alt="Blue rectangle" style="{props}width:90pt;height:48pt"><v:imagedata r:id="{rid}" cropleft="6553.6f"/></v:shape></w:pict>')
    if alternate:
        wrapper=parse_xml('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice Requires="wps"/><mc:Fallback/></mc:AlternateContent>')
        parent=drawing.getparent();parent.replace(drawing,wrapper)
        wrapper[0].append(drawing);wrapper[1].append(pict)
    else:drawing.getparent().replace(drawing,pict)
    p.add_run(' after.')
    file=tmp_path/'legacy.docx';doc.save(file);return file


@pytest.mark.parametrize('positioned',[False,True])
def test_legacy_embedded_raster_import_keeps_wording_asset_crop_width_and_alt(tmp_path,positioned):
    result=import_document(legacy_document(tmp_path,positioned=positioned))
    validate_document(result['document']);assert len(result['assets'])==1
    assert document_text(result['document'])=='Before  after.'
    picture=next(n for n in result['document']['content'][0]['content'] if n['type']=='inline_image')
    assert picture['attrs']['width']==120 and picture['attrs']['alt']=='Blue rectangle'
    assert picture['attrs']['crop']['left']==pytest.approx(10)
    assert any('positioned legacy' in w for w in result['warnings'])==positioned


def test_drawingml_and_vml_alternates_import_only_one_picture(tmp_path):
    result=import_document(legacy_document(tmp_path,alternate=True))
    assert len(result['assets'])==1
    assert len([n for n in result['document']['content'][0]['content'] if n['type']=='inline_image'])==1


def test_external_vml_images_fail_without_fetching(tmp_path):
    with pytest.raises(ValueError,match='linked externally'):
        import_document(legacy_document(tmp_path,external=True))
