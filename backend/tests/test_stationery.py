import copy
from pathlib import Path
import pytest
from docx import Document
from alder.stationery import PRESETS, create_stationery, validate_stock
from alder.models import validate_project, ValidationError
from alder.publication_layout import prepare_layout
from alder.publishing import build_export, import_document


@pytest.mark.parametrize('identity', [p['id'] for p in PRESETS])
def test_native_stationery_geometry_fields_and_independent_copies(tmp_path, identity):
    project=create_stationery(identity,'Addresses')
    before=copy.deepcopy(project)
    validate_stock(project)
    layout=prepare_layout(project)
    assert layout['settings']['stationeryPreset']==identity
    assert layout['document']['content'][0]['type']=='table'
    path=build_export(project,'docx',tmp_path,{'measuredRowSplits':[]})['path']
    word=Document(path)
    stock=next(s for s in PRESETS if s['id']==identity)
    assert abs(word.sections[0].page_width.mm-stock['paper'][0])<.03
    assert abs(word.sections[0].page_height.mm-stock['paper'][1])<.03
    assert len(word.tables)==1
    assert len(word.tables[0].rows)==len(stock['rows'])
    restored=import_document(Path(path))
    table=next(n for n in restored['document']['content'] if n['type']=='table')
    for row,height in zip(table['content'],stock['rows']):
        assert abs(row['attrs']['minHeight']-height*72/25.4)<.15
        for cell,width in zip(row['content'],stock['columns']):
            assert abs(cell['attrs']['colwidth'][0]-width*96/25.4)<.1
    assert project==before
    other=create_stationery(identity,'Independent')
    assert other['id']!=project['id']
    assert other['book']['chapters'][0]['id']!=project['book']['chapters'][0]['id']


@pytest.mark.parametrize('change',['width','margin','row','column','extra','gap','footer','columns'])
def test_changed_stock_refuses_print_without_blocking_save(change):
    p=create_stationery('labels-a4-21','Labels')
    table=p['book']['chapters'][0]['document']['content'][0]
    if change=='width':p['settings']['pageWidthMm']=220
    elif change=='margin':p['settings']['marginLeftMm']=10
    elif change=='row':table['content'].pop()
    elif change=='column':table['content'][0]['content'][0]['attrs']['colwidth']=[200]
    elif change=='extra':p['book']['chapters'][0]['document']['content'].append({'type':'paragraph','content':[{'type':'text','text':'Extra'}]})
    elif change=='gap':table['content'][0]['content'][1]['content'][0]['content']=[{'type':'text','text':'In gap'}]
    elif change=='footer':p['settings']['footer']=True
    else:p['settings']['columns']={'count':2,'gapMm':10,'separator':False}
    validate_project(p)
    with pytest.raises(ValueError,match='stationery layout'):
        prepare_layout(p)
    p['settings']['stationeryPreset']=None
    prepare_layout(p)


def test_invalid_stock_and_headless_print_require_explicit_resolution(tmp_path):
    with pytest.raises(ValueError,match='Unknown'):
        create_stationery('../fake','Name')
    with pytest.raises(ValidationError):
        create_stationery('envelope-dl','')
    p=create_stationery('envelope-dl','Envelope')
    for kind in ('docx','pdf'):
        with pytest.raises(ValueError,match='desktop'):
            build_export(p,kind,tmp_path)
    p['settings']['stationeryPreset']='missing'
    with pytest.raises(ValidationError,match='Unknown'):
        validate_project(p)


def test_saved_template_keeps_stock_only_with_its_layout(tmp_path):
    from alder.store import Store
    from alder.templates import Templates
    store=Store(tmp_path)
    original=store.insert(create_stationery('labels-a4-21','Labels'))
    library=Templates(store)
    for content in (True,False):
        saved=library.save(original['id'],str(content),content)
        new=library.create(saved['id'],'Copy')
        assert bool(new['settings'].get('stationeryPreset'))==content
        prepare_layout(new)


def test_fractional_label_gaps_survive_html_exchange(tmp_path):
    project=create_stationery('labels-a4-21','Labels')
    result=build_export(project,'html',tmp_path)
    restored=import_document(Path(result['path']))
    table=next(n for n in restored['document']['content'] if n['type']=='table')
    assert abs(table['content'][0]['content'][1]['attrs']['colwidth'][0]-2.5*96/25.4)<.0001
