import json
from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfReader
from alder.models import create_project, validate_project, ValidationError
from alder.page_geometry import page_geometry
from alder.publishing import build_export, import_document

CASES = json.loads((Path(__file__).parent / 'fixtures/page-geometry.json').read_text('utf-8'))


@pytest.mark.parametrize('fixture', CASES, ids=lambda case: case['name'])
def test_shared_page_geometry_contract(fixture):
    actual = page_geometry(fixture['settings'])
    for key, value in fixture['expected'].items():
        assert actual[key] == value


def project(settings):
    result = create_project(template='blank')
    result['settings'].update(documentKind='docx', includeTitle=False, footer=False, **settings)
    result['book'] = {'version': 1, 'chapters': [{'id':'geometry', 'title':'Document', 'document': {'type':'doc', 'content':[{'type':'paragraph', 'content':[{'type':'text', 'text':'GEOMETRY_MARKER'}]}]}, 'text':'', 'include':True}]}
    return validate_project(result)


@pytest.mark.parametrize('settings', [CASES[1]['settings'], CASES[2]['settings'], CASES[3]['settings']])
def test_page_geometry_matches_docx_pdf_html_and_docx_import(tmp_path, settings):
    p = project(settings)
    geometry = page_geometry(settings)
    word = build_export(p, 'docx', tmp_path)['path']
    section = Document(word).sections[0]
    assert section.page_width.mm == pytest.approx(geometry['width'], abs=.03)
    assert section.page_height.mm == pytest.approx(geometry['height'], abs=.03)
    for edge, value in geometry['margins'].items():
        actual = getattr(section, edge+'_margin').mm + (section.gutter.mm if edge == 'left' else 0)
        assert actual == pytest.approx(value, abs=.04)
    restored = page_geometry(import_document(Path(word))['settings'])
    for edge, value in geometry['margins'].items():
        assert restored['margins'][edge] == pytest.approx(value, abs=.04)
    pdf = PdfReader(build_export(p, 'pdf', tmp_path)['path'])
    assert float(pdf.pages[0].mediabox.width) == pytest.approx(geometry['width']*72/25.4, abs=.01)
    assert float(pdf.pages[0].mediabox.height) == pytest.approx(geometry['height']*72/25.4, abs=.01)
    positions = []
    pdf.pages[0].extract_text(visitor_text=lambda text, cm, tm, font, size: positions.append(cm[4]+tm[4]) if 'GEOMETRY_MARKER' in text else None)
    assert positions == pytest.approx([geometry['margins']['left']*72/25.4], abs=.02)
    html = Path(build_export(p, 'html', tmp_path)['path']).read_text('utf-8')
    assert f"size: {geometry['width']:g}mm {geometry['height']:g}mm;" in html
    assert 'margin: ' + ' '.join(f"{geometry['margins'][edge]:g}mm" for edge in ('top','right','bottom','left')) in html


@pytest.mark.parametrize('settings', [
    {'pageWidthMm':49}, {'pageHeightMm':1001}, {'marginTopMm':-1},
    {'gutterMm':True}, {'marginRightMm':float('nan')}, {'gutterMm':float('inf')},
    {'pageSize':'A5', 'marginLeftMm':80, 'marginRightMm':80},
])
def test_invalid_new_page_settings_are_rejected(settings):
    with pytest.raises(ValidationError): project(settings)


def test_legacy_impossible_margins_remain_readable_and_report_conversion(tmp_path):
    p = project(CASES[-1]['settings'])
    result = build_export(p, 'docx', tmp_path)
    assert any('reduced proportionally' in warning for warning in result['warnings'])
    section = Document(result['path']).sections[0]
    assert section.left_margin.mm == pytest.approx(64, abs=.03)


def test_mixed_docx_sections_retain_each_section_geometry(tmp_path):
    from docx.shared import Mm
    doc = Document()
    doc.sections[0].page_width = Mm(180)
    doc.sections[0].page_height = Mm(240)
    doc.add_paragraph('One')
    section = doc.add_section()
    section.page_width = Mm(240)
    section.page_height = Mm(180)
    doc.add_paragraph('Two')
    path = tmp_path/'sections.docx'
    doc.save(path)
    imported = import_document(path)
    assert imported['settings']['pageWidthMm'] == pytest.approx(180, abs=.03)
    section = next(n for n in imported['document']['content'] if n['type'] == 'section_break')
    assert section['attrs']['page']['pageWidthMm'] == pytest.approx(240, abs=.03)
    assert section['attrs']['page']['pageHeightMm'] == pytest.approx(180, abs=.03)
    assert not any('first section' in warning for warning in imported['warnings'])
