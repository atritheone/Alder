from copy import deepcopy
from pathlib import Path
from alder.heading_numbering import number_headings
from alder.publishing import build_export, import_document


def heading_project():
    from test_book import book_project
    p = book_project()
    p['settings'].update(documentKind='docx', includeTitle=False, includeToc=False)
    def heading(level, text, **attrs):
        return {'type': 'heading', 'attrs': {'level': level, 'headingNumbered': True, **attrs}, 'content': [{'type': 'text', 'text': text}]}
    p['book']['chapters'] = p['book']['chapters'][:1]
    p['book']['chapters'][0]['document'] = {'type': 'doc', 'content': [
        heading(1, 'First heading'), heading(2, 'First child'), heading(2, 'Second child'),
        heading(1, 'Restarted heading', headingStart=5), heading(2, 'Next child'),
        heading(1, 'Following heading'), heading(3, 'Skipped level'),
    ]}
    return p


def test_native_heading_numbers_remain_editable_after_docx_exchange(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    p = heading_project(); before = deepcopy(p)
    result = build_export(p, 'docx', tmp_path)
    word = Document(result['path'])
    assert all(paragraph._p.pPr.numPr is not None for paragraph in word.paragraphs)
    labels = [p._p.pPr.numPr.ilvl.val for p in word.paragraphs]
    assert labels == [0,1,1,0,1,0,2]
    assert word.paragraphs[0].text == 'First heading'
    imported = import_document(Path(result['path']))
    nodes = imported['document']['content']
    assert all(node['type'] == 'heading' for node in nodes)
    number_headings([{'blocks': nodes}])
    assert [n['attrs']['_headingLabel'] for n in nodes] == ['1','1.1','1.2','5','5.1','6','6.1.1']
    assert [n['attrs'].get('headingStart') for n in nodes] == [None,None,None,5,None,None,None]
    assert p == before


def test_html_numbers_never_duplicate_authored_heading_wording(tmp_path):
    p = heading_project()
    result = build_export(p, 'html', tmp_path)
    assert 'data-heading-marker="true">5.1' in Path(result['path']).read_text('utf-8')
    imported = import_document(Path(result['path']))
    nodes = [n for n in imported['document']['content'] if n['type'] == 'heading']
    assert nodes[0]['content'] == [{'type': 'text', 'text': 'First heading'}]
    assert nodes[3]['attrs']['headingStart'] == 5
    assert all(n['attrs']['headingNumbered'] for n in nodes)
