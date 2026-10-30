import base64
import copy
import io
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn

from alder.image_wrapping import image_wrapping, read_word_wrapping
from alder.models import create_project, validate_project, ValidationError
from alder.publishing import build_export, import_document


def fixture(side='left'):
    data = io.BytesIO()
    Image.new('RGB', (120, 180), '#263f5d').save(data, 'PNG')
    project = create_project(template='blank')
    project['settings'].update(documentKind='docx', includeTitle=False)
    project['assets'] = [{'id':'picture', 'name':'Picture.png', 'mime':'image/png', 'data':base64.b64encode(data.getvalue()).decode()}]
    picture = {'type':'inline_image', 'attrs':{'assetId':'picture', 'src':'asset:picture', 'width':96, 'alt':'A blue rectangle', 'wrap':side, 'wrapGap':16, 'crop':{'left':10,'top':0,'right':0,'bottom':0}}}
    project['book'] = {'version':1, 'chapters':[{'id':'main', 'title':'Text', 'include':True, 'document':{'type':'doc','content':[{'type':'paragraph','content':[picture,{'type':'text','text':'The picture follows this paragraph.'}]}, {'type':'paragraph','content':[{'type':'text','text':'The following paragraph also flows around it.'}]}]}}]}
    return validate_project(project)


@pytest.mark.parametrize('side', ['left', 'right'])
@pytest.mark.parametrize('format', ['docx', 'html', 'epub'])
def test_wrapped_picture_exchange_keeps_anchor_spacing_crop_and_description(tmp_path, side, format):
    project = fixture(side)
    original = copy.deepcopy(project)
    exported = build_export(project, format, tmp_path, {'validate':False})
    assert original == project
    restored = import_document(Path(exported['path']))
    picture = restored['document']['content'][0]['content'][0]
    assert picture['type'] == 'inline_image'
    assert picture['attrs']['wrap'] == side
    assert picture['attrs']['wrapGap'] == pytest.approx(16)
    assert picture['attrs']['crop']['left'] == 10
    assert picture['attrs']['alt'] == 'A blue rectangle'
    assert picture['attrs']['width'] == pytest.approx(96, abs=.001)
    if format == 'docx':
        document = Document(exported['path'])
        anchor = next(document.element.iter(qn('wp:anchor')))
        assert len(document.inline_shapes) == 0
        assert anchor.find(qn('wp:positionH')).get('relativeFrom') == 'column'
        assert anchor.find(qn('wp:positionV')).get('relativeFrom') == 'paragraph'
        assert anchor.find(qn('wp:wrapSquare')) is not None
        assert not any('floating Word picture' in warning for warning in restored['warnings'])


def test_unsupported_anchor_positions_are_reported_instead_of_misrepresented(tmp_path):
    exported = build_export(fixture(), 'docx', tmp_path, {'validate':False})
    document = Document(exported['path'])
    anchor = next(document.element.iter(qn('wp:anchor')))
    anchor.find(qn('wp:positionV')).find(qn('wp:posOffset')).text = '9525'
    assert read_word_wrapping(anchor) is None
    document.save(exported['path'])
    restored = import_document(Path(exported['path']))
    assert any('floating Word picture' in warning for warning in restored['warnings'])
    assert restored['document']['content'][0]['content'][0]['attrs'].get('wrap', 'none') == 'none'


def test_wrap_validation_and_direct_pdf_guard(tmp_path):
    project = fixture()
    for attrs in ({'wrap':'through'}, {'wrapGap':float('nan')}, {'wrapGap':True}, {'wrapGap':193}, {'wrapGap':-1}):
        with pytest.raises(ValueError):
            image_wrapping(attrs)
    with pytest.raises(ValueError, match='desktop PDF'):
        build_export(project, 'pdf', tmp_path)
    picture = project['book']['chapters'][0]['document']['content'][0]['content'][0]
    picture['type'] = 'image'
    with pytest.raises(ValidationError, match='anchored inside a paragraph'):
        validate_project(project)


def test_measured_oversized_rows_split_in_word_without_changing_the_source(tmp_path):
    project=fixture()
    blocks=project['book']['chapters'][0]['document']['content']
    row={'type':'table_row','attrs':{'blockId':'tall-row','allowSplit':False},'content':[{'type':'table_cell','content':blocks}]}
    project['book']['chapters'][0]['document']['content']=[{'type':'table','content':[row]}]
    project=validate_project(project);original=copy.deepcopy(project)
    exported=build_export(project,'docx',tmp_path,{'measuredRowSplits':['tall-row']})
    assert not list(Document(exported['path']).element.iter(qn('w:cantSplit')))
    assert any('taller than a page' in warning for warning in exported['warnings'])
    assert project==original
    kept=build_export(project,'docx',tmp_path,{'measuredRowSplits':[]})
    assert len(list(Document(kept['path']).element.iter(qn('w:cantSplit'))))==1
    with pytest.raises(ValueError,match='measured table-row'):
        build_export(project,'docx',tmp_path,{'measuredRowSplits':'tall-row'})
