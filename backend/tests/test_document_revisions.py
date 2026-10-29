from copy import deepcopy
from lxml import etree
import pytest
from alder.document_revisions import review_document, revision_summary, validate_revision, WordRevisionWriter, WordRevisionReader, W


def fixture():
    stamp = {'id': 'change_one', 'author': 'Reviewer', 'date': '2026-10-29T01:00:00Z'}
    return {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [
        {'type': 'text', 'text': 'Old', 'marks': [{'type': 'strong'}, {'type': 'revision_deletion', 'attrs': stamp}]},
        {'type': 'text', 'text': 'New', 'marks': [{'type': 'revision_insertion', 'attrs': stamp}]},
        {'type': 'text', 'text': ' wording.'},
    ]}]}


def text(node):
    return node.get('text', '') + ''.join(text(child) for child in node.get('content', []))


def test_original_final_and_markup_are_independent_non_destructive_views():
    doc = fixture(); before = deepcopy(doc)
    assert text(review_document(doc, 'final')) == 'New wording.'
    assert text(review_document(doc, 'original')) == 'Old wording.'
    assert text(review_document(doc, 'markup')) == 'OldNew wording.'
    assert review_document(doc, 'original')['content'][0]['content'][0]['marks'] == [{'type': 'strong'}]
    assert doc == before
    assert revision_summary(doc) == [{'id': 'change_one', 'author': 'Reviewer', 'date': '2026-10-29T01:00:00Z', 'inserted': 'New', 'deleted': 'Old'}]
    from alder.models import document_text
    from alder.document_references import reference_values
    assert document_text(doc) == 'New wording.'
    doc['content'][0].update(type='heading', attrs={'blockId': 'heading'})
    assert reference_values([doc])['heading'] == 'New wording.'


def test_word_runs_retain_deleted_wording_author_and_marks():
    from docx import Document
    doc = Document(); p = doc.add_paragraph(); writer = WordRevisionWriter()
    def render(paragraph, node):
        run = paragraph.add_run(node['text'])
        if any(mark['type'] == 'strong' for mark in node.get('marks', [])): run.bold = True
    for node in fixture()['content'][0]['content']:
        if not writer.inline(p, node, render): render(p, node)
    reader = WordRevisionReader(doc._element)
    deletion = p._p.find(f'{{{W}}}del'); insertion = p._p.find(f'{{{W}}}ins')
    assert deletion.find(f'.//{{{W}}}delText').text == 'Old'
    assert deletion.find(f'.//{{{W}}}b') is not None
    assert insertion.find(f'.//{{{W}}}t').text == 'New'
    a, b = reader.mark(deletion), reader.mark(insertion)
    assert a['attrs'] == b['attrs'] and a['attrs']['author'] == 'Reviewer'
    assert a['type'] == 'revision_deletion' and b['type'] == 'revision_insertion'


def test_nested_insertions_and_deletions_keep_separate_authorship():
    from docx import Document
    doc = Document(); p = doc.add_paragraph(); writer = WordRevisionWriter()
    node = fixture()['content'][0]['content'][1]
    node['marks'].append({'type': 'revision_deletion', 'attrs': {'id': 'change_two', 'author': 'Second reviewer', 'date': ''}})
    writer.inline(p, node, lambda paragraph, item: paragraph.add_run(item['text']))
    insertion = p._p.find(f'{{{W}}}ins'); deletion = insertion.find(f'{{{W}}}del')
    assert deletion.find(f'.//{{{W}}}delText').text == 'New'
    reader = WordRevisionReader(doc._element)
    assert reader.mark(insertion)['attrs']['author'] == 'Reviewer'
    assert reader.mark(deletion)['attrs']['author'] == 'Second reviewer'


def test_unsupported_structural_changes_are_never_silently_imported():
    for content in ('<w:moveFrom/>', '<w:rPrChange/>', '<w:rPr><w:del/></w:rPr>', '<w:trPr><w:ins/></w:trPr>'):
        root = etree.fromstring(f'<w:document xmlns:w="{W}">{content}</w:document>')
        with pytest.raises(ValueError, match='original file has not been changed'): WordRevisionReader(root)
    with pytest.raises(ValueError): validate_revision({'id': 'ok', 'author': 'Writer', 'date': 'not a date'})


def test_word_paragraph_breaks_keep_final_structure_and_restorable_formatting():
    from docx import Document
    from alder.document_revisions import word_review_paragraphs, restore_review_paragraphs
    stamp = {'id': 'change_break', 'author': 'Writer', 'date': '2026-10-29T01:00:00Z'}
    source = {'type': 'doc', 'content': [{'type': 'paragraph', 'attrs': {'blockId': 'left'}, 'content': [
        {'type': 'text', 'text': 'One'},
        {'type': 'revision_break', 'attrs': {'kind': 'deletion', 'stamp': stamp, 'rightType': 'heading', 'rightAttrs': {'blockId': 'right', 'level': 2, 'align': 'center'}}},
        {'type': 'text', 'text': 'Two'},
    ]}]}
    with pytest.raises(ValueError, match='differently formatted'): word_review_paragraphs(source)
    source['content'][0]['content'][1]['attrs'].update(rightType='paragraph', rightAttrs={'blockId': 'right'})
    native = word_review_paragraphs(source)
    assert [text(node) for node in native['content']] == ['One', 'Two']
    doc = Document(); writer = WordRevisionWriter()
    for node in native['content']:
        p = doc.add_paragraph(text(node))
        if node['attrs'].get('revisionBreak'): writer.paragraph_break(p, node['attrs']['revisionBreak'])
    reader = WordRevisionReader(doc._element)
    change = reader.paragraph_break(doc.paragraphs[0]._p)
    assert change['kind'] == 'deletion' and change['stamp']['author'] == 'Writer'
    restored = restore_review_paragraphs(native)
    assert restored == source


def test_inserted_native_paragraph_breaks_pair_with_the_next_paragraph():
    from alder.document_revisions import restore_review_paragraphs, word_review_paragraphs
    stamp = {'id': 'change_break', 'author': 'Writer', 'date': ''}
    native = {'type': 'doc', 'content': [
        {'type': 'paragraph', 'attrs': {'revisionBreak': {'kind': 'insertion', 'stamp': stamp}}, 'content': [{'type': 'text', 'text': 'One'}]},
        {'type': 'paragraph', 'attrs': {'blockId': 'right'}, 'content': [{'type': 'text', 'text': 'Two'}]},
    ]}
    expected = deepcopy(native)
    restored = restore_review_paragraphs(native)
    marker = restored['content'][0]['content'][-1]
    assert marker['attrs']['nextId'] == 'right'
    assert word_review_paragraphs(restored) == expected
    broken = {'type': 'doc', 'content': expected['content'][:1]}
    with pytest.raises(ValueError, match='no following paragraph'): restore_review_paragraphs(broken)


def project_fixture():
    from alder.models import create_project, validate_project
    from test_revision_formatting import fixture as formats
    project = create_project(template='blank')
    project['settings'].update(documentKind='docx', includeTitle=False, footer=False)
    doc = fixture()
    stamp = {'id': 'paragraph_change', 'author': 'Structure reviewer', 'date': '2026-10-29T03:00:00Z'}
    doc['content'] += [
        {'type': 'paragraph', 'attrs': {'blockId': 'joined'}, 'content': [{'type': 'text', 'text': 'Left'}, {'type': 'revision_break', 'attrs': {'kind': 'deletion', 'stamp': stamp, 'rightType': 'paragraph', 'rightAttrs': {'blockId': 'right'}}}, {'type': 'text', 'text': 'Right'}]},
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Split'}, {'type': 'revision_break', 'attrs': {'kind': 'insertion', 'stamp': stamp, 'nextId': 'next'}}]},
        {'type': 'paragraph', 'attrs': {'blockId': 'next'}, 'content': [{'type': 'text', 'text': ' boundary'}]},
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Formatting', 'marks': [{'type': 'em'}, {'type': 'strong'}, {'type': 'revision_format', 'attrs': formats()}]}]},
    ]
    project['book'] = {'version': 1, 'chapters': [{'id': 'chapter', 'title': 'Document', 'include': True, 'document': doc}]}
    return validate_project(project)


@pytest.mark.parametrize('format', ['html', 'epub', 'docx'])
def test_production_exchange_retains_text_paragraph_and_format_revisions(tmp_path, format):
    from pathlib import Path
    from alder.publishing import build_export, import_document
    from alder.models import validate_project, document_text
    from alder.document_comments import walk
    project = project_fixture(); before = deepcopy(project)
    result = build_export(project, format, tmp_path)
    assert project == before
    imported = import_document(Path(result['path']))['document']
    project['book']['chapters'][0]['document'] = imported
    imported = validate_project(project)['book']['chapters'][0]['document']
    assert document_text(review_document(imported, 'final')) == 'New wording.\nLeftRight\nSplit\n boundary\nFormatting'
    assert document_text(review_document(imported, 'original')) == 'Old wording.\nLeft\nRight\nSplit boundary\nFormatting'
    formats = [m for n in walk(imported) for m in n.get('marks', []) if m['type'] == 'revision_format']
    assert len(formats) == 1
    assert formats[0]['attrs']['records'][0]['stamp']['author'] == 'Format reviewer'
    assert {c['author'] for c in revision_summary(imported)} == {'Reviewer'}
    assert [n['attrs']['kind'] for n in walk(imported) if n['type'] == 'revision_break'] == ['deletion', 'insertion']


@pytest.mark.parametrize('format', ['txt', 'md', 'pdf'])
def test_plain_publications_use_final_text_without_accepting_authored_changes(tmp_path, format):
    from pathlib import Path
    from alder.publishing import build_export
    project = project_fixture(); before = deepcopy(project)
    result = build_export(project, format, tmp_path)
    if format == 'pdf':
        from pypdf import PdfReader
        value = '\n'.join(page.extract_text() for page in PdfReader(result['path']).pages)
    else: value = Path(result['path']).read_text(encoding='utf-8')
    assert 'New wording.' in value and 'Old' not in value and 'LeftRight' in value
    assert any('final review view' in warning for warning in result['warnings'])
    assert project == before


def test_production_word_import_refuses_unsupported_revisions_before_discarding_them(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from alder.publishing import import_document
    doc = Document(); p = doc.add_paragraph('Kept')
    p._p.append(OxmlElement('w:moveFrom'))
    path = tmp_path / 'unsupported.docx'; doc.save(path); before = path.read_bytes()
    with pytest.raises(ValueError, match='original file has not been changed'): import_document(path)
    assert path.read_bytes() == before
    doc = Document(); p = doc.add_paragraph('Left')
    WordRevisionWriter().paragraph_break(p, {'kind':'deletion','stamp':{'id':'break','author':'Writer','date':''}})
    doc.add_paragraph('Right','Heading 2'); doc.save(path); before = path.read_bytes()
    with pytest.raises(ValueError, match='different formatting'): import_document(path)
    assert path.read_bytes() == before


def test_duplicate_chapters_get_independent_revision_groups_and_revalidate_stably():
    from alder.models import validate_project
    from alder.document_comments import walk
    project = project_fixture(); chapter = deepcopy(project['book']['chapters'][0]); chapter['id'] = 'copy'; project['book']['chapters'].append(chapter)
    result = validate_project(project)
    def stamps(doc):
        found = set()
        for node in walk(doc):
            if node['type'] == 'revision_break': found.add(node['attrs']['stamp']['id'])
            for mark in node.get('marks', []):
                if mark['type'] in ('revision_insertion', 'revision_deletion'): found.add(mark['attrs']['id'])
                elif mark['type'] == 'revision_format': found.update(r['stamp']['id'] for r in mark['attrs']['records'])
        return found
    left, right = [stamps(chapter['document']) for chapter in result['book']['chapters']]
    assert len(left) == len(right) == 3 and not left & right
    assert validate_project(result)['book'] == result['book']
