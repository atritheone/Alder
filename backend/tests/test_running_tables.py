from pathlib import Path
from copy import deepcopy
import pytest
from docx import Document
from pypdf import PdfReader
from alder.models import create_project,validate_project,document_text
from alder.publishing import build_export,import_document


def fixture():
    p=create_project(template='blank')
    paragraph=lambda text:{'type':'paragraph','content':[{'type':'text','text':text}]}
    p['settings'].update(documentKind='docx',includeTitle=False,pageSize='A5',marginTopMm=42,marginMm=15,footer=False)
    table={'type':'table','attrs':{'align':'center'},'content':[
        {'type':'table_row','content':[
            {'type':'table_cell','attrs':{'rowspan':2,'colwidth':[200],'background':'#e2e9f3'},'content':[paragraph('HEADER COMPANY')]},
            {'type':'table_cell','attrs':{'colwidth':[246],'verticalAlign':'middle'},'content':[paragraph('HEADER DETAILS')]},
        ]},
        {'type':'table_row','content':[
            {'type':'table_cell','attrs':{'colwidth':[246]},'content':[{'type':'paragraph','content':[{'type':'text','text':'Page '},{'type':'document_field','attrs':{'fieldKind':'page'},'content':[{'type':'text','text':'1'}]}]}]},
        ]},
    ]}
    p['settings']['runningMatter']={'headerDistanceMm':7,'footerDistanceMm':8,'header':{'default':{'type':'doc','content':[paragraph('BEFORE TABLE'),table,paragraph('AFTER TABLE')]}},'footer':{'default':None}}
    p['book']={'version':1,'chapters':[{'id':'body','title':'Document','include':True,'document':{'type':'doc','content':[paragraph('BODY ONE'),{'type':'page_break'},paragraph('BODY TWO')]}}]}
    return validate_project(p)


@pytest.mark.parametrize('format',['docx','html','epub'])
def test_running_tables_preserve_structure_order_and_fields(tmp_path,format):
    p=fixture();before=deepcopy(p)
    result=build_export(p,format,tmp_path)
    assert p==before
    restored=import_document(Path(result['path']))
    story=restored['settings']['runningMatter']['header']['default']
    kinds=[n['type'] for n in story['content']]
    assert kinds[:3]==['paragraph','table','paragraph']
    table=story['content'][1]
    assert table['attrs']['align']=='center'
    assert table['content'][0]['content'][0]['attrs']['rowspan']==2
    assert table['content'][0]['content'][0]['attrs']['background']=='#e2e9f3'
    assert any(n['type']=='document_field' for n in table['content'][1]['content'][0]['content'][0]['content'])
    assert not any('Complex DOCX header' in warning for warning in restored['warnings'])
    assert 'HEADER' not in document_text(restored['document'])
    if format=='docx':
        word=Document(result['path']);header=word.sections[0].header
        assert len(header.tables)==1
        assert header.tables[0].cell(0,0).paragraphs[0].runs[0].font.size.pt==9
        assert header.tables[0].cell(0,1).paragraphs[0].runs[0].font.size.pt==9


def test_direct_pdf_retains_table_fields_and_every_cell_once_per_page(tmp_path):
    result=build_export(fixture(),'pdf',tmp_path)
    pages=PdfReader(result['path']).pages
    assert len(pages)==2
    for index,page in enumerate(pages):
        text=page.extract_text()
        for label in ['BEFORE TABLE','HEADER COMPANY','HEADER DETAILS','AFTER TABLE']:
            assert text.count(label)==1
        assert f'Page {index+1}' in text
