from copy import deepcopy
import pytest
from alder.revision_formatting import validate_format_log, format_marks, native_format_revision, run_marks
from alder.document_revisions import WordRevisionWriter, WordRevisionReader, review_document, W


def fixture():
    return {'base': [{'type': 'em'}], 'records': [{'stamp': {'id': 'format_one', 'author': 'Format reviewer', 'date': '2026-10-29T05:00:00Z'}, 'status': 'pending', 'patches': [{'type': 'strong', 'before': None, 'after': {}}]}]}


def render(paragraph, node):
    from docx.shared import Pt
    run = paragraph.add_run(node['text'])
    for mark in node.get('marks', []):
        if mark['type'] == 'strong': run.bold = True
        elif mark['type'] == 'em': run.italic = True
        elif mark['type'] == 'text_style':
            if mark['attrs'].get('fontFamily'): run.font.name = mark['attrs']['fontFamily']
            if mark['attrs'].get('fontSize'): run.font.size = Pt(mark['attrs']['fontSize'])


def test_attribute_patches_keep_later_accepted_edits_when_earlier_edits_are_rejected():
    log = {'base': [{'type': 'text_style', 'attrs': {'fontFamily': 'Cambria', 'fontSize': 12}}], 'records': [
        {'stamp': {'id': 'family', 'author': 'First', 'date': ''}, 'status': 'pending', 'patches': [{'type': 'text_style', 'before': {'fontFamily': 'Cambria', 'fontSize': 12}, 'after': {'fontFamily': 'Arial', 'fontSize': 12}}]},
        {'stamp': {'id': 'size', 'author': 'Second', 'date': ''}, 'status': 'accepted', 'patches': [{'type': 'text_style', 'before': {'fontFamily': 'Arial', 'fontSize': 12}, 'after': {'fontFamily': 'Arial', 'fontSize': 18}}]},
    ]}
    before = deepcopy(log)
    assert format_marks(log) == [{'type': 'text_style', 'attrs': {'fontFamily': 'Arial', 'fontSize': 18}}]
    assert format_marks(log, original=True) == [{'type': 'text_style', 'attrs': {'fontFamily': 'Cambria', 'fontSize': 18}}]
    assert log == before
    log['records'][0]['status'] = 'rejected'
    assert format_marks(log) == format_marks(log, original=True)


def test_native_formatting_revision_preserves_previous_marks_and_authorship():
    from docx import Document
    doc = Document(); paragraph = doc.add_paragraph(); log = fixture()
    node = {'type': 'text', 'text': 'Reviewed format', 'marks': format_marks(log) + [{'type': 'revision_format', 'attrs': log}]}
    assert WordRevisionWriter().inline(paragraph, node, render)
    reader = WordRevisionReader(doc._element)
    properties = paragraph._p.find(f'{{{W}}}r/{{{W}}}rPr')
    revision = reader.formatting(properties)
    assert revision['attrs']['base'] == [{'type': 'em'}]
    assert revision['attrs']['records'][0]['stamp']['author'] == 'Format reviewer'
    assert revision['attrs']['records'][0]['patches'] == [{'type': 'strong', 'before': None, 'after': {}}]
    imported = {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': node['text'], 'marks': run_marks(properties) + [revision]}]}]}
    assert review_document(imported, 'original')['content'][0]['content'][0]['marks'] == [{'type': 'em'}]
    assert {m['type'] for m in review_document(imported)['content'][0]['content'][0]['marks']} == {'em', 'strong'}


def test_nested_text_and_format_revisions_remain_independent():
    from docx import Document
    doc = Document(); p = doc.add_paragraph(); log = fixture()
    node = {'type': 'text', 'text': 'New', 'marks': [{'type': 'revision_insertion', 'attrs': {'id': 'new', 'author': 'Writer', 'date': ''}}, {'type': 'revision_format', 'attrs': log}]}
    WordRevisionWriter().inline(p, node, render)
    reader = WordRevisionReader(doc._element)
    insertion = p._p.find(f'{{{W}}}ins')
    assert reader.mark(insertion)['attrs']['author'] == 'Writer'
    assert reader.formatting(insertion.find(f'{{{W}}}r/{{{W}}}rPr'))['attrs']['records'][0]['stamp']['author'] == 'Format reviewer'


def test_unsupported_overlapping_word_formatting_and_unsafe_metadata_are_rejected():
    log = fixture(); extra = deepcopy(log['records'][0]); extra['stamp']['id'] = 'second'
    log['records'].append(extra)
    with pytest.raises(ValueError, match='Overlapping formatting'): native_format_revision(log)
    extra['status'] = 'accepted'
    assert native_format_revision(log)[0]['id'] == 'format_one'
    log = fixture(); log['base'].append({'type': 'link', 'attrs': {'href': 'javascript:alert(1)'}})
    with pytest.raises(ValueError, match='Unsafe'): validate_format_log(log)
    log = fixture(); log['records'][0]['patches'][0]['type'] = 'revision_insertion'
    with pytest.raises(ValueError, match='Unsupported'): validate_format_log(log)
