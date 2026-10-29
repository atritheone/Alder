import copy
import base64
import io
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from pypdf import PdfReader

from alder.models import create_project, validate_project, document_text
from alder.publishing import build_export, import_document


def fixture():
    project = create_project(template='blank')
    project['settings'].update(documentKind='docx', includeTitle=False, footer=False)
    data = io.BytesIO()
    Image.new('RGB', (96, 48), '#23405e').save(data, format='PNG')
    attrs = {'src': 'data:image/png;base64,' + base64.b64encode(data.getvalue()).decode(),
             'alt': 'Retained description', 'title': 'Retained title', 'width': 48, 'decorative': True}
    picture = {'type': 'inline_image', 'attrs': attrs}
    project['book'] = {'version': 1, 'chapters': [{'id': 'inline', 'title': 'Document', 'text': '', 'include': True,
        'document': {'type': 'doc', 'content': [
            {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'BEFORE ', 'marks': [{'type': 'strong'}]},
                                               picture, {'type': 'text', 'text': ' AFTER'}]},
            {'type': 'paragraph', 'content': [copy.deepcopy(picture)]},
            {'type': 'image', 'attrs': copy.deepcopy(attrs)},
        ]}}]}
    project['book']['chapters'][0]['text'] = document_text(project['book']['chapters'][0]['document'])
    return project


@pytest.mark.parametrize('format', ['html', 'docx', 'epub'])
def test_inline_placement_and_description_roundtrip(tmp_path, format):
    project = fixture()
    result = build_export(project, format, tmp_path, {'validate': False})
    imported = import_document(Path(result['path']))
    nodes = imported['document']['content']
    assert [n['type'] for n in nodes] == ['paragraph', 'paragraph', 'image']
    assert [n['type'] for n in nodes[0]['content']] == ['text', 'inline_image', 'text']
    assert nodes[0]['content'][0]['marks'][0]['type'] == 'strong'
    assert nodes[1]['content'][0]['type'] == 'inline_image'
    for node in (nodes[0]['content'][1], nodes[1]['content'][0], nodes[2]):
        assert node['attrs']['alt'] == 'Retained description'
        assert node['attrs']['title'] == 'Retained title'
        assert node['attrs']['decorative'] is True
        assert node['attrs']['width'] == pytest.approx(48, abs=.001)
    assert document_text(nodes[0]) == 'BEFORE  AFTER'
    if format == 'docx':
        doc = Document(result['path'])
        assert len(doc.inline_shapes) == 3


def test_word_inline_image_without_alder_metadata_stays_in_paragraph(tmp_path):
    doc = Document()
    p = doc.add_paragraph('Before ')
    data = io.BytesIO()
    Image.new('RGB', (80, 40), 'navy').save(data, format='PNG')
    data.seek(0)
    p.add_run().add_picture(data)
    p.add_run(' after.')
    path = tmp_path / 'word.docx'
    doc.save(path)
    nodes = import_document(path)['document']['content']
    assert len(nodes) == 1
    assert [n['type'] for n in nodes[0]['content']] == ['text', 'inline_image', 'text']


def test_pdf_embeds_inline_images_without_moving_surrounding_words_to_other_lines(tmp_path):
    result = build_export(fixture(), 'pdf', tmp_path)
    doc = PdfReader(result['path'])
    assert len(doc.pages) == 1
    positions = {}
    def visit(text, cm, tm, font, size):
        for token in ('BEFORE', 'AFTER'):
            if token in text:
                positions[token] = (tm[4], tm[5])
    doc.pages[0].extract_text(visitor_text=visit)
    assert positions['BEFORE'][1] == positions['AFTER'][1]
    assert len(doc.pages[0].images) == 1  # One shared image resource, three placements.
    assert not any('unsupported' in w or 'separate block' in w for w in result['warnings'])


def test_inline_image_validation_and_stable_identity():
    project = fixture()
    nodes = project['book']['chapters'][0]['document']['content']
    for image in (nodes[0]['content'][1], nodes[1]['content'][0], nodes[2]):
        image['attrs']['src'] = ''
    project = validate_project(project)
    nodes = project['book']['chapters'][0]['document']['content']
    assert nodes[0]['content'][1]['attrs']['blockId']
    assert nodes[0]['content'][1]['attrs']['blockId'] != nodes[1]['content'][0]['attrs']['blockId']


def test_saved_inline_assets_survive_reopening_recovery_and_templates(tmp_path):
    from alder.store import Store
    from alder.templates import Templates
    store = Store(tmp_path)
    project = store.create('Inline asset', 'blank')
    asset = store.store_asset(project['id'], 'picture.png', b'collected image', 'image/png')
    project['assets'] = [asset]
    project['book'] = fixture()['book']
    content = project['book']['chapters'][0]['document']['content']
    for picture in (content[0]['content'][1], content[1]['content'][0], content[2]):
        picture['attrs'].update(assetId=asset['id'], src=f'/api/projects/{project["id"]}/assets/{asset["id"]}')
    saved = store.update(project['id'], project, project['revision'])
    assert Store(tmp_path).get(saved['id']) == saved
    recovered = store.recover_version(saved['id'], store.versions(saved['id'])[0]['sequence'])
    library = Templates(store)
    template = library.save(saved['id'], 'Inline template')
    created = library.create(template['id'], 'Copy')
    for copy_project in (recovered, created):
        picture = copy_project['book']['chapters'][0]['document']['content'][0]['content'][1]
        assert picture['type'] == 'inline_image'
        assert picture['attrs']['src'] == f'/api/projects/{copy_project["id"]}/assets/{copy_project["assets"][0]["id"]}'
        assert store.asset_path(copy_project['id'], picture['attrs']['assetId'])[0].read_bytes() == b'collected image'
