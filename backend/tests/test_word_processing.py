from pathlib import Path
import pytest
from docx import Document
from docx.oxml.ns import qn
from alder.models import create_project, validate_project
from alder.publishing import build_export, import_document, plain_text, render_html


def paragraph(text, **attrs):
    return {"type": "paragraph", "attrs": attrs, "content": [{"type": "text", "text": text}]}


def manuscript(nodes):
    project = create_project(template="blank")
    project["settings"].update(documentKind="docx", includeTitle=False, footer=False)
    project["book"] = {"version": 1, "chapters": [{"id": "chapter-test", "title": "Test", "text": "", "document": {"type": "doc", "content": nodes}, "include": True}]}
    return validate_project(project)


@pytest.mark.parametrize('alignment', ['left', 'center', 'right'])
def test_image_dimensions_alignment_and_description_round_trip(tmp_path, alignment):
    import base64
    import io
    from PIL import Image
    raw = io.BytesIO()
    Image.new('RGB', (160, 80), '#426985').save(raw, 'PNG')
    project = manuscript([paragraph('Before'), paragraph('After')])
    project['assets'] = [{'id': 'illustration', 'name': 'Illustration', 'mime': 'image/png', 'data': base64.b64encode(raw.getvalue()).decode()}]
    attributes = {'assetId': 'illustration', 'src': 'asset:illustration', 'width': 192, 'align': alignment, 'alt': 'A useful description', 'title': 'Image title'}
    project['book']['chapters'][0]['document']['content'].insert(1, {'type': 'image', 'attrs': attributes})
    for format in ('docx', 'html'):
        result = build_export(project, format, tmp_path)
        imported = import_document(Path(result['path']))
        images = [node for node in imported['document']['content'] if node['type'] == 'image']
        assert len(images) == 1
        for key in ('width', 'align', 'alt', 'title'):
            if key == 'width':
                assert images[0]['attrs'][key] == pytest.approx(attributes[key], abs=1/9525)
            else:
                assert images[0]['attrs'][key] == attributes[key]
        assert imported['text'].index('Before') < imported['text'].index('After')
    doc = Document(build_export(project, 'docx', tmp_path)['path'])
    assert doc.inline_shapes[0].width.inches == pytest.approx(2)
    assert doc.inline_shapes[0].height.inches == pytest.approx(1)
    pdf = build_export(project, 'pdf', tmp_path)
    from pypdf import PdfReader
    assert len(PdfReader(pdf['path']).pages[0].images) == 1


@pytest.mark.parametrize('width', [0, -1, 10001, 'invalid', True])
def test_invalid_image_width_rejected(width):
    from alder.models import validate_document, ValidationError
    with pytest.raises(ValidationError, match='Image width'):
        validate_document({'type': 'doc', 'content': [{'type': 'image', 'attrs': {'width': width}}]})


def test_paragraph_properties_round_trip_and_font_runs(tmp_path):
    p = paragraph("Paragraph", fontFamily="Cambria", fontSize=13.5, leftIndent=10, rightIndent=24, firstLineIndent=-5, spaceBefore=9, spaceAfter=15, lineHeight=1.4)
    p["content"].append({"type": "text", "text": " coloured", "marks": [{"type": "text_style", "attrs": {"fontFamily": "Arial", "fontSize": 11.5, "color": "#123456"}}]})
    project = manuscript([p])
    html = render_html(project)
    assert "margin-right:24pt" in html and "margin-top:9pt" in html
    result = build_export(project, "docx", tmp_path)
    doc = Document(result["path"])
    assert doc.paragraphs[0].paragraph_format.right_indent.pt == 24
    imported = import_document(Path(result["path"]))["document"]["content"][0]
    for key in ("leftIndent", "rightIndent", "firstLineIndent", "spaceBefore", "spaceAfter", "lineHeight"):
        assert imported["attrs"][key] == p["attrs"][key]
    assert imported["content"][0]["marks"][0]["attrs"]["fontFamily"] == "Cambria"
    assert imported["content"][0]["marks"][0]["attrs"]["fontSize"] == 13.5
    assert imported["content"][1]["marks"][0]["attrs"] == p["content"][1]["marks"][0]["attrs"]


def test_docx_lists_use_numbering_and_reopen_as_nested_lists(tmp_path):
    listing = {"type": "ordered_list", "attrs": {"order": 4}, "content": [
        {"type": "list_item", "content": [paragraph("First"), {"type": "bullet_list", "content": [{"type": "list_item", "content": [paragraph("Nested")]}]}]},
        {"type": "list_item", "content": [paragraph("Second")]},
    ]}
    result = build_export(manuscript([listing]), "docx", tmp_path)
    doc = Document(result["path"])
    assert [p.text for p in doc.paragraphs] == ["First", "Nested", "Second"]
    assert all(p._p.pPr.numPr is not None for p in doc.paragraphs)
    imported = import_document(Path(result["path"]))["document"]["content"][0]
    assert imported["type"] == "ordered_list" and imported["attrs"]["order"] == 4
    assert imported["content"][0]["content"][1]["type"] == "bullet_list"
    assert imported["content"][1]["content"][0]["content"][0]["text"] == "Second"
    from pypdf import PdfReader
    pdf = build_export(manuscript([listing]), "pdf", tmp_path)
    text = "\n".join(page.extract_text() for page in PdfReader(pdf["path"]).pages)
    assert "\u2022 Nested" in text and "â" not in text


def test_docx_table_spans_and_widths_round_trip(tmp_path):
    def cell(text, **attrs):
        return {"type": "table_cell", "attrs": attrs, "content": [paragraph(text)]}
    table = {"type": "table", "content": [
        {"type": "table_row", "content": [cell("Tall", rowspan=2, colwidth=[160]), cell("Top", colwidth=[240])]},
        {"type": "table_row", "content": [cell("Bottom", colwidth=[240])]},
        {"type": "table_row", "content": [cell("Wide", colspan=2, colwidth=[160, 240])]},
    ]}
    result = build_export(manuscript([table]), "docx", tmp_path)
    assert not any("row spans" in w for w in result["warnings"])
    imported = import_document(Path(result["path"]))["document"]["content"][0]
    assert imported["content"][0]["content"][0]["attrs"]["rowspan"] == 2
    assert imported["content"][0]["content"][0]["attrs"]["colwidth"] == [160]
    assert imported["content"][2]["content"][0]["attrs"]["colspan"] == 2
    text = plain_text(imported)
    assert all(text.count(word) == 1 for word in ("Tall", "Top", "Bottom", "Wide"))
    from pypdf import PdfReader
    pdf = build_export(manuscript([table]), "pdf", tmp_path)
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(pdf["path"]).pages)
    assert all(pdf_text.count(word) == 1 for word in ("Tall", "Top", "Bottom", "Wide"))
    assert not any("flattened" in warning or "equal columns" in warning for warning in pdf["warnings"])


def test_plain_text_save_warns_only_when_rich_content_is_lost(tmp_path):
    plain = manuscript([paragraph("Plain document.")])
    assert not any("formatting" in warning for warning in build_export(plain, "txt", tmp_path)["warnings"])
    rich = manuscript([paragraph("Formatted document.", align="right")])
    assert any("formatting" in warning for warning in build_export(rich, "txt", tmp_path)["warnings"])


def test_plain_text_save_preserves_authored_spaces_and_empty_paragraphs(tmp_path):
    from alder.models import text_document
    text = "First line\n\nTrailing spaces  \n\n"
    result = build_export(manuscript(text_document(text)["content"]), "txt", tmp_path)
    assert Path(result["path"]).read_text("utf-8") == text


def test_heading_sizes_match_editor_contract_and_do_not_override_authored_size(tmp_path):
    from alder.publishing import DOCUMENT_FORMAT
    headings = [{"type": "heading", "attrs": {"level": level}, "content": [{"type": "text", "text": f"Heading {level}"}]} for level in range(1, 7)]
    headings.append({"type": "heading", "attrs": {"level": 1, "fontSize": 17}, "content": [{"type": "text", "text": "Custom heading"}]})
    project = manuscript(headings)
    doc = Document(build_export(project, 'docx', tmp_path)['path'])
    sizes = [p.runs[0].font.size.pt for p in doc.paragraphs]
    assert sizes == [*DOCUMENT_FORMAT['headingSizes'], 17]
    fonts = doc.styles['Heading 1'].element.rPr.rFonts
    assert fonts.get(qn('w:ascii')) == 'Cambria'
    assert not any('theme' in key.lower() for key in fonts.attrib)
    from pypdf import PdfReader
    actual = {}
    for page in PdfReader(build_export(project, 'pdf', tmp_path)['path']).pages:
        page.extract_text(visitor_text=lambda text, cm, tm, font, size: actual.update({text.strip(): size}) if text.strip() else None)
    assert actual['Heading 1'] == 24 and actual['Custom heading'] == 17


def test_pdf_uses_authored_margins_without_an_empty_section_heading(tmp_path):
    from pypdf import PdfReader
    project = manuscript([paragraph('Margin probe')])
    project['settings']['marginMm'] = 20
    result = build_export(project, 'pdf', tmp_path)
    page = PdfReader(result['path']).pages[0]
    positions = []
    page.extract_text(visitor_text=lambda text, cm, tm, font, size: positions.append((cm[4] + tm[4], cm[5] + tm[5])) if 'Margin probe' in text else None)
    margin = 20 * 72 / 25.4
    x, y = positions[0]
    assert abs(x - margin) < 0.1
    assert margin <= float(page.mediabox.height) - y <= margin + 20


@pytest.mark.parametrize('mode', ['exact', 'atLeast'])
def test_paragraph_flow_round_trip_preserves_on_and_off_overrides(tmp_path, mode):
    flags = ('keepWithNext', 'keepTogether', 'widowOrphan', 'pageBreakBefore')
    project = manuscript([
        paragraph('Keep these rules', lineSpacingMode=mode, lineSpacing=19.5, **dict.fromkeys(flags, True)),
        paragraph('Turn the rules off', lineSpacingMode='multiple', lineHeight=1.25, **dict.fromkeys(flags, False)),
    ])
    result = build_export(project, 'docx', tmp_path)
    imported = import_document(Path(result['path']))['document']['content']
    for key in flags:
        assert imported[0]['attrs'][key] is True
        assert imported[1]['attrs'][key] is False
    assert imported[0]['attrs']['lineSpacingMode'] == mode
    assert imported[0]['attrs']['lineSpacing'] == 19.5
    assert imported[1]['attrs']['lineSpacingMode'] == 'multiple'
    assert imported[1]['attrs']['lineHeight'] == 1.25
    assert 'lineSpacing' not in imported[1]['attrs']
    html = render_html(project)
    assert 'break-inside:avoid' in html and 'break-inside:auto' in html
    assert 'widows:2' in html and 'widows:1' in html
    html_result = build_export(project, 'html', tmp_path)
    from_html = import_document(Path(html_result['path']))['document']['content']
    assert from_html[0]['attrs']['lineSpacingMode'] == mode
    assert from_html[0]['attrs']['lineSpacing'] == 19.5
    assert from_html[1]['attrs']['lineHeight'] == 1.25
    for key in flags:
        assert from_html[0]['attrs'][key] is True
        assert from_html[1]['attrs'][key] is False


def test_pdf_page_break_before_and_oversized_kept_paragraph(tmp_path):
    from pypdf import PdfReader
    project = manuscript([
        paragraph('First page', pageBreakBefore=True),
        paragraph('Second page', pageBreakBefore=True),
        {**paragraph('', keepTogether=True, widowOrphan=True, lineSpacingMode='exact', lineSpacing=18),
         'content': [part for index in range(90) for part in ([{'type': 'hard_break'}] if index else []) + [{'type': 'text', 'text': f'KEPT{index:03}'}]]},
    ])
    pages = PdfReader(build_export(project, 'pdf', tmp_path)['path']).pages
    text = [page.extract_text() for page in pages]
    assert 'First page' in text[0] and 'Second page' not in text[0]
    assert 'Second page' in text[1]
    assert all('\n'.join(text).count(f'KEPT{index:03}') == 1 for index in range(90))
    assert len(pages) >= 4


@pytest.mark.parametrize('attrs', [{'lineSpacing': 0}, {'lineSpacing': float('nan')}, {'lineSpacingMode': 'unknown'}, {'keepTogether': 'false'}])
def test_invalid_paragraph_flow_is_rejected(attrs):
    from alder.models import ValidationError
    with pytest.raises(ValidationError):
        manuscript([paragraph('Invalid flow', **attrs)])


@pytest.mark.parametrize('marker, start, label', [('decimal', 4, '4'), ('lower-alpha', 27, 'aa'), ('upper-alpha', 2, 'B'), ('lower-roman', 9, 'ix'), ('upper-roman', 12, 'XII'), ('disc', 1, '\u2022'), ('circle', 1, '\u25cb'), ('square', 1, '\u25aa')])
def test_editable_list_markers_round_trip_and_pdf(tmp_path, marker, start, label):
    from pypdf import PdfReader
    ordered = marker not in ('disc', 'circle', 'square')
    project = manuscript([{'type': 'ordered_list' if ordered else 'bullet_list', 'attrs': {'order': start, 'listStyle': marker}, 'content': [{'type': 'list_item', 'content': [paragraph('List item')]}]}])
    docx = build_export(project, 'docx', tmp_path)
    imported = import_document(Path(docx['path']))['document']['content'][0]
    assert imported['attrs']['listStyle'] == marker
    if ordered: assert imported['attrs']['order'] == start
    html = build_export(project, 'html', tmp_path)
    imported_html = import_document(Path(html['path']))['document']['content'][0]
    assert imported_html['attrs']['listStyle'] == marker
    text = '\n'.join(page.extract_text() for page in PdfReader(build_export(project, 'pdf', tmp_path)['path']).pages)
    assert (label + ('. ' if ordered else ' ') + 'List item') in text


def test_word_implicit_flow_defaults_and_document_defaults_are_resolved():
    from alder.docx_support import paragraph_attributes
    from docx.oxml import OxmlElement
    doc = Document()
    p = doc.add_paragraph('Implicit Word defaults')
    attrs = paragraph_attributes(p)
    assert attrs['widowOrphan'] is True and attrs['keepTogether'] is False
    assert attrs['keepWithNext'] is False and attrs['pageBreakBefore'] is False
    assert attrs['lineHeight'] == 1.15 and attrs['spaceAfter'] == 10
    props = doc.styles.element.find('.//' + qn('w:pPrDefault') + '/' + qn('w:pPr'))
    prop = OxmlElement('w:widowControl')
    prop.set(qn('w:val'), '0')
    props.append(prop)
    assert paragraph_attributes(p)['widowOrphan'] is False
    p.paragraph_format.widow_control = True
    assert paragraph_attributes(p)['widowOrphan'] is True
